"""«Mühasibatlıq» kabinet bölməsi (sahib 2026-10-18) — RBAC, aylıq plan ↔ fakt, əhatə, drill-down, ixrac, sorğu büdcəsi.

Kilidlənən qaydalar:

* menyu + panel + ixrac mühasib (`accounting.view`), RİM rəhbəri, qəyyum, rektor, superadmin üçündür; müəllim / dekan /
  HR / tələbə görmür, fraqment 403, ixrac 403;
* fakt = sayılan jurnal dərslərinin saatı: dərsin öz müəllimi, boşdursa açılışın müəllimi; passiv açılış, bütün
  qeydiyyatları çıxarılmış («keçirilməyən») açılış və başqa semestrin dərsləri SAYILMIR;
* plan = paylanmış (distributed) tapşırığın lecture/seminar/lab bölgüsü; qaralama tapşırıq sayılmır; semestr planı
  aylara iş günlərinə mütənasib bölünür və cəm dəqiqdir;
* açarı alan UNIT rolu yalnız öz alt-ağacının müəllimlərini görür; drill-down əhatədən kənar müəllimi açmır;
* sorğu sayı müəllim sayından asılı deyil; panel tələbə adını göstərmir.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

from django.apps import apps as django_apps
from django.contrib.auth import get_user_model
from django.db import connection
from django.test import Client, RequestFactory, SimpleTestCase, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from apps.accounts.models import UserProfile
from apps.accounts.services.accounting import constants as C
from apps.accounts.services.accounting.panel import build_panel, resolve
from apps.accounts.services.accounting.periods import Semester, apportion, months, semesters, teaching_days
from apps.accounts.services.accounting.rows import apply_filters, plan_status
from apps.organizations.models import AcademicPeriod, Membership, Organization, OrgUnit, Role
from apps.registrar.models import CourseOffering, Enrollment, Lesson, Subject
from apps.workload.constants import Activity, Season, TaskStatus
from apps.workload.models import TeacherAssignment, TeacherWorkloadProfile, TeachingTask, TeachingTaskRow
from core.constants import AcademicPeriodType, OrganizationType, OrgUnitType
from core.rls import bypass_rls

User = get_user_model()
SECTION = "accounting"
YEAR = "2026/2027"
SEM = f"{YEAR}|fall"
OCT = "2026-10"


class _Period:
    def __init__(self, pk, start, end, exam=None, current=False, year=YEAR):
        self.pk, self.start_date, self.end_date, self.exam_session_start = pk, start, end, exam
        self.is_current, self.academic_year = current, year


class PeriodRuleTest(SimpleTestCase):
    """Planın aylara bölgüsü — sənədləşdirilmiş qayda (iş günlərinə mütənasib, cəm dəqiq)."""

    def _semester(self) -> Semester:
        return semesters([_Period(1, date(2026, 9, 15), date(2027, 1, 31), exam=date(2027, 1, 5), current=True)])[0]

    def test_window_stops_before_exam_session_and_months_include_exam_month(self):
        semester = self._semester()
        self.assertEqual((semester.key, semester.season), (SEM, "fall"))
        self.assertEqual(months(semester), ["2026-09", "2026-10", "2026-11", "2026-12", "2027-01"])
        self.assertEqual(
            teaching_days(semester), {"2026-09": 12, "2026-10": 22, "2026-11": 21, "2026-12": 23, "2027-01": 2}
        )

    def test_apportion_is_proportional_and_sums_exactly(self):
        shares = apportion(60, teaching_days(self._semester()))
        self.assertEqual(shares, {"2026-09": 9.0, "2026-10": 16.5, "2026-11": 15.7, "2026-12": 17.2, "2027-01": 1.6})
        self.assertAlmostEqual(sum(shares.values()), 60.0, places=6)
        self.assertEqual(apportion(0, {"2026-09": 3}), {"2026-09": 0.0})
        self.assertEqual(apportion(10, {}), {})

    def test_season_follows_workload_rule(self):
        spring = semesters([_Period(2, date(2027, 2, 10), date(2027, 6, 20))])[0]
        summer = semesters([_Period(3, date(2027, 7, 1), date(2027, 7, 30))])[0]
        self.assertEqual((spring.season, summer.season), ("spring", "summer"))

    def test_plan_status_band(self):
        self.assertEqual(plan_status(10, 9), C.STATUS_ON_PLAN)
        self.assertEqual(plan_status(10, 11), C.STATUS_ON_PLAN)
        self.assertEqual(plan_status(10, 8.9), C.STATUS_UNDER)
        self.assertEqual(plan_status(10, 11.1), C.STATUS_OVER)
        self.assertEqual(plan_status(0, 2), C.STATUS_NO_PLAN)
        self.assertEqual(plan_status(0, 0), C.STATUS_IDLE)


@override_settings(UNIVERSITY_MODE=True)
class AccountingSectionTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        owner = User.objects.create_user("ac_owner", "ac_owner@qku.edu.az", "pw")
        with bypass_rls():
            cls.org = Organization.objects.create(
                name="AC Univ", slug="ac-univ", org_type=OrganizationType.UNIVERSITY, owner=owner, status="active"
            )
            cls.f1 = cls._unit("Mühəndislik", OrgUnitType.FACULTY)
            cls.f2 = cls._unit("Humanitar", OrgUnitType.FACULTY)
            cls.c1 = cls._unit("İnformatika kafedrası", OrgUnitType.CHAIR, cls.f1)
            cls.c2 = cls._unit("Tarix kafedrası", OrgUnitType.CHAIR, cls.f2)
            cls.g1 = cls._unit("101A", OrgUnitType.GROUP, cls.f1)
            cls.fall = AcademicPeriod.objects.create(
                organization=cls.org,
                name="Payız",
                period_type=AcademicPeriodType.SEMESTER,
                academic_year=YEAR,
                start_date="2026-09-15",
                end_date="2027-01-31",
                exam_session_start="2027-01-05",
                is_current=True,
            )
            cls.spring = AcademicPeriod.objects.create(
                organization=cls.org,
                name="Yaz",
                period_type=AcademicPeriodType.SEMESTER,
                academic_year=YEAR,
                start_date="2027-02-10",
                end_date="2027-06-20",
            )
            cls.accountant = cls._member("ac_acc", "accountant")
            cls.rim = cls._member("ac_rim", "ikt_rehber")
            cls.trustee = cls._member("ac_trustee", "trustee")
            cls.rector = cls._member("ac_rector", "rector")
            cls.dean = cls._member("ac_dean", "dean", unit=cls.f1)
            cls.hr = cls._member("ac_hr", "hr")
            cls.student = cls._member("ac_student", "student", first="Tələbəzadə", last="Gizlinov")
            cls.superadmin = User.objects.create_superuser("ac_su", "ac_su@qku.edu.az", "pw")

            cls.t1 = cls._teacher("ac_t1", "Aysel", "Quliyeva", cls.c1)
            cls.t2 = cls._teacher("ac_t2", "Rauf", "Məmmədov", cls.c1)
            cls.t3 = cls._teacher("ac_t3", "Elçin", "Həsənov", cls.c2)
            TeacherWorkloadProfile.objects.create(
                organization=cls.org, teacher=cls.t1, academic_year=YEAR, staff_fraction="1.00"
            )
            TeacherWorkloadProfile.objects.create(
                organization=cls.org, teacher=cls.t2, academic_year=YEAR, staff_fraction="0.50"
            )

            # Plan: paylanmış tapşırıq (sayılır) + qaralama tapşırıq (sayılmır).
            task = TeachingTask.objects.create(
                organization=cls.org, chair=cls.c1, academic_year=YEAR, status=TaskStatus.DISTRIBUTED
            )
            row = TeachingTaskRow.objects.create(
                organization=cls.org,
                task=task,
                season=Season.FALL,
                subject_text="Alqoritmlər",
                lecture_total=60,
                seminar_total=30,
                lab_total=10,
                total_hours=100,
            )
            cls._assign(row, cls.t1, Activity.LECTURE, 60)
            cls._assign(row, cls.t1, Activity.LAB, 10)
            cls._assign(row, cls.t2, Activity.SEMINAR, 30)
            spring_row = TeachingTaskRow.objects.create(
                organization=cls.org,
                task=task,
                season=Season.SPRING,
                subject_text="Yaz",
                lecture_total=40,
                total_hours=40,
            )
            cls._assign(spring_row, cls.t1, Activity.LECTURE, 40)  # yaz semestri — payıza düşmür
            draft = TeachingTask.objects.create(
                organization=cls.org, chair=cls.c2, academic_year=YEAR, status=TaskStatus.DRAFT
            )
            draft_row = TeachingTaskRow.objects.create(
                organization=cls.org,
                task=draft,
                season=Season.FALL,
                subject_text="Tarix",
                lecture_total=20,
                total_hours=20,
            )
            cls._assign(draft_row, cls.t3, Activity.LECTURE, 20)

            # Fakt.
            s1 = Subject.objects.create(organization=cls.org, code="AC-1", name="Alqoritmlər")
            s2 = Subject.objects.create(organization=cls.org, code="AC-2", name="Boş fənn")
            s3 = Subject.objects.create(organization=cls.org, code="AC-3", name="Passiv fənn")
            s4 = Subject.objects.create(organization=cls.org, code="AC-4", name="Tarix")
            s5 = Subject.objects.create(organization=cls.org, code="AC-5", name="Yaz fənni")
            held = cls._offering(s1, cls.t1)
            dropped = cls._offering(s2, cls.t1, status=Enrollment.Status.DROPPED)
            inactive = cls._offering(s3, cls.t1, active=False)
            history = cls._offering(s4, cls.t3)
            spring_offering = cls._offering(s5, cls.t1, period=cls.spring)
            for day in (5, 12, 19):  # 3 mühazirə (dərsin öz müəllimi boş → açılışın müəllimi t1)
                cls._lesson(held, date(2026, 10, day), "lecture")
            for day in (6, 7, 13, 14, 20):  # bölünmüş fənn: seminarı t2 keçir
                cls._lesson(held, date(2026, 10, day), "seminar", instructor=cls.t2)
            cls._lesson(held, date(2026, 11, 3), "lab")
            cls._lesson(dropped, date(2026, 10, 8), "lecture")  # keçirilməyən fənn («Jurnal təmizliyi» qərarı)
            cls._not_held(dropped)
            cls._lesson(inactive, date(2026, 10, 9), "lecture")  # passiv jurnal
            cls._lesson(history, date(2026, 10, 21), "lecture")  # t3: planı yoxdur
            cls._lesson(spring_offering, date(2027, 3, 2), "lecture")  # başqa semestr

    # ── fikstur köməkçiləri ────────────────────────────────────────────────
    @classmethod
    def _unit(cls, name, kind, parent=None):
        return OrgUnit.objects.create(
            organization=cls.org, name=name, slug=f"ac-{abs(hash(name)) % 10**8}", unit_type=kind, parent=parent
        )

    @classmethod
    def _user(cls, username, first="", last=""):
        user = User.objects.create_user(username, f"{username}@qku.edu.az", "pw", first_name=first, last_name=last)
        UserProfile.objects.get_or_create(user=user)
        return user

    @classmethod
    def _member(cls, username, role, *, unit=None, first="", last=""):
        user = cls._user(username, first, last)
        Membership.objects.create(
            user=user,
            organization=cls.org,
            role=cls.org.roles.get(name=role),
            scope_unit=unit,
            is_primary=True,
            is_active=True,
        )
        return user

    @classmethod
    def _teacher(cls, username, first, last, chair):
        return cls._member(username, "teacher", unit=chair, first=first, last=last)

    @classmethod
    def _assign(cls, row, teacher, activity, hours):
        TeacherAssignment.objects.create(organization=cls.org, row=row, teacher=teacher, activity=activity, hours=hours)

    @classmethod
    def _offering(cls, subject, instructor, *, active=True, status=Enrollment.Status.ENROLLED, period=None):
        offering = CourseOffering.objects.create(
            organization=cls.org,
            subject=subject,
            period=period or cls.fall,
            group=cls.g1,
            instructor=instructor,
            lesson_hours=30,
            is_active=active,
        )
        Enrollment.objects.create(organization=cls.org, student=cls.student, offering=offering, status=status)
        return offering

    @classmethod
    def _not_held(cls, offering, category="not_held"):
        JournalHygieneDecision = django_apps.get_model("registrar", "JournalHygieneDecision")
        for enrollment in offering.enrollments.all():
            JournalHygieneDecision.objects.create(
                organization=cls.org,
                enrollment=enrollment,
                batch=uuid.uuid4(),
                category=category,
                reason_code="no_exam_evidence",
                from_status=Enrollment.Status.ENROLLED,
                to_status=Enrollment.Status.DROPPED,
                evidence={},
                evidence_digest="0" * 64,
                reason="Fənn qrupda keçirilməyib — test fiksturu.",
            )

    @classmethod
    def _lesson(cls, offering, day, kind, *, instructor=None, hours=2):
        Lesson.objects.create(
            organization=cls.org, offering=offering, date=day, kind=kind, hours=hours, instructor=instructor
        )

    def _client(self, user):
        client = Client()
        client.force_login(user)
        session = client.session
        session["active_organization"] = self.org.slug
        session.save()
        return client

    def _request(self, user, **params):
        request = RequestFactory().get("/accounts/profile/", {"section": SECTION, **params})
        request.user = user
        request.organization = self.org
        request.is_view_as = False
        return request

    def _rows(self, user=None, **params):
        params.setdefault("ac_semester", SEM)
        params.setdefault("ac_month", OCT)
        resolved = resolve(self._request(user or self.accountant, **params))
        return resolved, {row["username"]: row for row in apply_filters(resolved.rows, resolved.filters)}

    # ── RBAC ──────────────────────────────────────────────────────────────
    def test_visible_for_accountant_rim_trustee_rector_and_superadmin(self):
        for user in (self.accountant, self.rim, self.trustee, self.rector, self.superadmin):
            with self.subTest(user=user.username):
                client = self._client(user)
                page = client.get(reverse("accounts:profile"), {"section": SECTION, "ac_month": OCT})
                self.assertEqual(page.status_code, 200)
                self.assertContains(page, f'data-section="{SECTION}"')
                self.assertContains(page, "data-ac-root")
                self.assertContains(page, "accounting.css")
                fragment = client.get(reverse("accounts:profile_section_fragment", args=[SECTION]))
                self.assertEqual(fragment.status_code, 200)
                self.assertIn("data-ac-root", fragment.json()["html"])
                self.assertEqual(client.get(reverse("accounts:accounting_export")).status_code, 200)

    def test_hidden_and_refused_for_teacher_dean_hr_and_student(self):
        for user in (self.t1, self.dean, self.hr, self.student):
            with self.subTest(user=user.username):
                client = self._client(user)
                page = client.get(reverse("accounts:profile"), {"section": SECTION})
                self.assertNotContains(page, f'data-section="{SECTION}"')
                self.assertNotContains(page, "data-ac-root")
                fragment = client.get(reverse("accounts:profile_section_fragment", args=[SECTION]))
                self.assertEqual(fragment.status_code, 403)
                self.assertEqual(client.get(reverse("accounts:accounting_export")).status_code, 403)
                self.assertFalse(build_panel(self._request(user))["has_access"])

    # ── Rəqəmlər ──────────────────────────────────────────────────────────
    def test_october_held_vs_planned(self):
        resolved, by = self._rows()
        self.assertEqual(resolved.filters.month, OCT)
        self.assertEqual(set(by), {"ac_t1", "ac_t2", "ac_t3"})
        t1, t2, t3 = by["ac_t1"], by["ac_t2"], by["ac_t3"]
        # t1: 3 mühazirə × 2 saat; boş / passiv fənnin dərsləri və t2-nin seminarları DAXİL DEYİL.
        self.assertEqual((t1["held"], t1["kinds"], t1["lessons"]), (6, {"lecture": 6, "seminar": 0, "lab": 0}, 3))
        # Plan: mühazirə 60 → oktyabr 16,5; lab 10 → 2,7 (22 / 80 iş günü).
        self.assertEqual(t1["planned"], 19.2)
        self.assertEqual((t1["diff"], t1["pct"], t1["status"]), (-13.2, 31, C.STATUS_UNDER))
        self.assertEqual((t1["planned_sem"], t1["held_sem"]), (70, 8))  # noyabr lab daxil, yaz yükü xaric
        self.assertEqual(t1["staff_fraction"], 1.0)
        self.assertEqual((t1["faculty"], t1["chair"]), ("Mühəndislik", "İnformatika kafedrası"))
        # t2: bölünmüş fənnin seminarları (Lesson.instructor) — 5 × 2 = 10 saat, plan 30 → 8,2.
        self.assertEqual((t2["held"], t2["planned"], t2["status"]), (10, 8.2, C.STATUS_OVER))
        # t3: qaralama tapşırıq plan sayılmır.
        self.assertEqual((t3["held"], t3["planned"], t3["status"]), (2, 0.0, C.STATUS_NO_PLAN))

    def test_month_plans_sum_to_semester_plan(self):
        _resolved, by = self._rows()
        months_ = by["ac_t1"]["months"]
        self.assertEqual(list(months_), ["2026-09", "2026-10", "2026-11", "2026-12", "2027-01"])
        self.assertAlmostEqual(sum(cell["planned"] for cell in months_.values()), 70.0, places=6)
        self.assertEqual(months_["2026-11"]["kinds"]["lab"], 2)

    def test_whole_semester_and_filters(self):
        _resolved, by = self._rows(ac_month="all")
        self.assertEqual((by["ac_t1"]["planned"], by["ac_t1"]["held"]), (70.0, 8))
        cases = {
            "ac_faculty": (str(self.f1.pk), {"ac_t1", "ac_t2"}),
            "ac_chair": (str(self.c2.pk), {"ac_t3"}),
            "ac_q": ("məmmədov", {"ac_t2"}),
            "ac_deviation": ("1", {"ac_t1", "ac_t2", "ac_t3"}),
        }
        for param, (value, expected) in cases.items():
            with self.subTest(param=param):
                _resolved, by = self._rows(**{param: value})
                self.assertEqual(set(by), expected)
        resolved, by = self._rows(ac_semester="bogus", ac_month="1999-01", ac_faculty="x", ac_sort="evil")
        self.assertEqual(
            (resolved.filters.semester, resolved.filters.faculty, resolved.filters.sort), (SEM, "", "name")
        )
        self.assertIn(resolved.filters.month, months(resolved.semester))

    def test_spring_semester_isolated(self):
        _resolved, by = self._rows(ac_semester=f"{YEAR}|spring", ac_month="all")
        self.assertEqual((by["ac_t1"]["planned"], by["ac_t1"]["held"]), (40.0, 2))
        self.assertNotIn("ac_t2", by)

    def test_unit_scoped_holder_sees_only_its_subtree(self):
        dean_role = self.org.roles.get(name="dean")
        Role.objects.filter(pk=dean_role.pk).update(permissions=[*dean_role.permissions, "accounting.view"])
        resolved, by = self._rows(self.dean)
        self.assertTrue(resolved.actor.has_access)
        self.assertEqual(set(by), {"ac_t1", "ac_t2"})
        detail = build_panel(self._request(self.dean, ac_month=OCT, ac_teacher=str(self.t3.pk)))["detail"]
        self.assertFalse(detail["found"])  # əhatədən kənar müəllim açılmır

    def test_scoped_export_refuses_out_of_scope_teacher(self):
        from openpyxl import load_workbook

        dean_role = self.org.roles.get(name="dean")
        Role.objects.filter(pk=dean_role.pk).update(permissions=[*dean_role.permissions, "accounting.view"])
        client = self._client(self.dean)
        url = reverse("accounts:accounting_export")
        self.assertEqual(client.get(url, {"ac_month": OCT, "ac_teacher": self.t3.pk}).status_code, 404)
        self.assertEqual(client.get(url, {"ac_month": OCT, "ac_teacher": "999999"}).status_code, 404)
        own = client.get(url, {"ac_month": OCT, "ac_teacher": self.t1.pk})
        self.assertEqual(own.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(own.content)).sheetnames[-1], "Dərslər")
        listing = load_workbook(io.BytesIO(client.get(url, {"ac_month": OCT}).content))
        self.assertEqual(
            {row[1] for row in list(listing.worksheets[0].iter_rows(values_only=True))[2:]}, {"ac_t1", "ac_t2"}
        )

    def test_lessons_before_a_group_merge_still_count(self):
        """Noyabrda qrup birləşməsi / köçürmə (``dropped`` + ``superseded_by``) oktyabr saatlarını silmir."""
        subject = Subject.objects.create(organization=self.org, code="AC-M", name="Birləşən fənn")
        old_offering = self._offering(subject, self.t3)
        new_offering = CourseOffering.objects.create(
            organization=self.org,
            subject=subject,
            period=self.fall,
            group=self._unit("102A", OrgUnitType.GROUP, self.f1),
            instructor=self.t3,
        )
        successor = Enrollment.objects.create(organization=self.org, student=self.student, offering=new_offering)
        old_offering.enrollments.update(status=Enrollment.Status.DROPPED, superseded_by=successor)
        self._lesson(old_offering, date(2026, 10, 27), "seminar")
        _resolved, by = self._rows()
        self.assertEqual(by["ac_t3"]["held"], 4)  # 2 (Tarix) + 2 (birləşmədən əvvəlki seminar)
        self._not_held(old_offering, category="subject_exempt")  # «fəndən azad» qərarı — artıq sayılmır
        _resolved, by = self._rows()
        self.assertEqual(by["ac_t3"]["held"], 2)

    def test_accountant_with_only_two_keys_renders_section_and_export(self):
        role = self.org.roles.get(name="accountant")
        self.assertEqual(set(role.permissions) - {"application.create"}, {"accounting.view", "retake.payment_verify"})
        Role.objects.filter(pk=role.pk).update(permissions=["accounting.view", "retake.payment_verify"])
        client = self._client(self.accountant)
        html = client.get(reverse("accounts:profile_section_fragment", args=[SECTION]), {"ac_month": OCT}).json()[
            "html"
        ]
        self.assertIn("Quliyeva Aysel", html)
        self.assertIn("İnformatika kafedrası", html)  # kafedra seçimi və sətir yerləşdirməsi servisdən
        page = client.get(reverse("accounts:profile"), {"section": SECTION})
        self.assertNotContains(page, 'data-section="org-structure-tree"')
        self.assertNotContains(page, 'data-section="people-teachers"')
        self.assertEqual(client.get(reverse("accounts:accounting_export"), {"ac_month": OCT}).status_code, 200)

    # ── Panel ─────────────────────────────────────────────────────────────
    def test_panel_kpis_table_and_no_student_pii(self):
        panel = build_panel(self._request(self.accountant, ac_month=OCT))
        kpis = {tile["label"]: tile["value"] for tile in panel["kpis"]}
        self.assertEqual(kpis["Keçirilmiş saat"], "18")
        self.assertEqual(kpis["Plan"], "27,4")
        self.assertEqual((kpis["Plandan çox"], kpis["Plandan az"]), (1, 1))
        self.assertEqual(panel["table"]["total"], 3)
        html = (
            self._client(self.accountant)
            .get(reverse("accounts:profile_section_fragment", args=[SECTION]), {"ac_month": OCT})
            .json()["html"]
        )
        self.assertIn("Quliyeva Aysel", html)
        self.assertNotIn("Gizlinov", html)
        self.assertNotIn("ac_student", html)
        self.assertNotIn("style=", html.split("data-ac-root", 1)[1])

    def test_teacher_drill_down(self):
        panel = build_panel(self._request(self.accountant, ac_month=OCT, ac_teacher=str(self.t1.pk)))
        detail = panel["detail"]
        self.assertTrue(detail["found"])
        self.assertEqual(len(detail["months"]), 5)
        self.assertEqual(detail["lessons_total"], 3)
        self.assertEqual({item["kind"] for item in detail["lessons"]}, {"lecture"})
        self.assertEqual([(item["subject"], item["group"]) for item in detail["breakdown"]], [("Alqoritmlər", "101A")])
        html = (
            self._client(self.accountant)
            .get(
                reverse("accounts:profile_section_fragment", args=[SECTION]),
                {"ac_month": OCT, "ac_teacher": self.t1.pk},
            )
            .json()["html"]
        )
        self.assertIn("05.10.2026", html)
        self.assertNotIn("Gizlinov", html)
        missing = build_panel(self._request(self.accountant, ac_teacher="999999"))["detail"]
        self.assertFalse(missing["found"])

    # ── İxrac ─────────────────────────────────────────────────────────────
    def test_export_is_xlsx_filtered_and_audited(self):
        from openpyxl import load_workbook

        AuditLog = django_apps.get_model("audit", "AuditLog")
        response = self._client(self.accountant).get(
            reverse("accounts:accounting_export"), {"ac_month": OCT, "ac_chair": str(self.c1.pk)}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Müəllimlər", "Fənn və qrup"])
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
        header = rows[1]
        body = {row[1]: row for row in rows[2:]}
        self.assertEqual(set(body), {"ac_t1", "ac_t2"})
        self.assertEqual(body["ac_t1"][header.index("Fakt (saat)")], 6)
        self.assertEqual(body["ac_t2"][header.index("Plan (saat)")], 8.2)
        log = AuditLog.objects.filter(resource_type="accounts.accounting").latest("created_at")
        self.assertEqual((log.new_values["rows"], log.user_id), (2, self.accountant.pk))
        detail = self._client(self.accountant).get(
            reverse("accounts:accounting_export"), {"ac_month": OCT, "ac_teacher": self.t1.pk}
        )
        self.assertEqual(load_workbook(io.BytesIO(detail.content)).sheetnames[-1], "Dərslər")

    # ── Sorğu büdcəsi ─────────────────────────────────────────────────────
    def test_query_count_does_not_grow_with_teachers(self):
        def measure(**params):
            request = self._request(self.accountant, ac_month=OCT, **params)
            with CaptureQueriesContext(connection) as ctx:
                build_panel(request)
            return len(ctx.captured_queries)

        measure()  # isinmə: proses səviyyəli keşlər ölçüyə düşməsin
        small = measure()
        small_detail = measure(ac_teacher=str(self.t1.pk))
        subject = Subject.objects.create(organization=self.org, code="AC-BULK", name="Kütləvi")
        task = TeachingTask.objects.get(chair=self.c1, academic_year=YEAR)
        row = TeachingTaskRow.objects.create(
            organization=self.org,
            task=task,
            season=Season.FALL,
            subject_text="Bulk",
            lecture_total=240,
            total_hours=240,
        )
        with bypass_rls():
            for index in range(12):
                teacher = self._teacher(f"ac_bulk{index}", "Ad", f"Soyad{index}", self.c2)
                self._assign(row, teacher, Activity.LECTURE, 20)
                offering = CourseOffering.objects.create(
                    organization=self.org,
                    subject=subject,
                    period=self.fall,
                    group=self._unit(f"G{index}", OrgUnitType.GROUP, self.f2),
                    instructor=teacher,
                )
                Enrollment.objects.create(organization=self.org, student=self.student, offering=offering)
                self._lesson(offering, date(2026, 10, 2), "lecture")
        large = measure()
        self.assertEqual(small, large)
        self.assertLessEqual(large, 16)
        self.assertEqual(small_detail, measure(ac_teacher=str(self.t1.pk)))
