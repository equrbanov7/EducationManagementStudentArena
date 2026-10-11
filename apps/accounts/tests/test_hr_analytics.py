"""«Kadr analitikası» kabinet bölməsi (sahib 2026-10-11) — RBAC, rəqəmlər, əhatə, sorğu büdcəsi, ixrac.

Kilidlənən qaydalar:

* menyu bəndi + panel HR (`people.hr_analytics`), rektor (`*`), RİM rəhbəri və superadmin üçündür;
  müəllim / tələbə / dekan görmür, fraqment 403, ixrac 403;
* heyət = akademik üzvlük VƏ YA seçilmiş ilin yük profili; ştat vahidi, elmi ad (boşdursa vəzifədən),
  dərəcə, status, yük doluluğu, tədris fəaliyyəti və məlumat boşluqları fiksturdakı kimi sayılır;
* açarı alan UNIT rolu yalnız öz alt-ağacını görür;
* sorğu sayı müəllim sayından asılı deyil;
* ixrac XLSX verir və audit jurnalına yazılır.
"""

from __future__ import annotations

import importlib
import io
from datetime import date, datetime, timezone

from django.apps import apps as django_apps
from django.contrib.auth import get_user_model
from django.db import connection
from django.test import Client, RequestFactory, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from apps.accounts.models import UserProfile
from apps.accounts.services.hr_analytics import aggregates as agg
from apps.accounts.services.hr_analytics.filters import apply_filters
from apps.accounts.services.hr_analytics.panel import build_panel, resolve
from apps.organizations.models import AcademicPeriod, Membership, Organization, OrgUnit, Role
from apps.registrar.models import CourseOffering, Lesson, Subject
from apps.workload.constants import Activity, TaskStatus
from apps.workload.models import TeacherAssignment, TeacherWorkloadProfile, TeachingTask, TeachingTaskRow
from core.constants import AcademicPeriodType, OrganizationType, OrgUnitType
from core.rls import bypass_rls

User = get_user_model()
SECTION = "hr-analytics"
YEAR = "2026/2027"
LOGGED = datetime(2026, 9, 20, tzinfo=timezone.utc)


@override_settings(UNIVERSITY_MODE=True)
class HrAnalyticsTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        owner = User.objects.create_user("hr_owner", "hr_owner@qku.edu.az", "pw")
        with bypass_rls():
            cls.org = Organization.objects.create(
                name="HR Univ", slug="hr-univ", org_type=OrganizationType.UNIVERSITY, owner=owner, status="active"
            )
            cls.f1 = cls._unit("Mühəndislik", OrgUnitType.FACULTY)
            cls.f2 = cls._unit("Humanitar", OrgUnitType.FACULTY)
            cls.c1 = cls._unit("İnformatika kafedrası", OrgUnitType.CHAIR, cls.f1)
            cls.c2 = cls._unit("Riyaziyyat kafedrası", OrgUnitType.CHAIR, cls.f1)
            cls.c3 = cls._unit("Tarix kafedrası", OrgUnitType.CHAIR, cls.f2)
            cls.period = AcademicPeriod.objects.create(
                organization=cls.org,
                name="Payız",
                period_type=AcademicPeriodType.SEMESTER,
                academic_year=YEAR,
                start_date="2026-09-15",
                end_date="2027-01-31",
                is_current=True,
            )
            cls.hr = cls._member("hr_hr", "hr")
            cls.rim = cls._member("hr_rim", "ikt_rehber")
            cls.rector = cls._member("hr_rector", "rector")
            cls.dean = cls._member("hr_dean", "dean", unit=cls.f1)
            cls.student = cls._member("hr_student", "student")
            cls.superadmin = User.objects.create_superuser("hr_su", "hr_su@qku.edu.az", "pw")

            cls.t1 = cls._teacher("t1", "Aysel", "Quliyeva", cls.c1, title="docent", degree="phd", gender="female")
            cls.t2 = cls._teacher("t2", "Rauf", "Məmmədov", cls.c1, logged=False, initial=True)
            cls.t3 = cls._teacher("t3", "Elçin", "Həsənov", cls.c2, title="professor", degree="dsc")
            User.objects.filter(pk=cls.t3.pk).update(is_active=False)
            cls.t4 = cls._teacher("t4", "Nigar", "Əliyeva", cls.c3, title="assistant", role="lab_assistant")
            cls.t5 = cls._teacher("t5", "Aysel", "Quliyeva", cls.c3)  # t1 ilə eyni ad-soyad
            cls.t6 = cls._user("t6", "Samir", "Babayev")  # yalnız yük profili, üzvlük yoxdur
            cls.h1 = cls._teacher("h1", "Kamran", "Rzayev", cls.c2, title="docent", role="chair_head")
            # t1 həm də dekandır (inzibati vəzifə) — ikinci, qeyri-əsas üzvlük.
            Membership.objects.create(
                user=cls.t1,
                organization=cls.org,
                role=cls.org.roles.get(name="dean"),
                scope_unit=cls.f2,
                is_primary=False,
                is_active=True,
            )
            UserProfile.objects.filter(user=cls.t5).update(staff_position="Metodist")

            cls._profile(cls.t1, "1.00", position="dosent")
            cls._profile(cls.t1, "0.50", position="muellim", year="2025/2026")
            cls._profile(cls.t2, "0.50", position="bas_muellim")
            cls._profile(cls.t3, "1.00", external=True)
            cls._profile(cls.t6, "1.00")

            task = TeachingTask.objects.create(
                organization=cls.org, chair=cls.c1, academic_year=YEAR, status=TaskStatus.DISTRIBUTED
            )
            row = TeachingTaskRow.objects.create(
                organization=cls.org, task=task, subject_text="Alqoritmlər", lecture_total=800, total_hours=800
            )
            TeacherAssignment.objects.create(
                organization=cls.org, row=row, teacher=cls.t1, activity=Activity.LECTURE, hours=600
            )
            TeacherAssignment.objects.create(
                organization=cls.org, row=row, teacher=cls.t2, activity=Activity.LECTURE, hours=200
            )

            subject = Subject.objects.create(organization=cls.org, code="HR-1", name="Alqoritmlər")
            offering = CourseOffering.objects.create(
                organization=cls.org, subject=subject, period=cls.period, instructor=cls.t1, lesson_hours=30
            )
            for day in (1, 2):
                Lesson.objects.create(organization=cls.org, offering=offering, date=date(2026, 9, 15 + day), hours=2)

    # ── fikstur köməkçiləri ────────────────────────────────────────────────
    @classmethod
    def _unit(cls, name, kind, parent=None):
        return OrgUnit.objects.create(
            organization=cls.org, name=name, slug=f"hr-{abs(hash(name)) % 10**8}", unit_type=kind, parent=parent
        )

    @classmethod
    def _user(cls, username, first="", last="", *, logged=True, initial=False):
        user = User.objects.create_user(username, f"{username}@qku.edu.az", "pw", first_name=first, last_name=last)
        if logged:
            User.objects.filter(pk=user.pk).update(last_login=LOGGED)
        UserProfile.objects.update_or_create(user=user, defaults={"password_change_required": initial})
        return user

    @classmethod
    def _member(cls, username, role, *, unit=None):
        user = cls._user(username)
        Membership.objects.create(
            user=user,
            organization=cls.org,
            role=cls.org.roles.get(name=role),
            scope_unit=unit,
            is_primary=True,
            is_active=True,
        )
        return user

    @classmethod
    def _teacher(cls, username, first, last, chair, *, title="", degree="", gender="", role="teacher", **kw):
        user = cls._user(username, first, last, **kw)
        UserProfile.objects.filter(user=user).update(
            academic_title=title, academic_degree=degree, gender=gender or "unspecified"
        )
        Membership.objects.create(
            user=user,
            organization=cls.org,
            role=cls.org.roles.get(name=role),
            scope_unit=chair,
            is_primary=True,
            is_active=True,
        )
        return user

    @classmethod
    def _profile(cls, user, fraction, *, position="muellim", external=False, year=YEAR):
        TeacherWorkloadProfile.objects.create(
            organization=cls.org,
            teacher=user,
            academic_year=year,
            position=position,
            staff_fraction=fraction,
            is_external=external,
        )

    def _client(self, user):
        client = Client()
        client.force_login(user)
        session = client.session
        session["active_organization"] = self.org.slug
        session.save()
        return client

    def _request(self, user, **params):
        request = RequestFactory().get("/accounts/profile/", {"section": SECTION, **params})
        request.user = user
        request.organization = self.org
        request.is_view_as = False
        return request

    def _rows(self, user=None, **params):
        actor, filters, _options, rows = resolve(self._request(user or self.hr, **params))
        return actor, filters, apply_filters(rows, filters, with_issue=filters.tab == "people")

    def _by_name(self, rows):
        return {row["username"]: row for row in rows}

    # ── RBAC ──────────────────────────────────────────────────────────────
    def test_visible_for_hr_rector_rim_head_and_superadmin(self):
        for user in (self.hr, self.rector, self.rim, self.superadmin):
            with self.subTest(user=user.username):
                client = self._client(user)
                page = client.get(reverse("accounts:profile"), {"section": SECTION})
                self.assertEqual(page.status_code, 200)
                self.assertContains(page, f'data-section="{SECTION}"')
                self.assertContains(page, "data-hr-root")
                self.assertContains(page, "hr_analytics.css")
                fragment = client.get(reverse("accounts:profile_section_fragment", args=[SECTION]))
                self.assertEqual(fragment.status_code, 200)
                self.assertIn("data-hr-root", fragment.json()["html"])

    def test_hidden_and_refused_for_teacher_student_and_dean(self):
        for user in (self.t1, self.student, self.dean):
            with self.subTest(user=user.username):
                client = self._client(user)
                page = client.get(reverse("accounts:profile"), {"section": SECTION})
                self.assertNotContains(page, f'data-section="{SECTION}"')
                self.assertNotContains(page, "data-hr-root")
                fragment = client.get(reverse("accounts:profile_section_fragment", args=[SECTION]))
                self.assertEqual(fragment.status_code, 403)
                self.assertEqual(client.get(reverse("accounts:hr_analytics_export")).status_code, 403)
                self.assertFalse(build_panel(self._request(user))["has_access"])

    def test_hr_role_template_carries_the_key(self):
        self.assertIn("people.hr_analytics", self.org.roles.get(name="hr").permissions)
        module = importlib.import_module("apps.organizations.migrations.0058_hr_analytics_permission")
        role = self.org.roles.get(name="hr")
        Role.objects.filter(pk=role.pk).update(permissions=[p for p in role.permissions if p != "people.hr_analytics"])
        module.forward(django_apps, None)
        module.forward(django_apps, None)  # idempotent
        self.assertEqual(self.org.roles.get(name="hr").permissions.count("people.hr_analytics"), 1)

    # ── Rəqəmlər ──────────────────────────────────────────────────────────
    def test_population_and_rows(self):
        _actor, filters, rows = self._rows()
        self.assertEqual(filters.year, YEAR)
        by = self._by_name(rows)
        self.assertEqual(set(by), {"t1", "t2", "t3", "t4", "t5", "t6", "h1"})
        self.assertEqual((by["t1"]["faculty"], by["t1"]["chair"]), ("Mühəndislik", "İnformatika kafedrası"))
        self.assertEqual(by["t2"]["title"], "senior_teacher")  # yük profilinin vəzifəsindən
        self.assertTrue(by["t2"]["title_from_position"])
        self.assertEqual(by["t3"]["staff_type"], "hourly")
        self.assertEqual(by["t3"]["status"], "blocked")
        self.assertEqual(by["t4"]["staff_type"], "none")
        self.assertEqual(by["h1"]["chair"], "Riyaziyyat kafedrası")
        self.assertEqual(
            (by["t1"]["planned_load"], by["t1"]["fill_percent"], by["t1"]["load_band"]), (600, 120, "loaded")
        )
        self.assertEqual((by["t2"]["fill_percent"], by["t2"]["load_band"]), (40, "free"))
        self.assertEqual((by["t1"]["journals"], by["t1"]["lessons"], by["t1"]["held_hours"]), (1, 2, 4))
        self.assertEqual(by["t1"]["planned_hours"], 30)

    def test_staff_units_composition_and_kpis(self):
        _actor, _filters, rows = self._rows()
        staff = agg.staff_metrics(rows)
        self.assertEqual(
            (staff["full"], staff["half"], staff["hourly"], staff["none"], staff["units"]), (2, 1, 1, 3, 2.5)
        )
        comp = agg.composition_metrics(rows)
        # assistant, teacher, senior_teacher, docent, professor, none
        self.assertEqual(comp["titles"], [1, 1, 1, 2, 1, 1])
        self.assertEqual((comp["phd"], comp["dsc"]), (1, 1))
        kpis = {tile["key"]: tile["value"] for tile in agg.kpis(rows)}
        self.assertEqual((kpis["total"], kpis["active"], kpis["blocked"]), (7, 6, 1))
        self.assertEqual((kpis["never"], kpis["initial"]), (1, 1))
        tree = agg.tree(rows, agg.staff_metrics)
        faculty_names = [group["name"] for group in tree["groups"]]
        self.assertEqual(faculty_names[:2], ["Humanitar", "Mühəndislik"])
        self.assertEqual(tree["groups"][-1]["name"], "Fakültəsiz")  # t6

    def test_quality_issues(self):
        _actor, _filters, rows = self._rows()
        counts = {item["key"]: item["count"] for item in agg.quality(rows)}
        self.assertEqual(counts["no_title"], 1)  # t5; t6-nın adı yük profilinin vəzifəsindən
        self.assertEqual(counts["no_staff_unit"], 3)
        self.assertEqual(counts["no_chair"], 1)
        self.assertEqual(counts["duplicate_name"], 2)
        self.assertEqual(counts["no_membership"], 1)
        self.assertEqual(counts["never_logged_in"], 1)
        _actor, _filters, people = self._rows(hr_tab="people", hr_issue="duplicate_name")
        self.assertEqual({row["username"] for row in people}, {"t1", "t5"})

    def test_filters(self):
        cases = {
            "hr_faculty": (str(self.f1.pk), {"t1", "t2", "t3", "h1"}),
            "hr_chair": (str(self.c2.pk), {"t3", "h1"}),
            "hr_staff": ("hourly", {"t3"}),
            "hr_title": ("docent", {"t1", "h1"}),
            "hr_q": ("məmmədov", {"t2"}),
        }
        for param, (value, expected) in cases.items():
            with self.subTest(param=param):
                _actor, _filters, rows = self._rows(**{param: value})
                self.assertEqual({row["username"] for row in rows}, expected)
        _actor, filters, rows = self._rows(hr_year="1999/2000", hr_faculty="bogus", hr_tab="x")
        self.assertEqual((filters.year, filters.faculty, filters.tab, len(rows)), (YEAR, "", "overview", 7))

    def test_unit_scoped_holder_sees_only_its_subtree(self):
        dean_role = self.org.roles.get(name="dean")
        Role.objects.filter(pk=dean_role.pk).update(permissions=[*dean_role.permissions, "people.hr_analytics"])
        actor, _filters, rows = self._rows(self.dean)
        self.assertTrue(actor.has_access)
        self.assertEqual({row["username"] for row in rows}, {"t1", "t2", "t3", "h1"})

    def test_every_tab_renders(self):
        client = self._client(self.hr)
        for tab in ("overview", "staff", "composition", "workload", "activity", "quality", "people"):
            with self.subTest(tab=tab):
                response = client.get(reverse("accounts:profile_section_fragment", args=[SECTION]), {"hr_tab": tab})
                self.assertEqual(response.status_code, 200)
                html = response.json()["html"]
                self.assertIn('aria-current="page"', html)
                self.assertNotIn("style=", html.split("data-hr-root", 1)[1])

    # ── Sorğu büdcəsi ─────────────────────────────────────────────────────
    def test_query_count_does_not_grow_with_teachers(self):
        def measure():
            request = self._request(self.hr, hr_tab="people")
            with CaptureQueriesContext(connection) as ctx:
                build_panel(request)
            return len(ctx.captured_queries)

        measure()  # isinmə: proses səviyyəli keşlər (ContentType, rol scope-u) ölçüyə düşməsin
        small = measure()
        with bypass_rls():
            for index in range(12):
                user = self._teacher(f"bulk{index}", "Ad", f"Soyad{index}", self.c3)
                self._profile(user, "0.50")
        large = measure()
        self.assertEqual(small, large)
        self.assertLessEqual(large, 20)

    # ── İxrac ─────────────────────────────────────────────────────────────
    def test_export_is_xlsx_and_audited(self):
        from openpyxl import load_workbook

        AuditLog = django_apps.get_model("audit", "AuditLog")
        response = self._client(self.hr).get(reverse("accounts:hr_analytics_export"), {"hr_chair": str(self.c1.pk)})
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(workbook.sheetnames), 7)
        people = list(workbook.worksheets[0].iter_rows(values_only=True))
        self.assertIn("İnzibati vəzifə", people[0])
        names = [row[0] for row in people[1:]]
        self.assertEqual(len(names), 2)
        admin_col = people[0].index("İnzibati vəzifə")
        self.assertIn("Humanitar", next(row[admin_col] for row in people[1:] if row[1] == "t1"))
        history = list(workbook.worksheets[1].iter_rows(values_only=True))
        self.assertEqual(history[0][3], "Tədris ili")
        t1_years = [row[3] for row in history[1:] if row[1] == "t1"]
        self.assertEqual(t1_years, [YEAR, "2025/2026"])
        log = AuditLog.objects.filter(resource_type="accounts.hr_analytics").latest("created_at")
        self.assertEqual(log.new_values["rows"], 2)
        self.assertEqual(log.user_id, self.hr.pk)

    # ── Sahib 2026-10-11: hesab / inzibati vəzifə / dərəcə / band filtrləri, sıralama, tarixçə ───
    def test_account_admin_degree_and_band_filters(self):
        cases = {
            "hr_status": ("active", {"t1", "t2", "t4", "t5", "t6", "h1"}),
            "hr_status=stopped": ("stopped", {"t3"}),
            "hr_status=blocked": ("blocked", {"t3"}),
            "hr_status=never": ("never", {"t2"}),
            "hr_status=initial": ("initial", {"t2"}),
            "hr_admin": ("any", {"t1", "t5", "h1"}),
            "hr_admin=none": ("none", {"t2", "t3", "t4", "t6"}),
            "hr_admin=dean": ("dean", {"t1"}),
            "hr_admin=chair_head": ("chair_head", {"h1"}),
            "hr_degree": ("phd", {"t1"}),
            "hr_degree=none": ("none", {"t2", "t4", "t5", "t6", "h1"}),
            "hr_band": ("none", {"t3", "t4", "t5", "t6", "h1"}),
            "hr_q=admin": ("dean humanitar", {"t1"}),  # inzibati vəzifə + vahid adı axtarışdadır
        }
        for label, (value, expected) in cases.items():
            param = label.split("=")[0]
            with self.subTest(filter=label):
                _actor, _filters, rows = self._rows(**{param: value})
                self.assertEqual({row["username"] for row in rows}, expected)
        # Tanınmayan dəyərlər səssizcə boşalır (500 yox); rol adı yalnız datada varsa qəbul olunur.
        _actor, filters, rows = self._rows(hr_status="x", hr_admin="rector", hr_degree="y", hr_band="z")
        self.assertEqual((filters.status, filters.admin, filters.degree, filters.band, len(rows)), ("", "", "", "", 7))
        _actor, filters, _rows = self._rows(hr_admin="../etc")
        self.assertEqual(filters.admin, "")

    def test_admin_roles_and_staff_position_on_rows(self):
        _actor, _filters, rows = self._rows()
        by = self._by_name(rows)
        self.assertEqual([item["role"] for item in by["t1"]["admin_roles"]], ["dean"])
        self.assertEqual(by["t1"]["admin_roles"][0]["unit"], "Humanitar")
        self.assertIn("Humanitar", by["t1"]["admin_label"])
        self.assertEqual(by["h1"]["admin_role_names"], ["chair_head"])
        self.assertEqual(by["t5"]["admin_label"], "Metodist")
        self.assertEqual(by["t2"]["admin_label"], "")

    def test_people_sorting(self):
        def usernames(**params):
            panel = build_panel(self._request(self.hr, hr_tab="people", **params))
            return [row["username"] for row in panel["people"]["results"]]

        self.assertEqual(usernames()[0], "t6")  # ad üzrə (Babayev) — defolt
        by_staff = usernames(hr_sort="-staff")
        self.assertEqual(by_staff[:2], ["t1", "t6"])  # 1,00 · 1,00 (azalan: ad tərsinə)
        self.assertEqual(by_staff[2:4], ["t2", "t3"])  # 0,50 · saathesabı
        self.assertEqual(set(by_staff[4:]), {"t4", "t5", "h1"})  # profilsizlər sonda
        self.assertEqual(usernames(hr_sort="staff")[0], "t5")
        self.assertEqual(usernames(hr_sort="-load")[0], "t1")
        self.assertEqual(usernames(hr_sort="-activity")[0], "t1")
        self.assertEqual(usernames(hr_sort="status")[-1], "t3")
        self.assertEqual(usernames(hr_sort="admin")[:3], ["t1", "h1", "t5"])  # vəzifəsi olanlar əvvəl
        self.assertEqual(usernames(hr_sort="bogus"), usernames())
        self.assertEqual(usernames(hr_sort="-"), usernames())
        panel = build_panel(self._request(self.hr, hr_tab="people", hr_sort="-staff"))
        columns = {col["key"]: col for col in panel["people"]["columns"]}
        self.assertEqual(columns["staff"]["dir"], "descending")
        self.assertIn("hr_sort=staff", columns["staff"]["url"])  # ikinci klik istiqaməti çevirir
        self.assertEqual(columns["name"]["dir"], "")
        self.assertNotIn("hr_sort", columns["name"]["url"])  # defolt açar URL-ə yazılmır
        self.assertIn("hr_sort=-staff", panel["tabs"][-1]["url"])  # sıralama tab linklərində qalır

    def test_people_history_and_kpi_filter_wiring(self):
        panel = build_panel(self._request(self.hr, hr_tab="people", hr_status="active"))
        rows = {row["username"]: row for row in panel["people"]["results"]}
        timeline = rows["t1"]["history"]
        self.assertEqual([item["year"] for item in timeline], [YEAR, "2025/2026"])
        self.assertEqual([item["staff_short"] for item in timeline], ["1", "0,5"])
        self.assertTrue(timeline[0]["changed"])
        self.assertTrue(timeline[0]["current"])
        self.assertFalse(timeline[1]["changed"])
        self.assertEqual(rows["t1"]["history_changes"], 1)
        self.assertEqual(rows["t4"]["staff_short"], "—")
        self.assertEqual(rows["t2"]["staff_short"], "0,5")
        self.assertNotIn("t3", rows)  # hr_status=active — dayandırılmış hesab siyahıda yoxdur
        everyone = build_panel(self._request(self.hr, hr_tab="people"))["people"]["results"]
        self.assertEqual(next(row["staff_short"] for row in everyone if row["username"] == "t3"), "s/h")
        tiles = {tile["key"]: tile for tile in panel["kpis"]}
        self.assertEqual(tiles["active"]["filter"], "active")
        self.assertTrue(tiles["active"]["pressed"])
        self.assertFalse(tiles["blocked"]["pressed"])
        self.assertNotIn("filter", tiles["fill"])
        self.assertIn('"blocked": "stopped"', panel["kpi_filter_map"])
        self.assertTrue(panel["is_filtered"])

        client = self._client(self.hr)
        response = client.get(reverse("accounts:profile_section_fragment", args=[SECTION]), {"hr_tab": "people"})
        html = response.json()["html"]
        self.assertIn("hr_analytics.js", html)
        self.assertIn('data-hr-history-toggle="hr-history-%s"' % self.t1.pk, html)
        self.assertIn('name="hr_status"', html)
        self.assertIn('name="hr_admin"', html)
        self.assertIn('data-ems-kpi-filter="active"', html)
        self.assertIn('aria-sort="ascending"', html)
        self.assertIn("2025/2026", html)
