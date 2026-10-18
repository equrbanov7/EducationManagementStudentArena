"""«Mühasibatlıq» — müəllim pəncərəsi (modal), tarix aralığı rejimi, aralıqlı ixrac (sahib 2026-10-19).

Kilidlənən qaydalar:

* modal endpoint-i bölmə ilə EYNİ qapıdadır: açarsız 403, əhatədən kənar / mövcud olmayan müəllim 404; cavabda tələbə
  adı / istifadəçi adı, inline ``style=`` yoxdur; «Başqasının adından» rejimində ixrac düyməsi yoxdur;
* sorğu sayı müəllim və dərs sayından asılı deyil; «Daha çox» növbəti səhifəni qaytarır;
* aralığın planı = aylıq payların aralığa düşən iş günlərinə görə kəsilməsi: tam ay = ay rejimi, bütün semestr =
  semestr planı; semestrdən kənar tarixlər sıxılır, tərs tarixlər yerini dəyişir, pozuq tarix defolta düşür;
* KPI, cədvəl, modal və ixrac eyni aralığa tabedir; ``ac_from`` / ``ac_to`` paylaşılan linkdə aralığı açır.
"""

from __future__ import annotations

import io
import json
from datetime import date
from unittest import mock

from django.db import connection
from django.test import SimpleTestCase, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from apps.accounts.services.accounting import dataset
from apps.accounts.services.accounting.modal import teacher_modal
from apps.accounts.services.accounting.panel import build_panel, resolve_teacher
from apps.accounts.services.accounting.periods import apportion, semesters, teaching_days
from apps.accounts.services.accounting.ranges import (
    NOTE_CLAMPED,
    NOTE_INVALID,
    NOTE_OUTSIDE,
    NOTE_SWAPPED,
    default_range,
    month_parts,
    month_slice,
    normalize_range,
    range_day_counts,
    range_months,
    range_plan_kinds,
)
from apps.organizations.models import Role
from apps.registrar.models import CourseOffering, Enrollment, Subject
from apps.workload.constants import Activity, Season
from apps.workload.models import TeachingTask, TeachingTaskRow
from core.constants import OrgUnitType

from .test_accounting_2026_10_18 import OCT, SECTION, SEM, AccountingFixture, _Period


def _semester():
    return semesters([_Period(1, date(2026, 9, 15), date(2027, 1, 31), exam=date(2027, 1, 5), current=True)])[0]


class RangeRuleTest(SimpleTestCase):
    """Aralığın planı — aylıq paylardan iş günlərinə görə kəsmə (``ranges.py`` sənədi)."""

    PLAN = {"lecture": 60, "seminar": 0, "lab": 10}

    def _plan(self, start, end):
        semester = _semester()
        return range_plan_kinds(self.PLAN, teaching_days(semester), range_day_counts(semester, start, end))

    def test_full_month_equals_month_mode_and_full_semester_equals_semester_plan(self):
        semester = _semester()
        shares = {kind: apportion(hours, teaching_days(semester)) for kind, hours in self.PLAN.items()}
        october = self._plan(date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(
            october, {"lecture": shares["lecture"]["2026-10"], "seminar": 0.0, "lab": shares["lab"]["2026-10"]}
        )
        self.assertEqual(self._plan(semester.start, semester.end), {"lecture": 60.0, "seminar": 0.0, "lab": 10.0})

    def test_partial_months_are_prorated_by_teaching_days(self):
        semester = _semester()
        start, end = date(2026, 10, 12), date(2026, 11, 5)
        counts = range_day_counts(semester, start, end)
        self.assertEqual(counts, {"2026-10": 15, "2026-11": 4})  # B.e.–Cümə, pəncərə daxilində
        days = teaching_days(semester)
        lecture = apportion(60, days)
        expected = round(lecture["2026-10"] * 15 / 22 + lecture["2026-11"] * 4 / 21, 1)
        self.assertEqual(self._plan(start, end)["lecture"], expected)
        self.assertEqual(self._plan(start, end), {"lecture": 14.2, "seminar": 0.0, "lab": 2.3})
        parts = month_parts(60, days, counts)
        self.assertEqual(set(parts), {"2026-10", "2026-11"})
        self.assertAlmostEqual(sum(parts.values()), 11.25 + 15.7 * 4 / 21, places=6)

    def test_weekend_and_exam_session_days_carry_no_plan(self):
        self.assertEqual(self._plan(date(2026, 10, 3), date(2026, 10, 4)), {"lecture": 0.0, "seminar": 0.0, "lab": 0.0})
        self.assertEqual(self._plan(date(2027, 1, 10), date(2027, 1, 20)), {"lecture": 0.0, "seminar": 0.0, "lab": 0.0})

    def test_normalize_clamps_swaps_and_defaults(self):
        semester = _semester()
        self.assertEqual(
            normalize_range("2026-08-01", "2027-03-01", semester),
            (date(2026, 9, 15), date(2027, 1, 31), (NOTE_CLAMPED,)),
        )
        self.assertEqual(
            normalize_range("2026-10-31", "2026-10-01", semester),
            (date(2026, 10, 1), date(2026, 10, 31), (NOTE_SWAPPED,)),
        )
        self.assertEqual(
            normalize_range("bogus", "", semester, today=date(2026, 10, 18)),
            (date(2026, 9, 15), date(2026, 10, 18), (NOTE_INVALID,)),
        )
        self.assertEqual(normalize_range("", "", semester, today=date(2027, 6, 1))[:2], (semester.start, semester.end))
        self.assertEqual(
            normalize_range("01.10.2026", "15.10.2026", semester)[:2], (date(2026, 10, 1), date(2026, 10, 15))
        )

    def test_non_overlapping_range_falls_back_to_default_not_a_one_day_clamp(self):
        """Review B1: semestr dəyişəndə köhnə tarixlər bir günlük aralığa sıxılmırdı — indi defolt aralıq + qeyd."""
        semester = _semester()
        today = date(2026, 10, 18)
        for raw in (("2027-05-01", "2027-05-10"), ("2026-02-01", "2026-03-01"), ("2025-01-01", "2025-02-01")):
            with self.subTest(raw=raw):
                self.assertEqual(
                    normalize_range(*raw, semester, today=today),
                    (date(2026, 9, 15), date(2026, 10, 18), (NOTE_OUTSIDE,)),
                )
        # Qismən kəsişən aralıq əvvəlki kimi sıxılır.
        self.assertEqual(normalize_range("2027-01-20", "2027-03-01", semester)[2], (NOTE_CLAMPED,))

    def test_default_range_for_past_current_and_future_semesters(self):
        semester = _semester()
        self.assertEqual(default_range(semester, today=date(2026, 10, 18)), (date(2026, 9, 15), date(2026, 10, 18)))
        self.assertEqual(default_range(semester, today=date(2027, 6, 1)), (semester.start, semester.end))
        self.assertEqual(default_range(semester, today=date(2026, 8, 1)), (semester.start, semester.end))

    def test_range_months_and_slices(self):
        start, end = date(2026, 10, 12), date(2026, 12, 1)
        self.assertEqual(range_months(start, end), ["2026-10", "2026-11", "2026-12"])
        self.assertEqual(month_slice("2026-10", start, end), (date(2026, 10, 12), date(2026, 10, 31), True))
        self.assertEqual(month_slice("2026-11", start, end), (date(2026, 11, 1), date(2026, 11, 30), False))


RANGE = {"ac_month": "range", "ac_from": "2026-10-12", "ac_to": "2026-11-05"}


@override_settings(UNIVERSITY_MODE=True)
class AccountingDetailModalTest(AccountingFixture, TestCase):
    def _modal(self, user, teacher, **params):
        params.setdefault("ac_semester", SEM)
        return self._client(user).get(reverse("accounts:accounting_teacher_detail", args=[teacher]), params)

    # ── Tarix aralığı: panel, KPI, URL ────────────────────────────────────
    def test_range_rows_follow_the_prorating_rule(self):
        resolved, by = self._rows(**RANGE)
        self.assertTrue(resolved.filters.is_range)
        self.assertEqual((resolved.filters.date_from, resolved.filters.date_to), ("2026-10-12", "2026-11-05"))
        t1, t2 = by["ac_t1"], by["ac_t2"]
        # t1: 12.10 və 19.10 mühazirələri + 03.11 lab; 05.10 mühazirəsi aralıqdan kənardır.
        self.assertEqual((t1["held"], t1["kinds"], t1["lessons"]), (6, {"lecture": 4, "seminar": 0, "lab": 2}, 3))
        self.assertEqual((t1["planned"], t1["planned_kinds"]), (16.5, {"lecture": 14.2, "seminar": 0.0, "lab": 2.3}))
        self.assertEqual((t1["planned_sem"], t1["held_sem"]), (70, 8))  # semestr sütunu aralıqdan asılı deyil
        # t2: 13, 14, 20.10 seminarları (6 və 7.10 aralıqdan əvvəldir).
        self.assertEqual((t2["held"], t2["planned"]), (6, 7.1))

    def test_full_month_range_matches_month_mode(self):
        _resolved, month_rows = self._rows(ac_month=OCT)
        _resolved, range_rows = self._rows(ac_month="range", ac_from="2026-10-01", ac_to="2026-10-31")
        for username in ("ac_t1", "ac_t2", "ac_t3"):
            with self.subTest(teacher=username):
                for key in ("planned", "held", "kinds", "status"):
                    self.assertEqual(month_rows[username][key], range_rows[username][key])

    def test_out_of_semester_and_reversed_ranges_are_corrected_with_a_message(self):
        resolved, by = self._rows(ac_month="range", ac_from="2026-08-01", ac_to="2027-03-01")
        self.assertEqual((resolved.filters.date_from, resolved.filters.date_to), ("2026-09-15", "2027-01-31"))
        self.assertEqual((by["ac_t1"]["planned"], by["ac_t1"]["held"]), (70.0, 8))
        panel = build_panel(self._request(self.accountant, ac_month="range", ac_from="2026-08-01", ac_to="2027-03-01"))
        self.assertTrue(any("hüdudlarına" in text for text in panel["range_messages"]))
        resolved, _by = self._rows(ac_month="range", ac_from="2026-10-31", ac_to="2026-10-01")
        self.assertEqual((resolved.filters.date_from, resolved.filters.date_to), ("2026-10-01", "2026-10-31"))
        self.assertIn(NOTE_SWAPPED, resolved.filters.range_notes)

    def test_dates_alone_open_range_mode_and_params_are_shareable(self):
        resolved, _by = self._rows(ac_month="", ac_from="2026-10-12", ac_to="2026-11-05")
        self.assertTrue(resolved.filters.is_range)
        params = resolved.filters.as_params()
        self.assertEqual(
            (params["ac_month"], params["ac_from"], params["ac_to"]), ("range", "2026-10-12", "2026-11-05")
        )
        resolved, _by = self._rows(ac_month=OCT, ac_from="2026-10-12", ac_to="2026-11-05")
        self.assertFalse(resolved.filters.is_range)  # ay seçilibsə tarixlər nəzərə alınmır və URL-ə yazılmır
        self.assertNotIn("ac_from", resolved.filters.as_params())

    def test_range_panel_renders_date_fields_kpis_and_hint(self):
        panel = build_panel(self._request(self.accountant, ac_semester=SEM, **RANGE))
        self.assertEqual(panel["period_label"], "12.10.2026 – 05.11.2026")
        self.assertEqual((panel["range_min"], panel["range_max"]), ("2026-09-15", "2027-01-31"))
        kpis = {tile["label"]: tile["value"] for tile in panel["kpis"]}
        self.assertEqual(kpis["Keçirilmiş saat"], "14")  # t1 6 + t2 6 + t3 (21.10) 2
        self.assertEqual(kpis["Plan"], "23,6")  # 16,5 + 7,1
        html = (
            self._client(self.accountant)
            .get(reverse("accounts:profile_section_fragment", args=[SECTION]), {"ac_semester": SEM, **RANGE})
            .json()["html"]
        )
        self.assertIn('type="date" id="f-ac_from" value="2026-10-12"', html)
        self.assertIn('min="2026-09-15" max="2027-01-31"', html)
        self.assertIn('name="ac_from" value="2026-10-12"', html)
        self.assertIn('value="range" selected', html)
        self.assertIn("data-ac-detail-url", html)
        self.assertIn("accounting.js", html)
        self.assertNotIn("style=", html.split("data-ac-root", 1)[1])
        month_html = (
            self._client(self.accountant)
            .get(reverse("accounts:profile_section_fragment", args=[SECTION]), {"ac_month": OCT})
            .json()["html"]
        )
        self.assertIn('name="ac_from" value=""', month_html)  # ay rejimində gizli tarixlər boşdur
        self.assertNotIn('id="f-ac_from"', month_html)

    # ── Modal endpoint-i ──────────────────────────────────────────────────
    def test_modal_returns_teacher_detail_without_student_pii(self):
        response = self._modal(self.accountant, self.t1.pk, ac_month=OCT)
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])
        payload = response.json()
        self.assertEqual(payload["title"], "Quliyeva Aysel")
        self.assertIn("İnformatika kafedrası", payload["subtitle"])
        self.assertIn("Oktyabr 2026", payload["subtitle"])
        html = payload["html"]
        for needle in ("05.10.2026", "Alqoritmlər", "101A", "Mühazirə", "Plandan az", "data-ac-modal"):
            self.assertIn(needle, html)
        blob = json.dumps(payload, ensure_ascii=False)
        for secret in ("Gizlinov", "Tələbəzadə", "ac_student", "style="):
            self.assertNotIn(secret, blob)
        self.assertIn(f"ac_teacher={self.t1.pk}", payload["foot"])
        self.assertIn(reverse("accounts:accounting_export"), payload["foot"])

    def test_modal_refuses_without_key_and_hides_out_of_scope_teachers(self):
        for user in (self.t1, self.dean, self.hr, self.student):
            with self.subTest(user=user.username):
                self.assertEqual(self._modal(user, self.t1.pk).status_code, 403)
        dean_role = self.org.roles.get(name="dean")
        Role.objects.filter(pk=dean_role.pk).update(permissions=[*dean_role.permissions, "accounting.view"])
        self.assertEqual(self._modal(self.dean, self.t3.pk).status_code, 404)  # kafedra başqa fakültədədir
        self.assertEqual(self._modal(self.dean, 999999).status_code, 404)
        self.assertEqual(self._modal(self.dean, self.student.pk).status_code, 404)
        self.assertEqual(self._modal(self.dean, self.t1.pk).status_code, 200)

    def test_modal_in_view_as_is_read_only(self):
        request = self._request(self.accountant, ac_month=OCT)
        request.is_view_as = True
        result = teacher_modal(request, self.t1.pk)
        self.assertEqual(result.status, 200)
        self.assertEqual((result.context["export_url"], result.context["view_as"]), ("", True))

    def test_modal_range_shows_only_range_months_with_partial_slices(self):
        payload = self._modal(self.accountant, self.t1.pk, **RANGE).json()
        self.assertIn("12.10.2026 – 05.11.2026", payload["subtitle"])
        result = teacher_modal(self._request(self.accountant, ac_semester=SEM, **RANGE), self.t1.pk)
        detail = result.context["detail"]
        self.assertEqual([month["key"] for month in detail["months"]], ["2026-10", "2026-11"])
        self.assertEqual([month["partial"] for month in detail["months"]], ["12.10–31.10", "01.11–05.11"])
        self.assertEqual([month["planned_text"] for month in detail["months"]], ["13,1", "3,5"])
        self.assertEqual(result.context["lessons_total"], 3)
        self.assertEqual([row["held_text"] for row in detail["kind_rows"]], ["4", "0", "2"])

    def test_load_more_pages_lessons(self):
        subject = Subject.objects.create(organization=self.org, code="AC-LM", name="Çoxlu dərs")
        offering = self._offering(subject, self.t3)
        for day in range(1, 31):
            self._lesson(offering, date(2026, 10, day), "lecture")
            self._lesson(offering, date(2026, 10, day), "seminar")
        first = teacher_modal(self._request(self.accountant, ac_month=OCT), self.t3.pk).context
        self.assertEqual((len(first["lessons"]), first["lessons_total"], first["remaining"]), (50, 61, 11))
        self.assertIn("ac_part=lessons", first["more_url"])
        more = self._client(self.accountant).get(first["more_url"])
        self.assertEqual(more.status_code, 200)
        payload = more.json()
        self.assertEqual((payload["more_url"], payload["remaining"]), ("", 0))
        self.assertEqual(payload["rows"].count("<tr>"), 11)
        self.assertNotIn("Gizlinov", payload["rows"])

    def test_modal_query_count_is_bounded(self):
        def measure():
            client = self._client(self.accountant)
            url = reverse("accounts:accounting_teacher_detail", args=[self.t1.pk])
            client.get(url, {"ac_month": OCT})  # isinmə (sessiya, keşlər)
            with CaptureQueriesContext(connection) as ctx:
                self.assertEqual(client.get(url, {"ac_month": OCT}).status_code, 200)
            with CaptureQueriesContext(connection) as range_ctx:
                self.assertEqual(client.get(url, RANGE).status_code, 200)
            return len(ctx.captured_queries), len(range_ctx.captured_queries)

        small, small_range = measure()
        subject = Subject.objects.create(organization=self.org, code="AC-QB", name="Büdcə")
        for index in range(10):
            teacher = self._teacher(f"ac_qb{index}", "Ad", f"Büdcə{index}", self.c1)
            group = self._unit(f"QB{index}", OrgUnitType.GROUP, self.f1)
            offering = CourseOffering.objects.create(
                organization=self.org, subject=subject, period=self.fall, group=group, instructor=teacher
            )
            Enrollment.objects.create(organization=self.org, student=self.student, offering=offering)
            self._lesson(offering, date(2026, 10, 2), "lecture")
        held = CourseOffering.objects.get(subject__code="AC-1")
        for day in (6, 7, 8, 9, 13, 14, 15, 16):
            self._lesson(held, date(2026, 10, day), "lab")
        large, large_range = measure()
        self.assertEqual((small, small_range), (large, large_range))
        self.assertLessEqual(large, 24)  # ölçülüb: 20 (sessiya / auth / tenant middleware daxil)
        self.assertLessEqual(large_range, large + 1)  # aralıq rejimi: +1 GROUP BY

    def test_modal_reads_only_the_requested_teacher(self):
        """Review B3: pəncərə bütün təşkilatın sətirlərini qurmur — fakt / plan / şəxs yalnız həmin müəllimindir."""
        request = self._request(self.accountant, ac_semester=SEM, ac_month=OCT)
        with mock.patch.object(dataset, "people", wraps=dataset.people) as people:
            resolved = resolve_teacher(request, self.t1.pk)
        self.assertEqual([row["username"] for row in resolved.rows], ["ac_t1"])
        self.assertEqual(set(people.call_args.args[1]), {self.t1.pk})
        self.assertEqual((resolved.rows[0]["held"], resolved.rows[0]["planned"]), (6, 19.2))  # panel ilə eyni rəqəm
        self.assertEqual(resolve_teacher(self._request(self.accountant), 999999).rows, [])

    def test_modal_query_count_ignores_other_teachers(self):
        client = self._client(self.accountant)
        url = reverse("accounts:accounting_teacher_detail", args=[self.t1.pk])

        def measure():
            client.get(url, {"ac_month": OCT})  # isinmə
            with CaptureQueriesContext(connection) as ctx:
                self.assertEqual(client.get(url, {"ac_month": OCT}).status_code, 200)
            return len(ctx.captured_queries)

        before = measure()
        task = TeachingTask.objects.get(chair=self.c1, academic_year="2026/2027")
        row = TeachingTaskRow.objects.create(
            organization=self.org, task=task, season=Season.FALL, subject_text="X", lecture_total=900, total_hours=900
        )
        subject = Subject.objects.create(organization=self.org, code="AC-OT", name="Başqaları")
        for index in range(15):
            teacher = self._teacher(f"ac_ot{index}", "Ad", f"Başqa{index}", self.c1)
            self._assign(row, teacher, Activity.LECTURE, 30)
            offering = CourseOffering.objects.create(
                organization=self.org,
                subject=subject,
                period=self.fall,
                group=self._unit(f"OT{index}", OrgUnitType.GROUP, self.f1),
                instructor=teacher,
            )
            Enrollment.objects.create(organization=self.org, student=self.student, offering=offering)
            self._lesson(offering, date(2026, 10, 5), "lecture")
        self.assertEqual(measure(), before)

    def test_switching_semester_drops_stale_range_with_a_warning(self):
        """Review B1: payız tarixləri ilə yaz semestri seçiləndə bir günlük aralıq YOX — defolt aralıq + qeyd."""
        params = {
            "ac_semester": "2026/2027|spring",
            "ac_month": "range",
            "ac_from": "2026-10-12",
            "ac_to": "2026-11-05",
        }
        resolved, _by = self._rows(**params)
        self.assertIn(NOTE_OUTSIDE, resolved.filters.range_notes)
        start, end = default_range(resolved.semester)
        self.assertEqual((resolved.filters.date_from, resolved.filters.date_to), (start.isoformat(), end.isoformat()))
        self.assertNotEqual(start, end)
        panel = build_panel(self._request(self.accountant, **params))
        self.assertTrue(any("bu semestrə düşmür" in text for text in panel["range_messages"]))

    # ── İxrac ─────────────────────────────────────────────────────────────
    def test_export_honours_the_range_and_single_teacher(self):
        from openpyxl import load_workbook

        url = reverse("accounts:accounting_export")
        client = self._client(self.accountant)
        response = client.get(url, {"ac_semester": SEM, **RANGE})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertIn("12.10.2026 – 05.11.2026", rows[0][0])
        header, body = rows[1], {row[1]: row for row in rows[2:]}
        self.assertEqual(body["ac_t1"][header.index("Plan (saat)")], 16.5)
        self.assertEqual(body["ac_t1"][header.index("Fakt (saat)")], 6)
        self.assertEqual(body["ac_t2"][header.index("Fakt (saat)")], 6)
        single = load_workbook(io.BytesIO(client.get(url, {"ac_teacher": self.t1.pk, **RANGE}).content))
        teachers = list(single.worksheets[0].iter_rows(values_only=True))[2:]
        self.assertEqual([row[1] for row in teachers], ["ac_t1"])
        lessons = list(single["Dərslər"].iter_rows(values_only=True))[2:]
        self.assertEqual([row[0] for row in lessons], ["12.10.2026", "19.10.2026", "03.11.2026"])
