"""«Dəyərləndirilmiş nəticələr» — Excel ixracı, yeni süzgəclər, «Yoxlayan» qaydası (sahib 2026-10-27).

* Excel: imtahanın öz nəticə səhifəsi ilə EYNİ sütunlar (``apps.exams.services.results_export.HEADERS``),
  zal + kompüter (cəhdin qeydli kompüteri və ya final bileti), ata adı, fənn, kateqoriya, status;
  ixrac ekrandakı süzgəc / əhatə ilə eynidir; icazə qapısı bölmə ilə eynidir.
* Süzgəclər: «Qrup» və «Fənn» menyuda axtarışlı, «Növ»-də Midterm / Final / Quiz, «Dövr»-də
  istənilən tarix-saat aralığı.
* «Yoxlayan»: avtomatik test → «Avtomatik (sistem)», insan bal qoyubsa onun adı.
"""

from __future__ import annotations

import datetime
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from apps.accounts.models import ProfileRole
from apps.accounts.tests.test_review_queue_query_budget import _assign, _build_teacher_world
from apps.accounts.tests.test_review_results_redesign import _check_written
from apps.exams.domain.final_center import ROOM_SESSION_STATE_ACTIVE, TICKET_STATUS_WAITING
from apps.exams.models import (
    ExamAttempt,
    ExamGradeEvent,
    ExamRoom,
    ExamRoomComputer,
    ExamRoomSession,
    FinalExamTicket,
)
from apps.exams.services.results_export import HEADERS
from apps.registrar.models import Subject

User = get_user_model()
FRAGMENT = "accounts:profile_section_fragment"
EXPORT = "accounts:review_results_export"
AUTO = "Avtomatik (sistem)"


def _fmt(moment):
    return timezone.localtime(moment).strftime("%d.%m.%Y %H:%M")


def _build_export_world(slug):
    """Müəllim dünyası + midterm test (fənli), yazılı (biri insan tərəfindən yoxlanılıb), zal / kompüter / bilet."""
    world = _build_teacher_world(slug, 3)
    org, teacher = world["org"], world["teacher"]
    _check_written(world)
    test_exam, written_exam = world["exams"]["test"], world["exams"]["written"]
    subject = Subject.objects.create(organization=org, code=f"{slug}-S1", name=f"{slug} Riyaziyyat")
    test_exam.exam_type_extended = "midterm"
    test_exam.subject = subject
    test_exam.save(update_fields=["exam_type_extended", "subject"])
    written_exam.exam_type_extended = "quiz"
    written_exam.save(update_fields=["exam_type_extended"])

    student = world["students"][0]
    student.profile.patronymic = "Rəşid"
    student.profile.save(update_fields=["patronymic"])

    hall = ExamRoom.objects.create(organization=org, name=f"Zal {slug}", code=f"{slug}-Z", capacity=20)
    computer = ExamRoomComputer.objects.create(
        organization=org, room=hall, label="PC-7", seat_number=7, mac_address="AA:BB:CC:DD:EE:07"
    )
    test_attempts = list(ExamAttempt.objects.filter(exam=test_exam).order_by("user_id"))
    ExamAttempt.objects.filter(pk=test_attempts[0].pk).update(room=hall, room_computer=computer)

    # Final bileti ilə oturan tələbə: cəhddə zal yoxdur — sessiyanın zalı + yer → həmin zalın kompüteri.
    ticket_hall = ExamRoom.objects.create(organization=org, name=f"Zal B {slug}", code=f"{slug}-B", capacity=20)
    ExamRoomComputer.objects.create(
        organization=org, room=ticket_hall, label="PC-3", seat_number=3, mac_address="AA:BB:CC:DD:EE:03"
    )
    now = timezone.now()
    session = ExamRoomSession.objects.create(
        organization=org,
        room=ticket_hall,
        scheduled_start=now - datetime.timedelta(hours=4),
        scheduled_end=now + datetime.timedelta(hours=1),
        state=ROOM_SESSION_STATE_ACTIVE,
        started_at=now - datetime.timedelta(hours=4),
        created_by=teacher,
    )
    FinalExamTicket.objects.create(
        organization=org,
        session=session,
        exam=test_exam,
        student=test_attempts[1].user,
        attempt=test_attempts[1],
        seat_number=3,
        status=TICKET_STATUS_WAITING,
    )

    reviewer = User.objects.create_user(
        f"{slug}_rev", f"{slug}_rev@example.com", "pw", first_name="Rəy", last_name="Çi"
    )
    written_attempt = ExamAttempt.objects.filter(exam=written_exam, user=student).get()
    ExamGradeEvent.objects.create(attempt=written_attempt, grader=reviewer, old_score=None, new_score=3, max_points=5)
    world.update(
        test_attempts=test_attempts,
        written_attempt=written_attempt,
        reviewer=reviewer,
        subject=subject,
        hall=hall,
        ticket_hall=ticket_hall,
    )
    return world


class _Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.world = _build_export_world("rrx27")
        cls.org = cls.world["org"]
        cls.teacher = cls.world["teacher"]
        test_exam, written_exam = cls.world["exams"]["test"], cls.world["exams"]["written"]
        cls.test_kind, cls.written_kind = test_exam.get_exam_type_display(), written_exam.get_exam_type_display()
        cls.midterm, cls.quiz = (
            test_exam.get_exam_type_extended_display(),
            written_exam.get_exam_type_extended_display(),
        )

    def _login(self, user):
        self.client.force_login(user)
        session = self.client.session
        session["active_organization"] = self.org.slug
        session.save()

    def _items(self, user=None, **params):
        self._login(user or self.teacher)
        url = reverse(FRAGMENT, kwargs={"section": "review-results"})
        response = self.client.get(url, params, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 200)
        return response, list(response.context["evaluated_review_page_obj"].paginator.object_list)

    def _export(self, user=None, **params):
        self._login(user or self.teacher)
        response = self.client.get(reverse(EXPORT), params)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        header, body = rows[0], rows[1:]
        return header, [dict(zip(header, row)) for row in body]


class ReviewResultsExportTest(_Base):
    def test_export_has_unified_columns_hall_computer_and_checker(self):
        header, rows = self._export()
        self.assertEqual(tuple(header), HEADERS)
        self.assertEqual(len(rows), 6)  # 3 test + 3 yoxlanmış yazılı
        by_key = {(row["İstifadəçi adı"], row["Növ"]): row for row in rows}
        test_attempts = self.world["test_attempts"]
        seated = by_key[(test_attempts[0].user.username, self.test_kind)]
        self.assertEqual((seated["Zal"], seated["Kompüter"]), (self.world["hall"].name, "PC-7"))
        ticketed = by_key[(test_attempts[1].user.username, self.test_kind)]
        self.assertEqual((ticketed["Zal"], ticketed["Kompüter"]), (self.world["ticket_hall"].name, "PC-3"))
        nowhere = by_key[(test_attempts[2].user.username, self.test_kind)]
        self.assertIn(nowhere["Zal"], ("", None))
        for row in rows:
            if row["Növ"] == self.test_kind:
                self.assertEqual(row["Yoxlayan"], AUTO)
                self.assertEqual(row["Fənn"], self.world["subject"].name)
                self.assertEqual(row["Kateqoriya"], self.midterm)
                self.assertIsNotNone(row["Faiz"])
                self.assertTrue(row["Başlama"] and row["Bitmə"] and row["Müddət (s:dq:sn)"])
        student = self.world["students"][0]
        written = by_key[(student.username, self.written_kind)]
        self.assertEqual(written["Yoxlayan"], "Rəy Çi")
        self.assertEqual(written["Ata adı"], "Rəşid")
        self.assertEqual((written["Bal"], written["Maks. bal"], written["Faiz"]), (3, 5, 60))
        self.assertEqual(written["Kateqoriya"], self.quiz)
        self.assertTrue(written["Status"])

    def test_export_follows_on_screen_filters(self):
        _, items = self._items(evaluated_type="exam_midterm")
        header, rows = self._export(evaluated_type="exam_midterm")
        self.assertEqual(len(rows), len(items))
        self.assertEqual({row["Kateqoriya"] for row in rows}, {self.midterm})

        students = self.world["students"]
        names = {students[0].pk: "RRX G1", students[1].pk: "RRX G2", students[2].pk: "RRX G1"}
        with mock.patch("apps.registrar.student_group_names.current_group_names", return_value=names):
            _, items = self._items(evaluated_group="RRX G1")
            _, rows = self._export(evaluated_group="RRX G1")
        self.assertEqual(len(rows), len(items))
        self.assertEqual(len(rows), 4)
        self.assertEqual({row["Qruplar"] for row in rows}, {"RRX G1"})

    def test_export_url_on_page_carries_applied_filters(self):
        response, _ = self._items(evaluated_type="exam_midterm", evaluated_group="nope")
        export_url = response.context["evaluated_review_ui"]["export_url"]
        self.assertTrue(export_url.startswith(reverse(EXPORT)))
        self.assertIn("evaluated_type=exam_midterm", export_url)
        self.assertNotIn("evaluated_group", export_url)  # tanınmayan qrup tətbiq olunmur
        self.assertIn(export_url.replace("&", "&amp;"), response.json()["html"])

    def test_permission_scope(self):
        self.client.logout()
        response = self.client.get(reverse(EXPORT))
        self.assertEqual(response.status_code, 302)

        student = self.world["students"][0]
        self._login(student)
        response = self.client.get(reverse(EXPORT))
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("spreadsheetml", response.get("Content-Type", ""))

        other = User.objects.create_user("rrx27_t2", "rrx27_t2@example.com", "pw")
        _assign(other, self.org, ProfileRole.TEACHER, "teacher")
        _, rows = self._export(other)
        self.assertEqual(rows, [])  # başqa müəllimin nəticələri ixraca düşmür


class ReviewResultsFiltersTest(_Base):
    def test_filter_fields_group_subject_searchable_category_and_range(self):
        response, _ = self._items()
        fields = {field["name"]: field for field in response.context["evaluated_review_ui"]["filter_fields"]}
        self.assertTrue(fields["evaluated_group"]["searchable"])
        self.assertTrue(fields["evaluated_subject"]["searchable"])
        self.assertIn(
            self.world["subject"].name, [option["value"] for option in fields["evaluated_subject"]["options"]]
        )
        type_values = [option["value"] for option in fields["evaluated_type"]["options"]]
        for value in ("exams", "exam_test", "exam_written", "exam_midterm", "exam_final", "exam_quiz", "labs"):
            self.assertIn(value, type_values)
        self.assertIn("custom", [option["value"] for option in fields["evaluated_period"]["options"]])
        self.assertEqual(fields["evaluated_from"]["kind"], "datetime")
        self.assertEqual(fields["evaluated_to"]["range_select"], "evaluated_period")
        html = response.json()["html"]
        self.assertIn('name="evaluated_from"', html)
        self.assertIn("data-ems-dt-input", html)
        self.assertIn('data-ems-range-select="evaluated_period"', html)

    def test_category_filter(self):
        _, items = self._items(evaluated_type="exam_midterm")
        self.assertEqual({item["title"] for item in items}, {self.world["exams"]["test"].title})
        _, items = self._items(evaluated_type="exam_quiz")
        self.assertEqual({item["title"] for item in items}, {self.world["exams"]["written"].title})
        _, items = self._items(evaluated_type="exam_final")
        self.assertEqual(items, [])

    def test_subject_filter(self):
        _, items = self._items(evaluated_subject=self.world["subject"].name)
        self.assertEqual({item["title"] for item in items}, {self.world["exams"]["test"].title})
        response, items = self._items(evaluated_subject="Yoxdur")
        self.assertEqual(len(items), 6)  # tanınmayan fənn süzgəc deyil
        fields = {field["name"]: field for field in response.context["evaluated_review_ui"]["filter_fields"]}
        self.assertEqual(fields["evaluated_subject"]["value"], "")

    def test_datetime_range_filter(self):
        now = timezone.now()
        # Test cəhdləri ~2 saat əvvəl bitib; yazılılar 2 gün əvvəl yoxlanılıb.
        response, items = self._items(
            evaluated_from=_fmt(now - datetime.timedelta(hours=12)),
            evaluated_to=_fmt(now + datetime.timedelta(hours=1)),
        )
        self.assertEqual({item["exam_kind"] for item in items}, {"test"})
        self.assertEqual(response.context["evaluated_review_ui"]["period"], "custom")

        _, items = self._items(
            evaluated_period="custom",
            evaluated_from=_fmt(now - datetime.timedelta(days=3)),
            evaluated_to=_fmt(now - datetime.timedelta(days=1)),
        )
        self.assertEqual({item["exam_kind"] for item in items}, {"written"})

        # Yalnız tarix yazılan «son» — günün sonuna qədər.
        today = timezone.localtime(now).strftime("%d.%m.%Y")
        _, items = self._items(evaluated_from=_fmt(now - datetime.timedelta(hours=12)), evaluated_to=today)
        self.assertEqual(len(items), 3)

        # Hazır dövr aralığı üstələyir; oxunmayan mətn süzgəc deyil.
        response, items = self._items(evaluated_period="7", evaluated_from="bu gün", evaluated_to=today)
        self.assertEqual(len(items), 6)
        self.assertEqual(response.context["evaluated_review_ui"]["period"], "7")

        header, rows = self._export(
            evaluated_from=_fmt(now - datetime.timedelta(hours=12)),
            evaluated_to=_fmt(now + datetime.timedelta(hours=1)),
        )
        self.assertEqual({row["Növ"] for row in rows}, {self.test_kind})


class ReviewResultsCheckerTest(_Base):
    def test_checker_column_auto_vs_human(self):
        response, items = self._items()
        by_attempt = {item["attempt_id"]: item for item in items}
        for attempt in self.world["test_attempts"]:
            self.assertEqual(by_attempt[attempt.pk]["evaluator_display"], AUTO)
        self.assertEqual(by_attempt[self.world["written_attempt"].pk]["evaluator_display"], "Rəy Çi")
        others = [
            item
            for item in items
            if item["exam_kind"] == "written" and item["attempt_id"] != self.world["written_attempt"].pk
        ]
        self.assertTrue(others)
        self.assertTrue(all(item["evaluator_display"] == "" for item in others))  # kim yoxlayıb — məlum deyil
        teacher_name = self.teacher.get_full_name() or self.teacher.username
        self.assertNotIn(teacher_name, {item["evaluator_display"] for item in items})
        self.assertIn(AUTO, response.json()["html"])

    def test_human_regrade_of_test_shows_person(self):
        attempt = self.world["test_attempts"][0]
        ExamAttempt.objects.filter(pk=attempt.pk).update(graded_by=self.world["reviewer"])
        _, items = self._items()
        row = next(item for item in items if item["attempt_id"] == attempt.pk)
        self.assertEqual(row["evaluator_display"], "Rəy Çi")
        _, rows = self._export(evaluated_type="exam_test")
        self.assertIn("Rəy Çi", {row["Yoxlayan"] for row in rows})
