"""«Mühasibatlıq» → ödəniş dərəcələri və ilkin məbləğ (sahib 2026-10-19, tövsiyə olunan variant).

Kilidlənən qaydalar:

* məbləğ = dövrün keçirilmiş saatı × dərəcə (il × vəzifə × növ; növə xas yoxdursa «bütün növlər»), ``Decimal``, 2 rəqəm,
  ``ROUND_HALF_UP``; ay, bütün semestr və tarix aralığı rejimində eyni qayda;
* vəzifə (yük profili) və ya saatı olan növün dərəcəsi yoxdursa məbləğ BOŞDUR («Dərəcə təyin edilməyib»), 0 deyil;
* KPI «Hesablanmış məbləğ», cədvəl sütunu, modal və Excel («Məbləğ (AZN)», «Dərəcələr» vərəqi, audit-də dərəcə id-ləri);
* dərəcəni yalnız ``accounting.manage_fees`` TƏŞKİLAT əhatəsində dəyişir; view-as yazmır; yad sətir 404;
* sətir silinmir: dəyişiklik / «Ləğv et» arxivləyir; DB trigger-i DELETE-i və dərəcənin dəyişməsini bloklayır;
* sorğu sayı dərəcələrin və müəllimlərin sayından asılı deyil.
"""

from __future__ import annotations

import io
from decimal import Decimal

from django.apps import apps as django_apps
from django.core.exceptions import ValidationError
from django.db import IntegrityError, connection, transaction
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from apps.accounts.models import TeachingPayRate
from apps.accounts.services.accounting import rates
from apps.accounts.services.accounting.modal import teacher_modal
from apps.accounts.services.accounting.panel import build_panel
from apps.accounts.views.accounting_rates import accounting_rate_save
from apps.organizations.models import Membership, Organization, Role
from apps.workload.models import TeacherAssignment
from core.constants import OrganizationType, RoleScopeType
from core.rls import bypass_rls

from .test_accounting_2026_10_18 import OCT, SEM, YEAR, AccountingFixture

SAVE = "accounts:accounting_rate_save"


@override_settings(UNIVERSITY_MODE=True)
class PayRateTest(AccountingFixture, TestCase):
    def _rate(self, rate, *, position="muellim", activity="", year=YEAR, org=None):
        with bypass_rls():
            return TeachingPayRate.objects.create(
                organization=org or self.org,
                academic_year=year,
                position=position,
                activity=activity,
                hourly_rate=Decimal(rate),
            )

    def _pay(self, **params):
        _resolved, by = self._rows(**params)
        return {username: row["pay"] for username, row in by.items()}

    # ── Hesab ─────────────────────────────────────────────────────────────
    def test_amount_math_uses_decimal_half_up(self):
        self.assertEqual(rates.clean_rate("12,345"), Decimal("12.35"))
        self.assertEqual(rates.clean_rate("12.344"), Decimal("12.34"))
        self.assertEqual(rates.amount(7, Decimal("13.33")), Decimal("93.31"))
        self.assertEqual(rates.amount(Decimal("1.5"), Decimal("0.03")), Decimal("0.05"))  # 0,045 → 0,05
        for bad in ("", "abc", "0", "-5", "0.004", "10000.01", "NaN"):
            with self.subTest(bad=bad), self.assertRaises(ValidationError):
                rates.clean_rate(bad)

    def test_kind_rate_wins_over_all_kinds_rate(self):
        self._rate("20.00", activity="lecture")
        self._rate("15.50")
        pay = self._pay()
        self.assertEqual((pay["ac_t1"]["status"], pay["ac_t1"]["amount"]), (rates.PAY_OK, Decimal("120.00")))
        self.assertEqual(pay["ac_t1"]["kinds"]["lecture"]["rate"], Decimal("20.00"))
        # t2: 5 seminar × 2 saat × «bütün növlər» 15,50
        self.assertEqual(pay["ac_t2"]["amount"], Decimal("155.00"))

    def test_missing_position_and_missing_rate_leave_amount_empty_not_zero(self):
        self._rate("20.00", activity="lecture")
        pay = self._pay()
        self.assertEqual((pay["ac_t3"]["status"], pay["ac_t3"]["amount"]), (rates.PAY_NO_POSITION, None))
        self.assertEqual((pay["ac_t2"]["status"], pay["ac_t2"]["amount"]), (rates.PAY_NO_RATE, None))
        self.assertIn("vəzifə yoxdur", rates.missing_text(pay["ac_t3"]))
        # bütün semestr: t1-in laboratoriyası da var, lab dərəcəsi yoxdur → qismən, cəm boş
        semester = self._pay(ac_month="all")["ac_t1"]
        self.assertEqual(
            (semester["status"], semester["amount"], semester["missing"]), (rates.PAY_PARTIAL, None, ["lab"])
        )
        self.assertEqual(semester["kinds"]["lecture"]["amount"], Decimal("120.00"))
        self.assertIn("laboratoriya", rates.missing_text(semester))

    def test_no_rates_at_all_and_other_year_rates_do_not_count(self):
        self._rate("99.00", year="2025/2026")
        pay = self._pay()
        self.assertTrue(all(item["amount"] is None for item in pay.values()))

    def test_range_mode_uses_range_hours(self):
        self._rate("20.00")
        pay = self._pay(ac_month="range", ac_from="2026-10-01", ac_to="2026-10-12")
        self.assertEqual(pay["ac_t1"]["amount"], Decimal("80.00"))  # 5 və 12 oktyabr mühazirələri: 4 saat

    def test_hourly_paid_plan_is_labelled(self):
        with bypass_rls():
            TeacherAssignment.objects.filter(teacher=self.t2).update(is_hourly_paid=True)
        _resolved, by = self._rows()
        self.assertIn("saathesabı 30 saat", by["ac_t2"]["pay"]["hourly_note"])
        self.assertIn("ştat yükündədir", by["ac_t1"]["pay"]["hourly_note"])

    # ── Panel, modal, ixrac ───────────────────────────────────────────────
    def test_panel_kpi_table_and_rate_card(self):
        self._rate("20.00")
        panel = build_panel(self._request(self.accountant, ac_semester=SEM, ac_month=OCT))
        kpi = panel["kpis"][-1]
        self.assertEqual((kpi["label"], kpi["value"], kpi["unit"]), ("Hesablanmış məbləğ", "320,00", "AZN"))
        self.assertIn("dərəcəsi olmayan: 1", kpi["note"])
        by = {row["username"]: row for row in panel["table"]["results"]}
        self.assertEqual(by["ac_t1"]["pay_text"], "120,00")
        self.assertEqual(
            (by["ac_t3"]["pay_text"], by["ac_t3"]["pay_missing"]), ("", "Dərəcə təyin edilməyib (vəzifə yoxdur)")
        )
        self.assertTrue(panel["rates"]["can_manage"])
        self.assertEqual(len(panel["rates"]["active"]), 1)
        page = self._client(self.trustee).get(reverse("accounts:profile"), {"section": "accounting", "ac_month": OCT})
        self.assertContains(page, "Dərs saatının ödəniş dərəcələri")
        self.assertContains(page, "İlkin hesablama — rəsmi əmək haqqı hesabı deyil.")
        self.assertContains(page, "accounting_rates.js")
        self.assertNotContains(page, 'name="hourly_rate"')  # qəyyum görür, dəyişmir

    def test_modal_shows_kind_amounts(self):
        self._rate("20.00", activity="lecture")
        result = teacher_modal(self._request(self.accountant, ac_semester=SEM, ac_month=OCT), self.t1.pk)
        detail = result.context["detail"]
        self.assertEqual(detail["pay"]["amount_text"], "120,00")
        lecture = detail["kind_rows"][0]
        self.assertEqual((lecture["rate_text"], lecture["amount_text"]), ("20,00", "120,00"))
        response = self._client(self.accountant).get(
            reverse("accounts:accounting_teacher_detail", args=[self.t2.pk]), {"ac_semester": SEM, "ac_month": OCT}
        )
        self.assertIn("Dərəcə təyin edilməyib", response.json()["html"])

    def test_export_columns_rates_sheet_and_audit(self):
        from openpyxl import load_workbook

        rate = self._rate("20.00")
        response = self._client(self.accountant).get(reverse("accounts:accounting_export"), {"ac_month": OCT})
        workbook = load_workbook(io.BytesIO(response.content))
        rows = list(workbook["Müəllimlər"].iter_rows(values_only=True))
        self.assertIn("İlkin hesablama", rows[0][0])
        header = rows[1]
        body = {row[1]: row for row in rows[2:]}
        for column in ("Vəzifə", "Məbləğ (AZN)", "Məbləğ qeydi", "Saathesabı plan (semestr, saat)"):
            self.assertIn(column, header)
        amount = header.index("Məbləğ (AZN)")
        self.assertEqual(Decimal(str(body["ac_t1"][amount])), Decimal("120"))
        self.assertIsNone(body["ac_t3"][amount])
        self.assertEqual(body["ac_t3"][header.index("Məbləğ qeydi")], "Dərəcə təyin edilməyib (vəzifə yoxdur)")
        sheet = list(workbook["Dərəcələr"].iter_rows(values_only=True))
        self.assertEqual(sheet[2][-1], str(rate.pk))
        log = (
            django_apps.get_model("audit", "AuditLog")
            .objects.filter(resource_type="accounts.accounting")
            .latest("created_at")
        )
        self.assertEqual(log.new_values["pay_rates"], [str(rate.pk)])

    # ── İdarəetmə ─────────────────────────────────────────────────────────
    def _post(self, user, **data):
        payload = {"academic_year": YEAR, "position": "muellim", "activity": "", "hourly_rate": "18,50", **data}
        return self._client(user).post(reverse(SAVE), payload)

    def test_permission_matrix(self):
        self.assertEqual(self._post(self.accountant).status_code, 200)
        self.assertEqual(self._post(self.rim, hourly_rate="19").status_code, 200)
        for user in (self.trustee, self.dean, self.hr, self.t1, self.student):
            with self.subTest(user=user.username):
                self.assertEqual(self._post(user, hourly_rate="25").status_code, 403)
        with bypass_rls():
            unit_role = Role.objects.create(
                organization=self.org,
                name="t_unit_fees",
                display_name="unit fees",
                level=60,
                scope_type=RoleScopeType.UNIT,
                permissions=["accounting.view", "accounting.manage_fees"],
            )
            unit_user = self._user("ac_unit_fees")
            Membership.objects.create(
                user=unit_user, organization=self.org, role=unit_role, scope_unit=self.f1, is_active=True
            )
        self.assertEqual(self._post(unit_user, hourly_rate="26").status_code, 403)
        active = TeachingPayRate.objects.filter(organization=self.org, archived_at__isnull=True)
        self.assertEqual([str(row.hourly_rate) for row in active], ["19.00"])

    def test_view_as_never_writes(self):
        request = self._request(self.accountant)
        request.method = "POST"
        request.POST = {"academic_year": YEAR, "position": "muellim", "hourly_rate": "10"}
        request.is_view_as = True
        response = accounting_rate_save(request)
        self.assertEqual(response.status_code, 403)
        self.assertFalse(TeachingPayRate.objects.filter(organization=self.org).exists())

    def test_validation_errors(self):
        for data in ({"position": "rector"}, {"hourly_rate": "0"}, {"academic_year": "2030/2031"}, {"activity": "x"}):
            with self.subTest(data=data):
                response = self._post(self.accountant, **data)
                self.assertEqual(response.status_code, 400)
                self.assertFalse(response.json()["ok"])

    def test_change_archives_old_row_and_archive_never_deletes(self):
        AuditLog = django_apps.get_model("audit", "AuditLog")
        self.assertEqual(self._post(self.accountant, hourly_rate="10").status_code, 200)
        self.assertEqual(self._post(self.accountant, hourly_rate="10").status_code, 400)  # dəyişiklik yoxdur
        self.assertEqual(self._post(self.accountant, hourly_rate="12").status_code, 200)
        rows = list(TeachingPayRate.objects.filter(organization=self.org).order_by("created_at"))
        self.assertEqual([(str(r.hourly_rate), r.archive_reason) for r in rows], [("10.00", "replaced"), ("12.00", "")])
        archive = reverse("accounts:accounting_rate_archive", args=[rows[1].pk])
        self.assertEqual(self._client(self.trustee).post(archive).status_code, 403)
        self.assertEqual(self._client(self.accountant).post(archive).status_code, 200)
        self.assertEqual(self._client(self.accountant).post(archive).status_code, 404)  # artıq arxivdədir
        self.assertEqual(TeachingPayRate.objects.filter(organization=self.org).count(), 2)
        self.assertFalse(TeachingPayRate.objects.filter(organization=self.org, archived_at__isnull=True).exists())
        reasons = list(
            AuditLog.objects.filter(resource_type=rates.AUDIT_RESOURCE)
            .order_by("created_at")
            .values_list("reason", flat=True)
        )
        self.assertEqual(reasons, ["pay_rate_created", "pay_rate_replaced", "pay_rate_archived"])
        panel = build_panel(self._request(self.accountant, ac_semester=SEM, ac_month=OCT))
        self.assertEqual(len(panel["rates"]["history"]), 2)

    def test_other_organization_rate_is_404(self):
        with bypass_rls():
            other = Organization.objects.create(
                name="Başqa", slug="ac-other", org_type=OrganizationType.UNIVERSITY, owner=self.t3, status="active"
            )
        foreign = self._rate("30.00", org=other)
        response = self._client(self.accountant).post(reverse("accounts:accounting_rate_archive", args=[foreign.pk]))
        self.assertEqual(response.status_code, 404)

    def test_database_blocks_delete_and_rate_change(self):
        if connection.vendor != "postgresql":
            self.skipTest("PostgreSQL trigger-i")
        rate = self._rate("20.00")
        for statement in (
            "DELETE FROM accounts_teachingpayrate WHERE id = %s",
            "UPDATE accounts_teachingpayrate SET hourly_rate = 1 WHERE id = %s",
            "UPDATE accounts_teachingpayrate SET id = gen_random_uuid() WHERE id = %s",
            "UPDATE accounts_teachingpayrate SET archived_at = now(), archive_reason = 'bogus' WHERE id = %s",
            "UPDATE accounts_teachingpayrate SET archived_at = now() WHERE id = %s",  # səbəbsiz arxiv
        ):
            with self.subTest(statement=statement), bypass_rls():
                with self.assertRaises(IntegrityError), transaction.atomic(), connection.cursor() as cursor:
                    cursor.execute(statement, [rate.pk])
        with bypass_rls():
            with self.assertRaises(IntegrityError), transaction.atomic():
                TeachingPayRate.objects.create(
                    organization=self.org, academic_year=YEAR, position="muellim", hourly_rate=Decimal("5")
                )  # aktiv dublikat

    # ── Sorğu büdcəsi ─────────────────────────────────────────────────────
    def test_query_count_does_not_grow_with_rates(self):
        def measure():
            request = self._request(self.accountant, ac_semester=SEM, ac_month=OCT)
            with CaptureQueriesContext(connection) as ctx:
                build_panel(request)
            return len(ctx.captured_queries)

        measure()
        before = measure()
        self._rate("20.00")
        for position in ("professor", "dosent", "bas_muellim", "assistent"):
            for activity in ("", "lecture", "seminar", "lab"):
                self._rate("30.00", position=position, activity=activity)
        self.assertEqual(measure(), before)
