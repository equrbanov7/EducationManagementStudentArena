"""«Mühasibatlıq» — Excel ixracı (panelin cari filtr dəsti üzrə).

Vərəqlər:

1. «Müəllimlər» — süzülmüş BÜTÜN müəllimlər (səhifələmə yoxdur): fakültə, kafedra, ştat, dövrün planı,
   fakt (mühazirə / seminar / lab / cəm), fərq, icra %, status, semestr planı və semestr faktı;
2. «Fənn və qrup» — həmin müəllimlərin fənn × qrup × növ üzrə keçirilmiş saatı (1 GROUP BY);
3. «Dərslər» — YALNIZ drill-down açıqdırsa (``ac_teacher``): həmin müəllimin dərs siyahısı.

Hücrələr ``core.export_safety.sheet_append`` ilə yazılır (formula injection neytrallaşdırılır). Pul məbləği
YOXDUR — sahib saat istəyib; dərəcə/tarif qərarı açıq sualdır.
"""

from __future__ import annotations

import io

from django.utils.translation import pgettext

from core.export_safety import sheet_append
from core.ui import status_catalog

from ..hr_analytics.constants import STAFF_TYPES
from . import constants as C
from . import dataset
from .detail import lesson_list, period_bounds
from .filters import MONTH_ALL
from .formatting import staff_short
from .periods import month_label

_HEADER_FILL = "1F3A5F"
_HEADER_FONT = "FFFFFF"


def _kind_headers() -> list[str]:
    return [str(label) for _key, label in C.KINDS]


def _status_label(key: str) -> str:
    return str(status_catalog.label(C.STATUS_FAMILY, key))


def _period(resolved) -> str:
    if resolved.filters.month == MONTH_ALL:
        return resolved.semester.label
    return month_label(resolved.filters.month)


def _teachers_table(resolved, rows) -> dict:
    staff_labels = {key: str(label) for key, label in STAFF_TYPES}
    header = [
        pgettext(C.CTX, "Müəllim"),
        pgettext(C.CTX, "İstifadəçi adı"),
        pgettext(C.CTX, "Fakültə"),
        pgettext(C.CTX, "Kafedra"),
        pgettext(C.CTX, "Ştat"),
        pgettext(C.CTX, "Ştat növü"),
        pgettext(C.CTX, "Plan (saat)"),
        *(f"{label} ({pgettext(C.CTX, 'fakt')})" for label in _kind_headers()),
        pgettext(C.CTX, "Fakt (saat)"),
        pgettext(C.CTX, "Dərs sayı"),
        pgettext(C.CTX, "Fərq (saat)"),
        pgettext(C.CTX, "İcra, %"),
        pgettext(C.CTX, "Status"),
        pgettext(C.CTX, "Semestr planı"),
        pgettext(C.CTX, "Semestr faktı"),
    ]
    body = [
        [
            row["name"],
            row["username"],
            row["faculty"],
            row["chair"],
            staff_short(row),
            staff_labels.get(row["staff_type"], ""),
            row["planned"],
            *(row["kinds"][kind] for kind in C.KIND_KEYS),
            row["held"],
            row["lessons"],
            row["diff"],
            "" if row["pct"] is None else row["pct"],
            _status_label(row["status"]),
            row["planned_sem"],
            row["held_sem"],
        ]
        for row in rows
    ]
    return {"title": pgettext(C.CTX, "Müəllimlər"), "header": header, "rows": body}


def _breakdown_table(resolved, rows) -> dict:
    names = {row["id"]: row["name"] for row in rows}
    start, end = period_bounds(resolved.filters)
    items = dataset.subject_breakdown(resolved.actor.organization, resolved.semester, list(names), start=start, end=end)
    kinds = {key: str(label) for key, label in C.KINDS}
    body = [
        [
            names.get(item["teacher"], ""),
            item["offering__subject__name"] or "",
            item["offering__group__name"] or "",
            kinds.get(item["kind"], item["kind"]),
            int(item["lessons"] or 0),
            int(item["hours"] or 0),
        ]
        for item in sorted(items, key=lambda item: (names.get(item["teacher"], "").casefold(), str(item)))
    ]
    header = [
        pgettext(C.CTX, "Müəllim"),
        pgettext(C.CTX, "Fənn"),
        pgettext(C.CTX, "Qrup"),
        pgettext(C.CTX, "Növ"),
        pgettext(C.CTX, "Dərs sayı"),
        pgettext(C.CTX, "Saat"),
    ]
    return {"title": pgettext(C.CTX, "Fənn və qrup"), "header": header, "rows": body}


def _lessons_table(resolved) -> dict:
    items = lesson_list(resolved, int(resolved.filters.teacher), limit=C.EXPORT_LESSON_LIMIT)
    header = [
        pgettext(C.CTX, "Tarix"),
        pgettext(C.CTX, "Vaxt"),
        pgettext(C.CTX, "Qrup"),
        pgettext(C.CTX, "Fənn"),
        pgettext(C.CTX, "Növ"),
        pgettext(C.CTX, "Akademik saat"),
    ]
    body = [
        [
            item["date"].strftime("%d.%m.%Y"),
            item["time"].strftime("%H:%M") if item["time"] else "",
            item["group"],
            item["subject"],
            item["kind_label"],
            item["hours"],
        ]
        for item in items
    ]
    return {"title": pgettext(C.CTX, "Dərslər"), "header": header, "rows": body}


def build_tables(resolved, rows) -> list[dict]:
    tables = [_teachers_table(resolved, rows), _breakdown_table(resolved, rows)]
    # Əhatə: ``persons`` yalnız aktorun görə bildiyi müəllimlərdir — başqa id-nin dərsləri heç vaxt çıxmır.
    if resolved.filters.teacher and int(resolved.filters.teacher) in resolved.persons:
        tables.append(_lessons_table(resolved))
    return tables


def xlsx_bytes(tables, *, caption: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    fill = PatternFill("solid", fgColor=_HEADER_FILL)
    font = Font(name="Calibri", color=_HEADER_FONT, bold=True)
    for index, table in enumerate(tables):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = str(table["title"])[:31]
        sheet_append(sheet, [caption])
        sheet_append(sheet, table["header"])
        for cell in sheet[2]:
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in table["rows"]:
            sheet_append(sheet, ["" if value is None else value for value in row])
        sheet.freeze_panes = "A3"
        for column, header in enumerate(table["header"], start=1):
            values = [header, *(row[column - 1] for row in table["rows"][:300] if column - 1 < len(row))]
            width = min(max(len(str(value or "")) for value in values) + 2, 48)
            sheet.column_dimensions[get_column_letter(column)].width = max(width, 6)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_caption(resolved) -> str:
    return pgettext(C.CTX, "Mühasibatlıq · %(period)s · keçirilmiş dərs saatı və plan") % {"period": _period(resolved)}


__all__ = ["build_tables", "export_caption", "xlsx_bytes"]
