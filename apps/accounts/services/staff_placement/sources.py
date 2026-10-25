"""Mənbə faylları → «sübut» sətirləri (``Evidence``) — DB-yə YAZMIR, bölmələri yalnız ad üzrə tutuşdurur.

* **Forma 7** (``teacher_roster.parsing.read_workbook``) — kafedra bloku + müəllim + ştat vahidi;
* **struktur siyahısı** (``staff_roster.parse_rows`` — ``seed_staff_roster``-in oxuyucusu) — bölmə başlığı + şəxs +
  vəzifə; vəzifə → rol xəritəsi ``staff_roster.role_for``;
* **kafedra dərs yükü faylı** (``workload.public.parse_load_workbook`` — «Dərs yükü faylı» oxuyucusu) — kafedra
  adı faylda deyil, əmrdə verilir.

Paket (``--bundle``): zip + ``manifest.json``::

    {"forma7": "forma7.xlsx", "structure": "struktur.xlsx",
     "loads": [{"file": "loads/1.xlsx", "chair": "Tarix kafedrası"}],
     "chair_map": {"Psixologiya kafedrası": "Sosial iş kafedrası"}}

Zip diskə AÇILMIR — fayllar yaddaşda oxunur (ölçü həddi var).
"""

from __future__ import annotations

import io
import json
import zipfile
from dataclasses import dataclass, field

from .. import staff_roster
from ..teacher_roster.chair_teachers import person_name, split_name
from ..teacher_roster.matching import match_chairs
from ..teacher_roster.parsing import RATE_HOURLY, RATE_UNKNOWN, PersonName, parse_name, read_workbook

SOURCE_FORMA7 = "forma7"
SOURCE_STRUCTURE = "structure"
SOURCE_LOAD = "load"

KIND_STAFF = "staff"  # ştat / yarım ştat / əvəzçi
KIND_HOURLY = "hourly"  # saathesabı
KIND_UNKNOWN = "unknown_rate"
KIND_LOAD = "load"

MAX_BUNDLE_MEMBER = 20 * 1024 * 1024
MAX_BUNDLE_MEMBERS = 60


class SourceError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(code, message)
        self.code, self.message = code, message


@dataclass(frozen=True)
class Evidence:
    source: str
    name: PersonName
    raw_name: str
    section: str  # fayldakı kafedra/bölmə başlığı
    unit_id: str | None  # tutuşdurulmuş OrgUnit id (str)
    unit_name: str
    unit_type: str
    kind: str  # KIND_* (Forma 7 / yük) və ya struktur rolunun adı
    detail: str = ""  # vəzifə / ştat mətni / fayl adı

    @property
    def chair(self) -> str:
        return self.section


@dataclass
class Sources:
    forma7: list = field(default_factory=list)
    structure: list = field(default_factory=list)
    loads: list = field(default_factory=list)
    chair_map: dict = field(default_factory=dict)  # fayldakı kafedra adı → {"unit", "score", "source"}
    structure_chairs: set = field(default_factory=set)  # struktur siyahısında KAFEDRA kimi tanınan bölmə id-ləri
    skipped: list = field(default_factory=list)  # [(mənbə, mətn, səbəb)]

    @property
    def all(self) -> list:
        return [*self.forma7, *self.structure, *self.loads]


def _open(path_or_bytes):
    return io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else path_or_bytes


def _rate_kind(rate) -> str:
    if rate.kind == RATE_HOURLY:
        return KIND_HOURLY
    if rate.kind == RATE_UNKNOWN:
        return KIND_UNKNOWN
    return KIND_STAFF


def _type_of(chair_units, unit) -> str:
    if not unit:
        return ""
    return next((item.unit_type for item in chair_units if str(item.pk) == str(unit[0])), "")


def read_forma7(source, chair_units, overrides, result: Sources) -> None:
    """``chair_units`` — kafedra tipli aktiv ``OrgUnit``-lər; ``overrides`` — ``{fayl adı: bölmə adı}``."""
    roster = read_workbook(_open(source))
    mapping = match_chairs(roster.chairs, [(str(unit.pk), unit.name) for unit in chair_units], overrides)
    result.chair_map.update(mapping)
    for row in roster.rows:
        unit = (mapping.get(row.chair) or {}).get("unit")
        result.forma7.append(
            Evidence(
                source=SOURCE_FORMA7,
                name=row.name,
                raw_name=row.raw_name,
                section=row.chair,
                unit_id=unit[0] if unit else None,
                unit_name=unit[1] if unit else "",
                unit_type=_type_of(chair_units, unit),
                kind=_rate_kind(row.rate),
                detail=row.raw_rate,
            )
        )
    result.skipped.extend((SOURCE_FORMA7, text, reason) for _row, text, reason in roster.skipped)


def _structure_rows(source):
    import openpyxl

    book = openpyxl.load_workbook(_open(source), data_only=True, read_only=True)
    try:
        sheet = book.worksheets[0]
        return [(row[0], row[1] if len(row) > 1 else "") for row in sheet.iter_rows(values_only=True) if row]
    finally:
        book.close()


def read_structure(source, all_units, kafedra_types, result: Sources) -> None:
    """Struktur siyahısı: hər şəxs → (bölmə, gözlənilən rol). Bölmə ``staff_roster.match_unit`` ilə tapılır."""
    for person in staff_roster.parse_rows(_structure_rows(source)):
        name = parse_name(person["name"])
        if name is None:
            result.skipped.append((SOURCE_STRUCTURE, person["name"], "name_unparsed"))
            continue
        unit = staff_roster.match_unit(person["section"], all_units)
        role, _mapped = staff_roster.role_for(person["section"], person["position"])
        unit_type = unit.unit_type if unit is not None else ""
        if unit is not None and unit_type in kafedra_types:
            result.structure_chairs.add(str(unit.pk))
        result.structure.append(
            Evidence(
                source=SOURCE_STRUCTURE,
                name=name,
                raw_name=person["name"],
                section=person["section"],
                unit_id=str(unit.pk) if unit is not None else None,
                unit_name=unit.name if unit is not None else "",
                unit_type=unit_type,
                kind=role,
                detail=person["position"],
            )
        )


def read_load(source, chair_name: str, chair_units, overrides, result: Sources, *, label: str = "") -> None:
    """Kafedranın dərs yükü faylı — müəllim adları (təkrarsız); kafedra ``chair_name`` ilə verilir."""
    from apps.workload.public import parse_load_workbook

    mapping = match_chairs([chair_name], [(str(unit.pk), unit.name) for unit in chair_units], overrides)
    result.chair_map.update(mapping)
    unit = (mapping.get(chair_name) or {}).get("unit")
    try:
        entries, _skipped = parse_load_workbook(_open(source))
    except Exception as exc:  # noqa: BLE001 — oxunmayan fayl hesabatda qeyd olunur, əmr dayanmır
        result.skipped.append((SOURCE_LOAD, label or chair_name, f"file_unreadable: {exc}"[:200]))
        return
    seen = set()
    for entry in entries:
        text = " ".join(str(entry.teacher_text or "").split())
        if not text or text in seen:
            continue
        seen.add(text)
        words, _initials = split_name(text)
        name = person_name(words)
        if name is None:
            result.skipped.append((SOURCE_LOAD, text, "name_unparsed"))
            continue
        result.loads.append(
            Evidence(
                source=SOURCE_LOAD,
                name=name,
                raw_name=text,
                section=chair_name,
                unit_id=unit[0] if unit else None,
                unit_name=unit[1] if unit else "",
                unit_type=_type_of(chair_units, unit),
                kind=KIND_LOAD,
                detail=label,
            )
        )


def read_bundle(path):
    """``(forma7 baytları, struktur baytları, [(yük baytları, kafedra, ad)], chair_map)`` — zip + manifest."""
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise SourceError("bundle_unreadable", f"Paket zip kimi oxunmadı: {exc}") from None
    with archive:
        infos = {info.filename: info for info in archive.infolist() if not info.is_dir()}
        if len(infos) > MAX_BUNDLE_MEMBERS:
            raise SourceError("bundle_too_many_files", "Paketdə həddən çox fayl var.")

        def blob(name):
            info = infos.get(str(name or ""))
            if info is None:
                raise SourceError("bundle_member_missing", f"Paketdə fayl yoxdur: {name!r}")
            if info.file_size > MAX_BUNDLE_MEMBER:
                raise SourceError("bundle_member_too_big", f"Fayl çox böyükdür: {name!r}")
            return archive.read(info)

        try:
            manifest = json.loads(blob("manifest.json").decode("utf-8"))
        except ValueError as exc:
            raise SourceError("bundle_manifest_invalid", f"manifest.json oxunmadı: {exc}") from None
        forma7 = blob(manifest["forma7"]) if manifest.get("forma7") else None
        structure = blob(manifest["structure"]) if manifest.get("structure") else None
        loads = [(blob(item["file"]), str(item["chair"]), str(item["file"])) for item in manifest.get("loads") or []]
        chair_map = {str(key): str(value) for key, value in (manifest.get("chair_map") or {}).items()}
    return forma7, structure, loads, chair_map


__all__ = [
    "KIND_HOURLY",
    "KIND_LOAD",
    "KIND_STAFF",
    "KIND_UNKNOWN",
    "SOURCE_FORMA7",
    "SOURCE_LOAD",
    "SOURCE_STRUCTURE",
    "Evidence",
    "SourceError",
    "Sources",
    "read_bundle",
    "read_forma7",
    "read_load",
    "read_structure",
]
