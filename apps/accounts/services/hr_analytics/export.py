"""«Kadr analitikası» Excel ixracı — cari filtr dəstinin cədvəlləri (sahib 2026-10-11).

Vərəqlər: müəllimlər (sətir-sətir, inzibati vəzifə ilə), ştat tarixçəsi (müəllim × tədris ili — bütün
illərin yük profilləri), ştat vahidləri, tərkib, dərs yükü, tədris fəaliyyəti, məlumat keyfiyyəti. Mətn xanaları ``core.export_safety.sheet_append`` ilə formula-neytrallaşdırılır; əlaqə
məlumatı (telefon, e-poçt, FİN) YOXDUR; cins/yaş yalnız demoqrafiya icazəsi ilə. Audit ixrac view-undadır.
"""

from __future__ import annotations

import io

from django.utils.translation import pgettext

from core.export_safety import sheet_append

from . import aggregates as agg
from . import constants as C

_HEADER_FILL = "2563EB"  # --ems-primary-600
_HEADER_FONT = "FFFFFF"


def _yes(value) -> str:
    return pgettext(C.CTX, "Bəli") if value else pgettext(C.CTX, "Xeyr")


def _people_table(rows, *, demographics: bool) -> dict:
    labels = agg._labels()
    issue_labels = {key: str(label) for key, label in C.ISSUES}
    header = [
        pgettext(C.CTX, "Ad Soyad Ata adı"),
        pgettext(C.CTX, "İstifadəçi adı"),
        pgettext(C.CTX, "Fakültə"),
        pgettext(C.CTX, "Kafedra"),
        pgettext(C.CTX, "Elmi ad / vəzifə"),
        pgettext(C.CTX, "Elmi dərəcə"),
        pgettext(C.CTX, "Ştat növü"),
        pgettext(C.CTX, "Ştat payı"),
        pgettext(C.CTX, "İnzibati vəzifə"),
        pgettext(C.CTX, "Norma (saat)"),
        pgettext(C.CTX, "Bölgü üzrə yük (saat)"),
        pgettext(C.CTX, "Doluluq %"),
        pgettext(C.CTX, "Jurnal"),
        pgettext(C.CTX, "Keçirilmiş dərs"),
        pgettext(C.CTX, "Keçirilmiş saat"),
        pgettext(C.CTX, "Plan saat (jurnal)"),
        pgettext(C.CTX, "Hesab statusu"),
        pgettext(C.CTX, "Heç daxil olmayıb"),
        pgettext(C.CTX, "Aktivləşdirilməyib"),
    ]
    if demographics:
        header += [pgettext(C.CTX, "Cins"), pgettext(C.CTX, "Yaş")]
    header.append(pgettext(C.CTX, "Problemlər"))
    gender_labels = {"male": pgettext(C.CTX, "Kişi"), "female": pgettext(C.CTX, "Qadın")}
    body = []
    for row in sorted(rows, key=lambda item: item["name"].casefold())[: C.EXPORT_LIMIT]:
        line = [
            row["name"],
            row["username"],
            row["faculty"],
            row["chair"],
            labels["title"].get(row["title"], ""),
            "" if row["degree"] == C.DEGREE_NONE else labels["degree"].get(row["degree"], ""),
            labels["staff"].get(row["staff_type"], ""),
            row["staff_fraction"],
            row["admin_label"],
            row["norm_hours"] if row["load_band"] != "none" else None,
            row["planned_load"],
            row["fill_percent"] if row["load_band"] != "none" else None,
            row["journals"],
            row["lessons"],
            row["held_hours"],
            row["planned_hours"],
            str(C.STATUS_LABELS.get(row["status"], "")),
            _yes(row["never_logged_in"]),
            _yes(row["not_activated"]),
        ]
        if demographics:
            line += [gender_labels.get(row["gender"], ""), row["age"]]
        line.append("; ".join(issue_labels[key] for key in row["issues"]))
        body.append(line)
    return {"title": pgettext(C.CTX, "Müəllimlər"), "header": header, "rows": body}


def _history_table(rows, organization) -> dict:
    """Ştat tarixçəsi — hər müəllim × tədris ili bir sətir (1 sorğu, ``dataset.profile_history``)."""
    from .dataset import profile_history

    labels = agg._labels()
    yes = pgettext(C.CTX, "Bəli")
    header = [
        pgettext(C.CTX, "Ad Soyad Ata adı"),
        pgettext(C.CTX, "İstifadəçi adı"),
        pgettext(C.CTX, "Kafedra"),
        pgettext(C.CTX, "Tədris ili"),
        pgettext(C.CTX, "Vəzifə (yük profili)"),
        pgettext(C.CTX, "Ştat növü"),
        pgettext(C.CTX, "Ştat payı"),
        pgettext(C.CTX, "Norma (saat)"),
        pgettext(C.CTX, "Əvvəlki ilə nisbətən dəyişib"),
    ]
    ordered = sorted(rows, key=lambda item: item["name"].casefold())[: C.EXPORT_LIMIT]
    history = profile_history(organization, [row["id"] for row in ordered]) if organization is not None else {}
    body = []
    for row in ordered:
        for item in history.get(row["id"], []):
            body.append(
                [
                    row["name"],
                    row["username"],
                    row["chair"],
                    item["year"],
                    labels["position"].get(item["position"], item["position"]),
                    labels["staff"].get(item["staff_type"], ""),
                    item["staff_fraction"] if item["staff_type"] not in (C.STAFF_HOURLY, C.STAFF_NONE) else None,
                    item["norm"] or None,
                    yes if item["changed"] else "",
                ]
            )
    return {"title": pgettext(C.CTX, "Ştat tarixçəsi"), "header": header, "rows": body}


def _tree_table(title, rows, metrics, columns) -> dict:
    """``columns``: [(başlıq, metrics-dən dəyər çıxaran funksiya)] — kafedra + fakültə cəmi + ümumi cəm."""
    data = agg.tree(rows, metrics)
    total_label = pgettext(C.CTX, "Cəmi")
    header = [pgettext(C.CTX, "Fakültə"), pgettext(C.CTX, "Kafedra")] + [label for label, _fn in columns]
    body = []
    for group in data["groups"]:
        for chair in group["chairs"]:
            body.append([group["name"], chair["name"]] + [fn(chair) for _label, fn in columns])
        body.append([group["name"], total_label] + [fn(group["total"]) for _label, fn in columns])
    body.append([pgettext(C.CTX, "Universitet üzrə"), total_label] + [fn(data["total"]) for _label, fn in columns])
    return {"title": title, "header": header, "rows": body}


def _pick(key, index=None):
    """Sütun dəyəri çıxaran funksiya (dövr dəyişənini erkən bağlayır)."""
    if index is None:
        return lambda metrics: metrics[key]
    return lambda metrics: metrics[key][index]


def build_tables(rows, *, demographics: bool, organization=None) -> list[dict]:
    labels = agg._labels()
    teachers = (pgettext(C.CTX, "Müəllim"), lambda m: m["teachers"])
    staff_columns = [teachers] + [(labels["staff"][key], _pick(key)) for key in C.STAFF_KEYS]
    staff_columns.append((pgettext(C.CTX, "Ştat vahidi"), lambda m: m["units"]))
    title_keys = (*C.TITLE_ORDER, C.TITLE_NONE)
    composition_columns = [teachers] + [
        (labels["title"][key], _pick("titles", index)) for index, key in enumerate(title_keys)
    ]
    composition_columns += [
        (labels["degree"]["phd"], lambda m: m["phd"]),
        (labels["degree"]["dsc"], lambda m: m["dsc"]),
        (pgettext(C.CTX, "Elmi dərəcəli %"), lambda m: m["degree_percent"]),
    ]
    workload_columns = [
        teachers,
        (pgettext(C.CTX, "Yükü olan"), lambda m: m["with_load"]),
        (pgettext(C.CTX, "Norma (saat)"), lambda m: m["norm"]),
        (pgettext(C.CTX, "Bölgü üzrə yük (saat)"), lambda m: m["planned"]),
        (pgettext(C.CTX, "Saathesabı saat"), lambda m: m["hourly_paid"]),
        (pgettext(C.CTX, "Doluluq %"), lambda m: m["fill"]),
    ] + [(str(label), _pick("bands", index)) for index, (_key, label) in enumerate(C.LOAD_BANDS)]
    activity_columns = [
        teachers,
        (pgettext(C.CTX, "Dərs deyən"), lambda m: m["teaching"]),
        (pgettext(C.CTX, "Jurnal"), lambda m: m["journals"]),
        (pgettext(C.CTX, "Keçirilmiş dərs"), lambda m: m["lessons"]),
        (pgettext(C.CTX, "Keçirilmiş saat"), lambda m: m["held"]),
        (pgettext(C.CTX, "Plan saat (jurnal)"), lambda m: m["planned"]),
        (pgettext(C.CTX, "İcra %"), lambda m: m["done"]),
    ]
    quality = {
        "title": pgettext(C.CTX, "Məlumat keyfiyyəti"),
        "header": [pgettext(C.CTX, "Problem"), pgettext(C.CTX, "Say")],
        "rows": [[item["label"], item["count"]] for item in agg.quality(rows)],
    }
    return [
        _people_table(rows, demographics=demographics),
        _history_table(rows, organization),
        _tree_table(pgettext(C.CTX, "Ştat vahidləri"), rows, agg.staff_metrics, staff_columns),
        _tree_table(pgettext(C.CTX, "Tərkib"), rows, agg.composition_metrics, composition_columns),
        _tree_table(pgettext(C.CTX, "Dərs yükü"), rows, agg.workload_metrics, workload_columns),
        _tree_table(pgettext(C.CTX, "Tədris fəaliyyəti"), rows, agg.activity_metrics, activity_columns),
        quality,
    ]


def xlsx_bytes(tables) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    fill = PatternFill("solid", fgColor=_HEADER_FILL)
    font = Font(name="Calibri", color=_HEADER_FONT, bold=True)
    for index, table in enumerate(tables):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = str(table["title"])[:31]
        sheet_append(sheet, table["header"])
        for cell in sheet[1]:
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in table["rows"]:
            sheet_append(sheet, ["" if value is None else value for value in row])
        sheet.freeze_panes = "A2"
        for column, header in enumerate(table["header"], start=1):
            values = [header, *(row[column - 1] for row in table["rows"][:300] if column - 1 < len(row))]
            width = min(max(len(str(value or "")) for value in values) + 2, 48)
            sheet.column_dimensions[get_column_letter(column)].width = max(width, 6)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["build_tables", "xlsx_bytes"]
