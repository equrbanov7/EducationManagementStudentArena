"""«Forma 7» müəllim siyahısının OXUNMASI və normallaşdırılması (sahib 2026-10-11) — DB-siz, saf funksiyalar.

Sahib: «Sənə aktiv olaraq olan bütün müəllimləri atdım — hal-hazırda kim işləyirsə o. Digər hesablar
bloklanıb dayandırılacaq.» Fayl kafedra bloklarından ibarətdir:

    «<ad> kafedrası»                      ← blok başlığı
    N | Soyad, ad, ata adı | Elmi dərəcə | Ştat vahidi   ← sütun başlığı
    1 | Soyad Ad Ata [oğlu|qızı]  | …    | ştat         ← şəxs sətirləri
      | Cəmi                                            ← blokun sonu

Dəyərlər səliqəsizdir — burada hamısı kanonik formaya salınır:

* ad: «Soyad Ad Ata», «Soyad Ad Ata oğlu/q.», «Soyad Ad (baş müəllim)», bəzən «Ad Soyad»;
* elmi ad / dərəcə: «Professor,f.ü.f.d», «Kafedra müdiri,dosent.», «s.e.f.d.,dos», «T.ü.e.d. prof.»;
* ştat: «Ştat/Tam/1,0 ştat» = 1, «0,5 ştat/yarım ştat/y/ş» = 0,5, «s/h, S/H, Saat hesabı» = saathesabı (kənar).

Tanınmayan dəyər UYDURULMUR — ``None`` qayıdır və hesabatda göstərilir.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal

# Kodlar ``UserProfile.AcademicTitle`` / ``AcademicDegree`` ilə eynidir (modul DB-siz qalsın deyə sətir).
TITLE_PROFESSOR = "professor"
TITLE_DOCENT = "docent"
TITLE_SENIOR = "senior_teacher"
TITLE_TEACHER = "teacher"
TITLE_ASSISTANT = "assistant"
DEGREE_PHD = "phd"
DEGREE_DSC = "dsc"

#: Bir şəxs faylda bir neçə dəfə gəlirsə ən yüksək ad götürülür.
TITLE_RANK = {TITLE_ASSISTANT: 1, TITLE_TEACHER: 2, TITLE_SENIOR: 3, TITLE_DOCENT: 4, TITLE_PROFESSOR: 5}
DEGREE_RANK = {DEGREE_PHD: 1, DEGREE_DSC: 2}

RATE_FULL = "full"
RATE_HALF = "half"
RATE_HOURLY = "hourly"
RATE_UNKNOWN = "unknown"

#: Ata adından sonra gələn «oğlu/qızı» (və qısaltmaları) — müqayisədə atılır.
_PATRONYMIC_SUFFIXES = frozenset({"oglu", "ogly", "qizi", "kizi", "gizi", "o", "q"})

_FOLD = str.maketrans({"ə": "e", "ı": "i", "ş": "s", "ç": "c", "ğ": "g", "ö": "o", "ü": "u", "â": "a", "î": "i"})
_YEAR_RE = re.compile(r"(20\d\d)\s*[/\-–]\s*(20\d\d)")


def fold(text) -> str:
    """Registr/diakritik/boşluq dözümlü açar: «Qazı-zadə Elnarə» → «qazi zade elnare»."""
    value = unicodedata.normalize("NFC", str(text or "")).replace("İ", "i").replace("I", "ı").lower()
    value = unicodedata.normalize("NFC", unicodedata.normalize("NFD", value).replace("̇", ""))
    value = value.translate(_FOLD)
    return " ".join(re.sub(r"[^a-z0-9]+", " ", value).split())


def key(text) -> str:
    """Ayırıcısız açar (ad müqayisəsi üçün): «Qazı-zadə» == «Qazızadə» == «qazi zade»."""
    return fold(text).replace(" ", "")


def patronymic_key(text) -> str:
    """Ata adının açarı — «Adil qızı», «Adil q.», «Adil» eyni açarı verir."""
    tokens = [token for token in fold(text).split() if token not in _PATRONYMIC_SUFFIXES]
    return "".join(tokens)


@dataclass(frozen=True)
class PersonName:
    last: str
    first: str
    patronymic: str = ""
    note: str = ""  # mötərizədəki qeyd: «(baş müəllim)», «(kaf.müdiri)»
    two_part: bool = False  # «Soyad Ad» (və ya «Ad Soyad») — ata adı yoxdur

    @property
    def last_key(self) -> str:
        return key(self.last)

    @property
    def first_key(self) -> str:
        return key(self.first)

    @property
    def patronymic_key(self) -> str:
        return patronymic_key(self.patronymic)

    @property
    def display(self) -> str:
        return " ".join(part for part in (self.last, self.first, self.patronymic) if part)


def parse_name(raw) -> PersonName | None:
    """«Soyad Ad Ata [oğlu]» → :class:`PersonName`; bir sözlük/boş mətn → ``None``."""
    text = unicodedata.normalize("NFC", str(raw or "")).strip()
    notes = [" ".join(chunk.split()) for chunk in re.findall(r"\(([^)]*)\)", text)]
    text = re.sub(r"\([^)]*\)", " ", text)
    tokens = [token.strip(",;") for token in text.split() if token.strip(",;.")]
    while len(tokens) > 2 and key(tokens[-1]) in _PATRONYMIC_SUFFIXES:
        tokens.pop()
    if len(tokens) < 2:
        return None
    note = "; ".join(chunk for chunk in notes if chunk)
    if len(tokens) == 2:
        return PersonName(last=tokens[0], first=tokens[1], note=note, two_part=True)
    return PersonName(last=tokens[0], first=tokens[1], patronymic=" ".join(tokens[2:]), note=note)


def parse_title_degree(text, extra: str = "") -> tuple[str | None, str | None]:
    """Elmi ad + elmi dərəcə: («Kafedra müdiri,dosent.», «i.ü.f.d») → (``docent``, ``phd``).

    Vəzifə sözləri (kafedra müdiri, dekan) ELMİ AD deyil — nəzərə alınmır; «doktorant», «dissertant»,
    «-» heç nə vermir.
    """
    raw = " ".join(part for part in (str(text or ""), str(extra or "")) if part)
    lowered = unicodedata.normalize("NFC", raw).replace("İ", "i").replace("I", "ı").lower()
    words = fold(raw)
    compact = re.sub(r"\s+", "", lowered.replace(",", "."))

    title = None
    if re.search(r"\bprof", words):
        title = TITLE_PROFESSOR
    elif re.search(r"\bdos(?:ent)?\b", words):
        title = TITLE_DOCENT
    elif re.search(r"\bbas\s*muel", words):
        title = TITLE_SENIOR
    elif re.search(r"\bassist", words):
        title = TITLE_ASSISTANT
    elif re.search(r"\bmuellim\b", words):
        title = TITLE_TEACHER

    degree = None
    if "elmler doktoru" in words or re.search(r"(?:^|[.\-])e\.d(?:\.|$|[^a-z])", compact) or compact == "e.d":
        degree = DEGREE_DSC
    elif "felsefe doktoru" in words or "phd" in words.replace(" ", "") or "f.d" in compact:
        degree = DEGREE_PHD
    elif re.search(r"(?:^|[.\-])e\.n(?:\.|$|[^a-z])", compact):  # «i.e.n» — elmlər namizədi = fəlsəfə doktoru
        degree = DEGREE_PHD
    return title, degree


@dataclass(frozen=True)
class Rate:
    kind: str
    fraction: Decimal | None = None
    external: bool | None = None


def parse_rate(text) -> Rate:
    """«Ştat vahidi» sütunu → :class:`Rate` (tanınmasa ``RATE_UNKNOWN``)."""
    words = fold(text)
    compact = words.replace(" ", "")
    raw_compact = re.sub(r"\s+", "", str(text or "").lower())
    if not compact:
        return Rate(RATE_UNKNOWN)
    if compact == "sh" or "saathesab" in compact or re.fullmatch(r"s[/\\.]h\.?", raw_compact):
        return Rate(RATE_HOURLY, Decimal("0"), True)
    if re.search(r"0[.,]5", raw_compact) or "yarim" in compact or re.match(r"y/", raw_compact):
        return Rate(RATE_HALF, Decimal("0.5"), False)
    if re.search(r"1[.,]5", raw_compact):
        return Rate(RATE_FULL, Decimal("1.5"), False)
    if compact.startswith(("stat", "tam")) or re.search(r"1[.,]0", raw_compact):
        return Rate(RATE_FULL, Decimal("1"), False)
    return Rate(RATE_UNKNOWN)


@dataclass(frozen=True)
class RosterRow:
    row_no: int
    chair: str
    seq: str
    raw_name: str
    name: PersonName
    raw_degree: str
    raw_rate: str
    title: str | None
    degree: str | None
    rate: Rate


@dataclass
class RosterFile:
    academic_year: str = ""
    chairs: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    skipped: list = field(default_factory=list)  # [(sətir, mətn, səbəb)]


def _cell(row, index) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    value = row[index]
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _is_chair_header(first: str, rest: list) -> bool:
    return bool(first) and not any(rest) and "kafedra" in fold(first)


def parse_sheet_rows(rows) -> RosterFile:
    """Vərəqin sətirləri (``values_only`` tuple-ları) → :class:`RosterFile`."""
    result = RosterFile()
    chair = ""
    for row_no, row in enumerate(rows, start=1):
        cells = [_cell(row, index) for index in range(4)]
        first, name_cell = cells[0], cells[1]
        if not any(cells):
            continue
        if not result.academic_year:
            match = _YEAR_RE.search(" ".join(cells))
            if match and int(match.group(2)) == int(match.group(1)) + 1:
                result.academic_year = f"{match.group(1)}/{match.group(2)}"
        if _is_chair_header(first, cells[1:]):
            chair = first
            result.chairs.append(chair)
            continue
        if fold(first) == "n" or fold(name_cell) == "cemi":
            continue
        if not re.fullmatch(r"\d+\.?", first):
            if chair and name_cell:
                result.skipped.append((row_no, name_cell, "no_sequence_number"))
            continue
        if not chair:
            result.skipped.append((row_no, name_cell, "outside_chair_block"))
            continue
        name = parse_name(name_cell)
        if name is None:
            result.skipped.append((row_no, name_cell, "name_unparsed"))
            continue
        title, degree = parse_title_degree(cells[2], name.note)
        result.rows.append(
            RosterRow(
                row_no=row_no,
                chair=chair,
                seq=first.rstrip("."),
                raw_name=name_cell,
                name=name,
                raw_degree=cells[2],
                raw_rate=cells[3],
                title=title,
                degree=degree,
                rate=parse_rate(cells[3]),
            )
        )
    return result


def read_workbook(path, sheet: str = "") -> RosterFile:
    """xlsx faylını oxuyur; vərəq adı verilməyibsə «Aktiv - Passiv», yoxdursa birinci vərəq."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        names = workbook.sheetnames
        wanted = sheet or ("Aktiv - Passiv" if "Aktiv - Passiv" in names else names[0])
        if wanted not in names:
            raise ValueError(f"sheet_not_found: {wanted!r} ({', '.join(names)})")
        return parse_sheet_rows(workbook[wanted].iter_rows(values_only=True))
    finally:
        workbook.close()


__all__ = [
    "DEGREE_DSC",
    "DEGREE_PHD",
    "DEGREE_RANK",
    "RATE_FULL",
    "RATE_HALF",
    "RATE_HOURLY",
    "RATE_UNKNOWN",
    "TITLE_RANK",
    "PersonName",
    "Rate",
    "RosterFile",
    "RosterRow",
    "fold",
    "key",
    "parse_name",
    "parse_rate",
    "parse_sheet_rows",
    "parse_title_degree",
    "patronymic_key",
    "read_workbook",
]
