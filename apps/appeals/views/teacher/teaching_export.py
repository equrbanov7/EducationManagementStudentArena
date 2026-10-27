"""«Fənnimə aid apellyasiyalar» — Excel ixracı (ekrandakı süzülmüş siyahının özü, səhifələməsiz; sahib 2026-10-27).

Xanalar ``core.export_safety.sheet_cell`` ilə yazılır (formula neytrallaşdırması — CSV/XLSX injection).
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from django.utils import timezone
from django.utils.translation import pgettext

from core.export_safety import sheet_cell

from .teaching_rows import CTX

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
#: Tək faylda ən çox sətir (təşkilat miqyaslı nəzarət görünüşündə yaddaşı məhdudlaşdırır).
MAX_EXPORT_ROWS = 5000


def _columns():
    return [
        ("index", "№", 6),
        ("student_short", pgettext(CTX, "Tələbə"), 26),
        ("patronymic", pgettext(CTX, "Ata adı"), 16),
        ("student_username", pgettext(CTX, "İstifadəçi adı"), 18),
        ("group", pgettext(CTX, "Qrup"), 14),
        ("subject", pgettext(CTX, "Fənn"), 28),
        ("exam_title", pgettext(CTX, "İmtahan"), 30),
        ("category_label", pgettext(CTX, "Kateqoriya"), 12),
        ("item_count", pgettext(CTX, "Sual sayı"), 10),
        ("accepted", pgettext(CTX, "Qəbul"), 9),
        ("rejected", pgettext(CTX, "Rədd"), 9),
        ("status_label", pgettext(CTX, "Status"), 18),
        ("score_before", pgettext(CTX, "İlkin bal"), 11),
        ("score_after", pgettext(CTX, "Yeni bal"), 11),
        ("delta", pgettext(CTX, "Dəyişmə"), 10),
        ("created_at", pgettext(CTX, "Göndərilib"), 17),
        ("decided_at", pgettext(CTX, "Qərar tarixi"), 17),
        ("decided_by", pgettext(CTX, "Qərarı verən müəllim"), 26),
        ("reviewer", pgettext(CTX, "Yadda saxlayan"), 22),
        ("note", pgettext(CTX, "İmtahan mərkəzinin qeydi"), 40),
    ]


#: Rəqəm kimi yazılan sütunlar (mətn «+1» formula qoruyucusundan apostrofla keçirdi).
_NUMERIC = frozenset({"score_before", "score_after", "delta"})


def _number(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    return int(number) if number.is_integer() else round(number, 2)


def _cell(value):
    if isinstance(value, datetime):
        value = timezone.localtime(value) if timezone.is_aware(value) else value
        return value.replace(tzinfo=None)
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    return str(value)


def build_xlsx(rows, *, truncated=False) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Appeals"
    columns = _columns()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    wrap = Alignment(vertical="center", wrap_text=True)
    for column, (_key, header, width) in enumerate(columns, start=1):
        cell = sheet_cell(sheet, row=1, column=column, value=header)
        cell.font, cell.fill, cell.alignment = header_font, header_fill, wrap
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row_index, row in enumerate(rows, start=2):
        values = {**row, "index": row_index - 1}
        for column, (key, _header, _width) in enumerate(columns, start=1):
            value = _cell(_number(values.get(key)) if key in _NUMERIC else values.get(key))
            cell = sheet_cell(sheet, row=row_index, column=column, value=value)
            cell.alignment = wrap
            if isinstance(value, datetime):
                cell.number_format = "DD.MM.YYYY HH:MM"
    if truncated:
        note = pgettext(CTX, "Fayl ilk %(n)s sətirlə məhdudlaşdırılıb — süzgəcləri daraldıb yenidən yükləyin.")
        cell = sheet_cell(sheet, row=len(rows) + 3, column=2, value=note % {"n": MAX_EXPORT_ROWS})
        cell.font = Font(bold=True, color="B91C1C")
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename() -> str:
    return f"apellyasiyalar-{timezone.localtime().strftime('%Y%m%d-%H%M')}.xlsx"


__all__ = ["MAX_EXPORT_ROWS", "XLSX_CONTENT_TYPE", "build_xlsx", "export_filename"]
