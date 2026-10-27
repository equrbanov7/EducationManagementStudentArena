"""2026-10-27 (sahib) — «Fənnimə aid apellyasiyalar» kabinet bölməsi (`?section=teaching-appeals`).

* bölmə ayrıca səhifə deyil — kabinetin içində (sol menyu qalır) açılır; köhnə ``/appeals/teaching/`` yönləndirir;
* əhatə: müəllim — yalnız öz açılışının (sahib / slot / bölünmüş jurnal payı) tələbələrinin həmin fənn üzrə
  apellyasiyaları; inzibati rolu olan (dekan) dərs deyən müəllim də görür; superadmin / RİM rəhbəri — bütün
  təşkilat; başqa müəllim və tələbə — heç nə;
* süzgəclər (status, kateqoriya, fənn, axtarış, tarix), çekməcə endpoint-inin icazəsi, Excel ixracı.
"""

from __future__ import annotations

from io import BytesIO
from urllib.parse import unquote

from django.test import Client
from django.urls import reverse

from apps.appeals.models import Appeal, AppealItem
from apps.appeals.services.teaching_scope import scoped_appeals, teaching_section_visible
from apps.appeals.tests.test_key_correction_2026_10_13 import _AJAX, COMMENT
from apps.appeals.tests.test_review_teacher_document_2026_10_21 import User, _ReviewSetup
from apps.organizations.models import Membership
from apps.registrar.models import CourseOffering, JournalShare
from core.rls import bypass_rls

SECTION_URL = "/accounts/profile/?section=teaching-appeals"


class _SectionSetup(_ReviewSetup):
    def setUp(self):
        super().setUp()
        with bypass_rls():
            self.rim_head = User.objects.create_user("ta_rim", "ta_rim@qku.edu.az", "pw")
            Membership.objects.create(
                user=self.rim_head,
                organization=self.org,
                role=self.org.roles.get(name="ikt_rehber"),
                is_primary=True,
                is_active=True,
            )

    def _section(self, user, query=""):
        return self._client(user).get(SECTION_URL + query)

    def _allowed(self, user):
        response = self._client(user).get(reverse("accounts:profile"))
        self.assertEqual(response.status_code, 200)
        return set(response.context["allowed_sections"])


class CabinetSectionTests(_SectionSetup):
    def test_section_renders_inside_cabinet_with_sidebar(self):
        response = self._section(self.subject_teacher)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-profile-section-panel="teaching-appeals"', html=False)
        self.assertContains(response, 'data-section="teaching-appeals"', html=False)  # sol menyu bəndi
        self.assertContains(response, "sidebar-menu-link", html=False)
        self.assertContains(response, self.exam.title)
        self.assertContains(response, "appeals/js/teaching_appeals.js", html=False)
        self.assertContains(response, "data-ta-open-appeal", html=False)
        self.assertNotContains(response, "data-appeals-review-form", html=False)  # qərar hüququ yoxdur

    def test_ajax_fragment_endpoint_serves_section(self):
        url = reverse("accounts:profile_section_fragment", args=["teaching-appeals"])
        response = self._client(self.subject_teacher).get(url, **_AJAX)
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.exam.title, response.content.decode())

    def test_old_url_redirects_to_section_keeping_filters(self):
        response = self._client(self.subject_teacher).get(
            reverse("appeals:teaching_appeals") + f"?status=pending&open={self.appeal.id}"
        )
        self.assertEqual(response.status_code, 302)
        location = unquote(response["Location"])
        self.assertIn("section=teaching-appeals", location)
        self.assertIn("ta_status=pending", location)
        self.assertIn(f"ta_open={self.appeal.id}", location)

    def test_open_param_marks_panel_for_auto_drawer(self):
        response = self._section(self.subject_teacher, f"&ta_open={self.appeal.id}")
        self.assertContains(response, f'data-ta-open="{self.appeal.id}"', html=False)

    def test_notification_link_points_to_section(self):
        from apps.appeals.services.notify import _teaching_link

        self.assertEqual(
            _teaching_link(self.appeal), f"/accounts/profile/?section=teaching-appeals&ta_open={self.appeal.id}"
        )


class ScopeTests(_SectionSetup):
    def test_subject_teacher_sees_own_and_other_teacher_sees_none(self):
        self.assertIn("teaching-appeals", self._allowed(self.subject_teacher))
        self.assertIn(self.appeal, scoped_appeals(self.subject_teacher, self.org)[0])
        other, org_wide = scoped_appeals(self.other_teacher, self.org)
        self.assertFalse(org_wide)
        self.assertFalse(other.exists())
        response = self._section(self.other_teacher)
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, self.exam.title)

    def test_teacher_with_admin_role_still_sees_menu_and_own_subjects(self):
        """Dekan (müəllim rolu yoxdur) açılışın jurnal sahibidirsə — bölmə görünür və öz fənnini görür."""
        CourseOffering.objects.filter(pk=self.offering.pk).update(instructor=self.dean)
        self.assertTrue(teaching_section_visible(self.dean, self.org))
        self.assertIn("teaching-appeals", self._allowed(self.dean))
        self.assertIn(self.appeal, scoped_appeals(self.dean, self.org)[0])
        self.assertContains(self._section(self.dean), self.exam.title)
        # Dərs deməyən dekan menyunu görmür.
        self.assertNotIn("teaching-appeals", self._allowed(self.other_dean))

    def test_journal_share_counts_as_teaching(self):
        with bypass_rls():
            JournalShare.objects.create(organization=self.org, offering=self.offering, teacher=self.dean)
        self.assertTrue(teaching_section_visible(self.dean, self.org))
        self.assertIn(self.appeal, scoped_appeals(self.dean, self.org)[0])

    def test_journal_share_with_students_and_inactive_share(self):
        with bypass_rls():
            share = JournalShare.objects.create(organization=self.org, offering=self.offering, teacher=self.dean)
            share.enrollments.set(self.offering.enrollments.filter(student=self.student))
        self.assertIn(self.appeal, scoped_appeals(self.dean, self.org)[0])
        JournalShare.objects.filter(pk=share.pk).update(is_active=False)
        self.assertFalse(scoped_appeals(self.dean, self.org)[0].exists())
        self.assertFalse(teaching_section_visible(self.dean, self.org))

    def test_rim_head_and_superadmin_see_all_appeals(self):
        for user in (self.rim_head, self.reviewer):
            with self.subTest(user=user.username):
                self.assertIn("teaching-appeals", self._allowed(user))
                queryset, org_wide = scoped_appeals(user, self.org)
                self.assertTrue(org_wide)
                self.assertIn(self.appeal, queryset)
                response = self._section(user)
                self.assertContains(response, self.exam.title)
                self.assertContains(response, "nəzarət görünüşü")

    def test_student_has_no_section(self):
        self.assertFalse(teaching_section_visible(self.student, self.org))
        self.assertNotIn("teaching-appeals", self._allowed(self.student))
        response = self._section(self.student)
        self.assertNotContains(response, 'data-profile-section-panel="teaching-appeals"', html=False)


class FilterTests(_SectionSetup):
    def test_status_filter(self):
        self.assertContains(self._section(self.subject_teacher, "&ta_status=pending"), self.exam.title)
        self.assertNotContains(self._section(self.subject_teacher, "&ta_status=accepted_any"), self.exam.title)

    def test_category_filter(self):
        self.exam.exam_type_extended = "final"
        self.exam.save(update_fields=["exam_type_extended"])
        self.assertContains(self._section(self.subject_teacher, "&ta_category=final"), self.exam.title)
        self.assertNotContains(self._section(self.subject_teacher, "&ta_category=midterm"), self.exam.title)

    def test_subject_and_search_filters(self):
        self.assertContains(self._section(self.subject_teacher, f"&ta_subject={self.subject.pk}"), self.exam.title)
        self.assertNotContains(self._section(self.subject_teacher, "&ta_subject=999999"), self.exam.title)
        self.assertContains(self._section(self.subject_teacher, "&ta_q=a3_student"), self.exam.title)
        self.assertNotContains(self._section(self.subject_teacher, "&ta_q=yoxdur_bele"), self.exam.title)

    def test_date_filter_and_invalid_values_ignored(self):
        self.assertNotContains(self._section(self.subject_teacher, "&ta_from=2999-01-01"), self.exam.title)
        self.assertContains(self._section(self.subject_teacher, "&ta_from=bad&ta_status=bogus"), self.exam.title)
        # Formatı düz, tarixi səhv dəyər və unicode rəqəm 500 verməməlidir (süzgəc sadəcə atılır).
        for query in ("&ta_from=2026-02-30", "&ta_to=2026-13-01", "&ta_subject=%C2%B2"):
            with self.subTest(query=query):
                self.assertContains(self._section(self.subject_teacher, query), self.exam.title)

    def test_kpis_present(self):
        response = self._section(self.subject_teacher)
        self.assertContains(response, 'data-ems-kpi-filter="pending"', html=False)
        self.assertContains(response, "Orta bal dəyişməsi")


class DrawerTests(_SectionSetup):
    def _detail(self, user, ajax=True):
        url = reverse("appeals:teaching_appeal_detail", args=[self.appeal.id])
        return self._client(user).get(url, **(_AJAX if ajax else {}))

    def test_subject_teacher_gets_drawer_html(self):
        response = self._detail(self.subject_teacher)
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["ok"])
        self.assertIn(COMMENT, payload["html"])
        self.assertIn("Tarixçə", payload["html"])
        self.assertIn(self.exam.title, payload["subtitle"])

    def test_drawer_shows_decision_and_hides_document_download_from_teacher(self):
        from apps.appeals.tests.test_review_teacher_document_2026_10_21 import _pdf

        self.assertEqual(self._review_post(document=_pdf()).status_code, 200)
        html = self._detail(self.subject_teacher).json()["html"]
        self.assertIn("Haqlıdır", html)
        self.assertIn("qerar.pdf", html)
        self.assertIn("yalnız imtahan mərkəzi endirə bilər", html)
        self.assertIn("Səbinə Fənnova", html)

    def test_out_of_scope_users_get_404(self):
        for user in (self.other_teacher, self.student):
            with self.subTest(user=user.username):
                self.assertEqual(self._detail(user).status_code, 404)

    def test_rim_head_can_open_drawer(self):
        self.assertEqual(self._detail(self.rim_head).status_code, 200)

    def test_non_ajax_redirects_to_section(self):
        response = self._detail(self.subject_teacher, ajax=False)
        self.assertEqual(response.status_code, 302)
        self.assertIn(f"ta_open={self.appeal.id}", response["Location"])

    def test_anonymous_redirected_to_login(self):
        response = Client().get(reverse("appeals:teaching_appeal_detail", args=[self.appeal.id]), **_AJAX)
        self.assertEqual(response.status_code, 302)


class ExportTests(_SectionSetup):
    def _rows(self, user, query=""):
        from openpyxl import load_workbook

        response = self._client(user).get(reverse("appeals:teaching_appeals_export") + query)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment;", response["Content-Disposition"])
        sheet = load_workbook(BytesIO(response.content)).active
        return [row for row in sheet.iter_rows(min_row=2, values_only=True)]

    def test_teacher_exports_own_appeals(self):
        rows = self._rows(self.subject_teacher)
        self.assertEqual(len(rows), 1)
        self.assertIn(self.exam.title, rows[0])
        self.assertIn("a3_student", rows[0])

    def test_export_respects_filters_and_scope(self):
        self.assertEqual(self._rows(self.subject_teacher, "?ta_status=rejected"), [])
        self.assertEqual(self._rows(self.other_teacher), [])
        self.assertEqual(len(self._rows(self.rim_head)), Appeal.objects.filter(organization=self.org).count())

    def test_export_marks_truncation(self):
        from unittest import mock

        with mock.patch("apps.appeals.views.teacher.teaching.MAX_EXPORT_ROWS", 0):
            response = self._client(self.subject_teacher).get(reverse("appeals:teaching_appeals_export"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Export-Truncated"], "0")

    def test_export_shows_score_change_after_decision(self):
        self.assertEqual(self._review_post().status_code, 200)
        self.assertEqual(AppealItem.objects.get(pk=self.item.pk).status, "accepted")
        row = self._rows(self.subject_teacher)[0]
        self.assertEqual(row[12:15], (1, 2, 1))  # ilkin bal → yeni bal, dəyişmə (rəqəm)
