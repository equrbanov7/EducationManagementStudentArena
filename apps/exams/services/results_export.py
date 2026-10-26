"""Vahid nəticə ixracı (Excel) — sahib 2026-10-27.

«Dəyərləndirilmiş nəticələr» (`/accounts/profile/?section=review-results`) və imtahanın öz
nəticə səhifəsi (`/exams/<slug>/results/`) EYNİ sütunlarla ixrac edir: ``COLUMNS``. Sətirlər
çağıranın verdiyi sıradadır (ekrandakı süzgəc + sıralama) — burada süzgəc yoxdur.

* :func:`exam_attempt_rows` — imtahan cəhdlərindən sətir lüğətləri (bal / maks / faiz, test
  sayğacları, «Yoxlayan», zal + kompüter, müddət, status, uzaqlaşdırma səbəbi, apellyasiya).
  Sorğu sayı cəhd sayından asılı DEYİL (prefetch + tək xəritə sorğuları).
* :func:`build_results_xlsx` — sətirləri `.xlsx`-ə yazır; mətn xanaları formula kimi şərh
  olunmasın deyə ``core.export_safety.sheet_cell``-dən keçir (audit F-07).

Başlıqlar Azərbaycanca sabitdir (mövcud ixracın davamı; Excel sütun adları dildən asılı deyil).
Tarixlər yerli zonaya (Asia/Baku) çevrilib tz-siz yazılır — Excel UI ilə eyni saatı göstərir.
Loglara heç bir şəxsi məlumat yazılmır.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from django.db.models import prefetch_related_objects
from django.utils import timezone
from django.utils.text import slugify

from core.export_safety import sheet_cell
from core.person_names import clean_name_part

from .results_export_meta import checked_by_map, location_map

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: (açar, başlıq, en, mətn-sola) — hər iki səhifənin ixracı bu sırada.
COLUMNS = (
    ("index", "#", 5, False),
    ("groups", "Qruplar", 26, True),
    ("parent_group", "Ana qrup", 18, True),
    ("student", "Ad Soyad", 26, True),
    ("username", "İstifadəçi adı", 20, True),
    ("patronymic", "Ata adı", 16, True),
    ("email", "E-poçt", 28, True),
    ("title", "İmtahan / iş", 30, True),
    ("subject", "Fənn", 26, True),
    ("kind", "Növ", 16, True),
    ("category", "Kateqoriya", 14, True),
    ("score", "Bal", 8, False),
    ("max_score", "Maks. bal", 10, False),
    ("percent", "Faiz", 8, False),
    ("appeal_bonus", "Apel. bonus", 11, False),
    ("correct", "Düzgün", 9, False),
    ("wrong", "Səhv", 8, False),
    ("unanswered", "Cavabsız", 10, False),
    ("delivered", "Verilmiş sual", 13, False),
    ("checked_by", "Yoxlayan", 24, True),
    ("checked_at", "Yoxlanma tarixi", 18, True),
    ("started_at", "Başlama", 18, True),
    ("finished_at", "Bitmə", 18, True),
    ("duration", "Müddət (s:dq:sn)", 15, False),
    ("hall", "Zal", 18, True),
    ("computer", "Kompüter", 16, True),
    ("attempt_number", "Cəhd №", 8, False),
    ("status", "Status", 16, True),
    ("removal_reason", "Uzaqlaşdırma səbəbi", 30, True),
    ("appeal", "Apellyasiya", 18, True),
)
HEADERS = tuple(header for _key, header, _width, _left in COLUMNS)


def _number(value):
    if value in (None, ""):
        return ""
    number = float(value)
    return int(number) if number.is_integer() else round(number, 2)


def format_duration(seconds) -> str:
    if seconds is None:
        return ""
    total = max(int(seconds), 0)
    return f"{total // 3600:02d}:{(total % 3600) // 60:02d}:{total % 60:02d}"


def patronymic_map(user_ids) -> dict[int, str]:
    """``user_id`` → ata adı (profil; idxal yer tutucuları boş sayılır) — tək sorğu, ORM əlaqəsi ilə."""
    from django.contrib.auth import get_user_model

    ids = {user_id for user_id in user_ids if user_id}
    if not ids:
        return {}
    rows = get_user_model().objects.filter(pk__in=ids).values_list("pk", "profile__patronymic")
    return {pk: clean_name_part(value) for pk, value in rows if value}


def subject_title(exam) -> str:
    """İmtahanın fənni — akademik fənn, yoxdursa kursun adı."""
    if exam.subject_id and exam.subject is not None:
        return exam.subject.name
    return exam.course.title if exam.course_id and exam.course is not None else ""


def _score_numbers(attempts):
    """``attempt_id`` → bal / maks / faiz / bonus / sayğaclar (test: snapshot + apellyasiya bonusu)."""
    from apps.exams import score_adjustments
    from apps.exams.views.teacher.results._test_scores import attach_legacy_selected_options

    from .manual_grading import answer_max_points
    from .result_calculation import calculate_test_attempt_result

    prefetch_related_objects(attempts, "answers__question__options")
    tests = [attempt for attempt in attempts if attempt.exam.exam_type == "test"]
    if tests:
        attach_legacy_selected_options(tests)
    bonuses = score_adjustments.bonus_map([attempt.id for attempt in tests]) if tests else {}
    numbers = {}
    for attempt in attempts:
        if attempt.exam.exam_type == "test":
            result = calculate_test_attempt_result(attempt, answers=attempt.answers.all())
            bonus = bonuses.get(attempt.id) or 0
            if bonus:
                result = score_adjustments.apply_bonus(result, bonus)
            numbers[attempt.id] = {
                "score": _number(result.score),
                "max_score": _number(result.max_score),
                "percent": _number(result.percentage),
                "appeal_bonus": _number(bonus) if bonus else 0,
                "correct": result.correct_count,
                "wrong": result.wrong_count,
                "unanswered": result.unanswered_count,
                "delivered": result.delivered_count,
            }
            continue
        max_score = sum(answer_max_points(answer) for answer in attempt.answers.all())
        score = attempt.teacher_score
        percent = round(float(score) * 100.0 / float(max_score), 1) if score is not None and max_score else ""
        numbers[attempt.id] = {
            "score": _number(score),
            "max_score": max_score or "",
            "percent": _number(percent),
            "appeal_bonus": "",
            "correct": attempt.correct_count,
            "wrong": attempt.wrong_count,
            "unanswered": "",
            "delivered": attempt.correct_count + attempt.wrong_count,
        }
    return numbers


def exam_attempt_rows(attempts, *, group_lookup=None, now=None) -> dict[int, dict]:
    """``attempt_id`` → sətir lüğəti (``COLUMNS`` açarları; «#» yazılarkən qoyulur).

    ``group_lookup(user_id) → (qruplar, ana qrup)`` — çağıranın qrup mənbəyi (imtahan səhifəsi:
    kohort + reyestr indeksi; «Dəyərləndirilmiş nəticələr»: sətrin reyestr qrupu).
    """
    from apps.exams import score_adjustments
    from apps.exams.views.teacher.results._helpers import _attempt_effective_duration, _attempt_effective_finish

    from .supervision import attach_attempt_interventions

    attempts = [attempt for attempt in attempts if attempt.pk]
    if not attempts:
        return {}
    now = now or timezone.now()
    attach_attempt_interventions(attempts)
    numbers = _score_numbers(attempts)
    checkers = checked_by_map(attempts)
    locations = location_map(attempts)
    appeals = score_adjustments.status_by_attempt([attempt.id for attempt in attempts])
    patronymics = patronymic_map(attempt.user_id for attempt in attempts)

    rows = {}
    for attempt in attempts:
        exam = attempt.exam
        finish, _inferred = _attempt_effective_finish(attempt, now=now)
        groups, parent = group_lookup(attempt.user_id) if group_lookup else ("", "")
        intervention = getattr(attempt, "exam_intervention", None) or {}
        hall, computer = locations.get(attempt.id, ("", ""))
        appeal = appeals.get(attempt.id)
        rows[attempt.id] = {
            "groups": groups,
            "parent_group": parent,
            "student": attempt.user.get_full_name() or attempt.user.username,
            "username": attempt.user.username,
            "patronymic": patronymics.get(attempt.user_id, ""),
            "email": attempt.user.email or "",
            "title": exam.title,
            "subject": subject_title(exam),
            "kind": exam.get_exam_type_display(),
            "category": exam.get_exam_type_extended_display() if exam.exam_type_extended else "",
            **numbers[attempt.id],
            "checked_by": checkers.get(attempt.id, ""),
            "checked_at": attempt.teacher_checked_at,
            "started_at": attempt.started_at,
            "finished_at": finish,
            "duration": format_duration(_attempt_effective_duration(attempt, finish)),
            "hall": hall,
            "computer": computer,
            "attempt_number": attempt.attempt_number,
            "status": attempt.get_status_display(),
            "removal_reason": (intervention.get("reason") if intervention.get("is_terminal") else "") or "",
            "appeal": appeal["label"] if appeal else "",
        }
    return rows


def _cell_value(value):
    if isinstance(value, datetime):
        value = timezone.localtime(value) if timezone.is_aware(value) else value
        return value.replace(tzinfo=None)
    if value is None:
        return ""
    return str(value) if not isinstance(value, (int, float)) else value


def build_results_xlsx(rows, *, sheet_title: str, filename_stem: str):
    """(filename, content_type, bytes) — ``rows``: ``COLUMNS`` açarlı lüğətlər, verilən sırada."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    # Vərəq adında `/ \ ? * [ ] :` qadağandır (openpyxl ValueError) — boşluqla əvəzlənir.
    safe_title = "".join(" " if ch in "/\\?*[]:" else ch for ch in sheet_title).strip("'").strip()
    sheet.title = ((safe_title[:28] + "…") if len(safe_title) > 30 else safe_title) or "Results"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for column, (_key, header, width, _left) in enumerate(COLUMNS, start=1):
        cell = sheet_cell(sheet, row=1, column=column, value=header)
        cell.font, cell.fill, cell.alignment = header_font, header_fill, center
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 28

    for row_index, row in enumerate(rows, start=2):
        values = {**row, "index": row_index - 1}
        for column, (key, _header, _width, is_text) in enumerate(COLUMNS, start=1):
            value = _cell_value(values.get(key, ""))
            cell = sheet_cell(sheet, row=row_index, column=column, value=value)
            cell.alignment = left if is_text else center
            if isinstance(value, datetime):
                cell.number_format = "DD.MM.YYYY HH:MM"
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")  # fayl adı yerli saatla (Asia/Baku)
    filename = f"{slugify(filename_stem) or 'results'}-{stamp}.xlsx"
    return filename, XLSX_CONTENT_TYPE, buffer.getvalue()


__all__ = [
    "COLUMNS",
    "HEADERS",
    "XLSX_CONTENT_TYPE",
    "build_results_xlsx",
    "exam_attempt_rows",
    "format_duration",
    "patronymic_map",
    "subject_title",
]
