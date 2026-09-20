"""Sual şablonunun Word (.docx) variantı (sahib 2026-09-21).

Müəllim TXT-ni açmağa öyrəşməyib: «Word-də açıb baxa bilsin, necə yazmaq lazım
olduğunu görsün». Bu modul MÖVCUD TXT şablon mətnini (bax
``question_bank/_helpers._question_bank_template_txt`` və
``question_library/_shared._bank_template_*``) olduğu kimi .docx-ə çevirir:

* ``#`` ilə başlayan sətirlər — izah/təlimat (boz, kursiv), prefiks SAXLANILIR ki,
  fayl redaktə olunub geri yüklənəndə import parseri onları eyni qaydada ötürsün;
* qalan sətirlər — nümunə suallar (adi abzas), düz cavab işarəsi (``*B)`` /
  ``Cavab: B``) dəyişmir → DOCX importu ilə round-trip mümkündür.
"""

from __future__ import annotations

import io

from django.utils.translation import pgettext

_CTX = "exams.template.test_question_bank"


def build_template_docx(template_text: str, *, title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.core_properties.title = title

    heading = document.add_heading(title, level=1)
    heading.runs[0].font.size = Pt(16)

    intro = document.add_paragraph()
    run = intro.add_run(
        pgettext(
            _CTX,
            "Aşağıdakı nümunə suallara baxın, öz suallarınızı EYNİ formada bu faylda yazın və faylı sistemə yükləyin. "
            "«#» ilə başlayan sətirlər izahdır — import zamanı nəzərə alınmır (silə də bilərsiniz).",
        )
    )
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x5F, 0x6D)

    for line in (template_text or "").splitlines():
        text = line.rstrip()
        if not text:
            document.add_paragraph()
            continue
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.font.size = Pt(11)
        if text.lstrip().startswith("#"):
            run.italic = True
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        elif text.lstrip().startswith("*"):
            run.bold = True

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
