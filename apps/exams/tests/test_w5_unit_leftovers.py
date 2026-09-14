"""W5 `w5left` (2026-09-14): reyestr qrupu (`Exam.allowed_units`) qalıq oxuyanları.

W4 hesabatı «Orkestrator üçün qalanlar» 2–3: imtahan mərkəzi statistikası,
apellyasiya statistikası, «yenidən şans» bölməsi və müəllim nəticələri
qrup/fakültə/kafedra filtrlərini yalnız köhnə kohortla (`StudentGroup.org_unit`)
qururdu; `ExamForm.clean()` istisnaları yalnız kohort üzvləri ilə süzürdü.

Fixture: Fakültə «Dizayn» → Kafedra «Qrafika» → İxtisas «Qrafik dizayn» →
Qrup «634 ing» (reyestr); paralel kohort «Kohort-A» (`org_unit` = kafedra).
"""

from datetime import timedelta

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import Client, RequestFactory, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import ProfileRole
from apps.accounts.views.profile._sections.exam_chance import build_exam_chance_section
from apps.appeals.constants import APPEAL_STATUS_PENDING
from apps.appeals.models import Appeal
from apps.exams.domain.unit_scope_filters import split_group_filter_values, unit_filter_value
from apps.exams.forms import ExamForm
from apps.exams.models import Exam, ExamAttempt, StudentExamAttemptGrant, StudentGroup
from apps.exams.tests.test_exam_center_policy import PASSWORD, _assign_user_to_org
from apps.organizations.models import Organization, OrgUnit
from apps.registrar.models import Curriculum, Program, StudentAcademicRecord
from core.constants import OrganizationType, OrgUnitType
from core.rls import bypass_rls

User = get_user_model()


def _login(user, organization):
    client = Client()
    client.force_login(user)
    session = client.session
    session["active_organization"] = organization.slug
    session.save()
    return client


class _RegistryFixture(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user("w5l_owner", "w5l_owner@test.az", PASSWORD)
        cls.org = Organization.objects.create(
            name="W5L Uni", org_type=OrganizationType.UNIVERSITY, owner=cls.owner, status="active", is_active=True
        )
        cls.center = User.objects.create_user("w5l_center", "w5l_center@test.az", PASSWORD)
        _assign_user_to_org(cls.center, cls.org, ProfileRole.MEMBER, "exam_center_head")
        cls.teacher = User.objects.create_user("w5l_teacher", "w5l_teacher@test.az", PASSWORD)
        _assign_user_to_org(cls.teacher, cls.org, ProfileRole.TEACHER, "teacher")
        cls.student = User.objects.create_user("w5l_student", "w5l_student@test.az", PASSWORD, first_name="Mərvi")
        _assign_user_to_org(cls.student, cls.org, ProfileRole.STUDENT, "student")
        cls.student2 = User.objects.create_user("w5l_student2", "w5l_student2@test.az", PASSWORD, first_name="Aysel")
        _assign_user_to_org(cls.student2, cls.org, ProfileRole.STUDENT, "student")
        cls.outsider = User.objects.create_user("w5l_outsider", "w5l_outsider@test.az", PASSWORD, first_name="Kənar")
        _assign_user_to_org(cls.outsider, cls.org, ProfileRole.STUDENT, "student")

        with bypass_rls():
            cls.faculty = OrgUnit.objects.create(
                organization=cls.org, name="Dizayn", slug="w5l-fac", unit_type=OrgUnitType.FACULTY
            )
            cls.faculty_other = OrgUnit.objects.create(
                organization=cls.org, name="Biznes", slug="w5l-fac2", unit_type=OrgUnitType.FACULTY
            )
            cls.chair = OrgUnit.objects.create(
                organization=cls.org, parent=cls.faculty, name="Qrafika", slug="w5l-chair", unit_type=OrgUnitType.CHAIR
            )
            cls.chair_other = OrgUnit.objects.create(
                organization=cls.org,
                parent=cls.faculty_other,
                name="Marketinq",
                slug="w5l-chair2",
                unit_type=OrgUnitType.CHAIR,
            )
            cls.specialty = OrgUnit.objects.create(
                organization=cls.org,
                parent=cls.chair,
                name="Qrafik dizayn",
                slug="w5l-spec",
                unit_type=OrgUnitType.SPECIALTY,
            )
            cls.group = OrgUnit.objects.create(
                organization=cls.org, parent=cls.specialty, name="634 ing", slug="w5l-634", unit_type=OrgUnitType.GROUP
            )
            cls.group_other = OrgUnit.objects.create(
                organization=cls.org,
                parent=cls.chair_other,
                name="701 biz",
                slug="w5l-701",
                unit_type=OrgUnitType.GROUP,
            )
            program = Program.objects.create(organization=cls.org, code="DZ", name="Dizayn", absence_limit_percent=25)
            curriculum = Curriculum.objects.create(organization=cls.org, program=program, admission_year=2024)
            for student, unit in ((cls.student, cls.group), (cls.student2, cls.group), (cls.outsider, cls.group_other)):
                StudentAcademicRecord.objects.create(
                    organization=cls.org,
                    student=student,
                    program=program,
                    curriculum=curriculum,
                    group=unit,
                    admission_year=2024,
                )

        cls.cohort = StudentGroup.objects.create(
            teacher=cls.teacher, organization=cls.org, name="Kohort-A", org_unit=cls.chair
        )
        cls.cohort.students.add(cls.outsider)

    def _exam(self, title="W5L Final", **kwargs):
        defaults = {
            "title": title,
            "author": self.teacher,
            "organization": self.org,
            "exam_type": "test",
            "exam_type_extended": "final",
            "is_active": True,
            "is_public": False,
            "start_datetime": timezone.now() - timedelta(hours=1),
            "end_datetime": timezone.now() + timedelta(days=2),
        }
        defaults.update(kwargs)
        return Exam.objects.create(**defaults)

    @staticmethod
    def _attempt(exam, user, status="submitted"):
        now = timezone.now()
        return ExamAttempt.objects.create(
            user=user,
            exam=exam,
            status=status,
            attempt_number=ExamAttempt.objects.filter(user=user, exam=exam).count() + 1,
            started_at=now,
            finished_at=now,
        )


class SplitGroupFilterValuesTests(TestCase):
    def test_accepts_cohort_ints_and_unit_uuids(self):
        cohorts, units = split_group_filter_values(
            "12, unit:7f6b0c2a-2a5e-4d7e-9c3b-1a2b3c4d5e6f,7F6B0C2A-2A5E-4D7E-9C3B-1A2B3C4D5E6F,zibil,,3"
        )
        self.assertEqual(cohorts, [12, 3])
        self.assertEqual(units, ["7f6b0c2a-2a5e-4d7e-9c3b-1a2b3c4d5e6f"])


class ExamCenterStatsUnitTests(_RegistryFixture):
    """`exam_center_stats_data` reyestr qrupunu qrup/fakültə/kafedra filtrlərində görür."""

    def setUp(self):
        self.exam = self._exam()
        self.exam.allowed_units.add(self.group)
        self._attempt(self.exam, self.student)
        self.cohort_exam = self._exam(title="W5L Kohort Final")
        self.cohort_exam.allowed_groups.add(self.cohort)
        self._attempt(self.cohort_exam, self.outsider)
        self.client = _login(self.center, self.org)
        self.url = reverse("exams:exam_center_stats_data")

    def _data(self, **params):
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200)
        return response.json()

    def test_group_filter_accepts_unit_value(self):
        data = self._data(groups=unit_filter_value(self.group.pk))
        self.assertEqual(data["summary"]["total"], 1)
        self.assertEqual(data["results"][0]["exam"], "W5L Final")
        # Kohort int id-si əvvəlki kimi işləyir; ikisi birlikdə → hər ikisi.
        self.assertEqual(self._data(groups=str(self.cohort.id))["summary"]["total"], 1)
        both = self._data(groups=f"{self.cohort.id},{self.group.pk}")
        self.assertEqual(both["summary"]["total"], 2)

    def test_faculty_department_and_legacy_units_filters_include_registry_group(self):
        for params in (
            {"faculties": str(self.faculty.pk)},
            {"departments": str(self.chair.pk)},
            {"units": str(self.faculty.pk)},
        ):
            with self.subTest(params=params):
                data = self._data(**params)
                # Reyestr imtahanı + eyni kafedradakı kohort imtahanı (parity).
                self.assertEqual(data["summary"]["total"], 2, params)
                self.assertEqual({row["exam"] for row in data["results"]}, {"W5L Final", "W5L Kohort Final"})
        self.assertEqual(self._data(faculties=str(self.faculty_other.pk))["summary"]["total"], 0)
        self.assertEqual(self._data(departments=str(self.chair_other.pk))["summary"]["total"], 0)

    def test_row_columns_show_registry_group_and_ancestors(self):
        data = self._data(groups=unit_filter_value(self.group.pk))
        row = data["results"][0]
        self.assertEqual(row["group"], "634 ing")
        self.assertEqual(row["kafedra"], "Qrafika")
        self.assertEqual(row["faculty"], "Dizayn")

    def test_query_count_independent_of_row_count(self):
        def count_queries():
            with CaptureQueriesContext(connection) as ctx:
                self._data(faculties=str(self.faculty.pk))
            return len(ctx.captured_queries)

        count_queries()  # isinmə: sessiya/icazə keşləri ilk sorğuda dolur
        baseline = count_queries()
        for _ in range(3):
            self._attempt(self.exam, self.student2)
            self._attempt(self.cohort_exam, self.outsider)
        self.assertEqual(count_queries(), baseline)

    def test_export_includes_registry_group_column(self):
        response = self.client.get(reverse("exams:exam_center_stats_export"), {"groups": str(self.group.pk)})
        self.assertEqual(response.status_code, 200)
        from io import BytesIO

        from openpyxl import load_workbook

        ws = load_workbook(BytesIO(response.content)).active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][2], "634 ing")
        self.assertEqual(rows[0][3], "Qrafika")
        self.assertEqual(rows[0][4], "Dizayn")


class AppealStatsUnitTests(_RegistryFixture):
    def setUp(self):
        self.exam = self._exam()
        self.exam.allowed_units.add(self.group)
        attempt = self._attempt(self.exam, self.student)
        Appeal.objects.create(
            attempt=attempt, exam=self.exam, student=self.student, organization=self.org, status=APPEAL_STATUS_PENDING
        )
        self.client = _login(self.center, self.org)
        self.url = reverse("appeals:appeal_stats_data")

    def _data(self, **params):
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200)
        return response.json()

    def test_filters_and_group_column(self):
        for params in (
            {"groups": unit_filter_value(self.group.pk)},
            {"faculties": str(self.faculty.pk)},
            {"departments": str(self.chair.pk)},
            {"units": str(self.chair.pk)},
        ):
            with self.subTest(params=params):
                data = self._data(**params)
                self.assertEqual(data["summary"]["total"], 1, params)
                self.assertEqual(data["results"][0]["group"], "634 ing")
        self.assertEqual(self._data(faculties=str(self.faculty_other.pk))["summary"]["total"], 0)

    def test_query_count_independent_of_row_count(self):
        def count_queries():
            with CaptureQueriesContext(connection) as ctx:
                self._data(faculties=str(self.faculty.pk))
            return len(ctx.captured_queries)

        count_queries()  # isinmə
        baseline = count_queries()
        for _ in range(3):
            attempt = self._attempt(self.exam, self.student2)
            Appeal.objects.create(
                attempt=attempt,
                exam=self.exam,
                student=self.student2,
                organization=self.org,
                status=APPEAL_STATUS_PENDING,
            )
        self.assertEqual(count_queries(), baseline)


class ExamChanceSectionUnitTests(_RegistryFixture):
    def _section(self, **params):
        request = RequestFactory().get("/", params)
        request.user = self.center
        section = {}
        build_exam_chance_section(
            request,
            section,
            active_organization=self.org,
            allowed_sections={"exam-chance"},
            active_section="exam-chance",
        )
        return section

    def test_faculty_and_kafedra_filters_include_registry_group_exams(self):
        exam = self._exam()
        exam.allowed_units.add(self.group)
        other = self._exam(title="W5L Other")
        other.allowed_units.add(self.group_other)

        titles = {e.title for e in self._section(chance_faculty=str(self.faculty.pk))["exams"]}
        self.assertEqual(titles, {"W5L Final"})
        titles = {
            e.title
            for e in self._section(chance_faculty=str(self.faculty.pk), chance_kafedra=str(self.chair.pk))["exams"]
        }
        self.assertEqual(titles, {"W5L Final"})
        self.assertEqual(self._section(chance_faculty=str(self.faculty.pk))["exam_count"], 1)

    def test_student_search_matches_registry_group_name(self):
        results = self._section(chance_student_q="634")["student_results"]
        self.assertEqual({u.username for u in results}, {"w5l_student", "w5l_student2"})


class TeacherResultsUnitGroupTests(_RegistryFixture):
    def setUp(self):
        self.exam = self._exam()
        self.exam.allowed_units.add(self.group)
        self.exam.allowed_groups.add(self.cohort)
        self._attempt(self.exam, self.student)
        self._attempt(self.exam, self.outsider)
        self.client = _login(self.teacher, self.org)
        self.url = reverse("exams:teacher_exam_results", args=[self.exam.slug])

    def test_available_groups_include_registry_group_and_filter_by_it(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        options = {option.id: option.name for option in response.context["available_groups"]}
        # Kohort + təyin olunmuş reyestr qrupu + cəhd edən `outsider`-in cari qrupu
        # (kohort semantikası ilə eyni: iştirakçıların üzv olduğu qruplar da siyahıdadır).
        self.assertEqual(
            options,
            {
                str(self.cohort.id): "Kohort-A",
                unit_filter_value(self.group.pk): "634 ing",
                unit_filter_value(self.group_other.pk): "701 biz",
            },
        )

        response = self.client.get(self.url, {"group": unit_filter_value(self.group.pk)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row["attempt"].user_id for row in response.context["attempts_data"]], [self.student.id])
        self.assertEqual(response.context["group_filter"], unit_filter_value(self.group.pk))

        response = self.client.get(self.url, {"group": str(self.cohort.id)})
        self.assertEqual([row["attempt"].user_id for row in response.context["attempts_data"]], [self.outsider.id])

        # Cəhd edənin cari qrupu ilə də süzülür (kohort semantikası).
        response = self.client.get(self.url, {"group": unit_filter_value(self.group_other.pk)})
        self.assertEqual([row["attempt"].user_id for row in response.context["attempts_data"]], [self.outsider.id])

        # Yad tenant / siyahıda olmayan / zibil dəyər → filtr tətbiq olunmur (əvvəlki davranış).
        for raw in ("unit:zibil", "999999", unit_filter_value(self.faculty.pk)):
            response = self.client.get(self.url, {"group": raw})
            self.assertEqual(len(response.context["attempts_data"]), 2, raw)
            self.assertEqual(response.context["group_filter"], "")

    def test_export_lists_registry_group_names(self):
        from io import BytesIO

        from openpyxl import load_workbook

        response = self.client.get(reverse("exams:export_exam_results_xlsx", args=[self.exam.slug]))
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(BytesIO(response.content)).active
        # Sütunlar: #, Qruplar, Ad Soyad, İstifadəçi adı, …
        by_username = {row[3]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(by_username["w5l_student"], "634 ing")
        self.assertEqual(by_username["w5l_outsider"], "701 biz, Kohort-A")  # kohort + cari reyestr qrupu

    def test_group_second_chance_accepts_registry_group(self):
        response = self.client.post(
            reverse("exams:grant_extra_attempt_group", args=[self.exam.slug]),
            {"group_id": unit_filter_value(self.group.pk), "extra_attempts": "1"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200, response.content[:200])
        self.assertEqual(response.json()["granted_count"], 2)
        self.assertEqual(
            set(StudentExamAttemptGrant.objects.filter(exam=self.exam).values_list("student_id", flat=True)),
            {self.student.id, self.student2.id},
        )
        response = self.client.post(
            reverse("exams:grant_extra_attempt_group", args=[self.exam.slug]),
            {"group_id": "unit:zibil"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 400)


class ExamFormExclusionUnitTests(_RegistryFixture):
    """T2: `excluded_users` reyestr qrupu tələbələri ilə də süzülür."""

    def _payload(self, **extra):
        payload = {"title": "W5L Form", "description": "", "exam_type": "test", "random_question_count": "10"}
        payload.update(extra)
        return payload

    def test_registry_group_student_can_be_excluded(self):
        form = ExamForm(
            data=self._payload(
                allowed_units=[str(self.group.pk)],
                excluded_users=[str(self.student.id), str(self.outsider.id)],
            ),
            user=self.teacher,
            organization=self.org,
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(
            set(form.cleaned_data["excluded_users"].values_list("id", flat=True)),
            {self.student.id},  # yad qrupun tələbəsi (`outsider`) istisna siyahısına düşmür
        )

    def test_cohort_only_keeps_legacy_behaviour(self):
        form = ExamForm(
            data=self._payload(
                allowed_groups=[str(self.cohort.id)],
                excluded_users=[str(self.student.id), str(self.outsider.id)],
            ),
            user=self.teacher,
            organization=self.org,
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(set(form.cleaned_data["excluded_users"].values_list("id", flat=True)), {self.outsider.id})

    def test_without_any_group_exclusions_are_dropped(self):
        form = ExamForm(
            data=self._payload(excluded_users=[str(self.student.id)]), user=self.teacher, organization=self.org
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["excluded_users"].count(), 0)

    def test_wizard_create_persists_registry_group_exclusion(self):
        client = _login(self.teacher, self.org)
        start = timezone.localtime() - timedelta(hours=1)
        response = client.post(
            reverse("exams:create_exam") + "?modal=1",
            {
                "modal": "1",
                "title": "W5L Wizard Excl",
                "description": "",
                "exam_type": "test",
                "exam_type_extended": "quiz",
                "random_question_count": "10",
                "start_datetime": start.strftime("%Y-%m-%dT%H:%M"),
                "end_datetime": (start + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M"),
                "total_duration_minutes": "60",
                "allowed_units": [str(self.group.pk)],
                "excluded_users": [str(self.student.id)],
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200, response.content[:300])
        exam = Exam.objects.get(title="W5L Wizard Excl")
        self.assertEqual(list(exam.excluded_users.values_list("id", flat=True)), [self.student.id])
        self.assertFalse(exam.can_user_see(self.student))
        self.assertTrue(exam.can_user_see(self.student2) or not exam.is_active)
