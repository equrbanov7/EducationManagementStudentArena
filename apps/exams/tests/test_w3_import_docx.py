"""W3 2026-09-14 — DOCX idxalı: şəkil + düstur (uçdan-uca) və təhlükəsizlik.

DOCX fikstürü testin içində python-docx ilə qurulur (kiçik PNG + əl ilə
qoyulmuş ``m:oMath``); binar fikstür repo-ya düşmür.
"""

from __future__ import annotations

import io
import tempfile
import zipfile
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils.translation import pgettext

from lxml import etree
from PIL import Image

from apps.accounts.models import ProfileRole
from apps.exams.models import BankQuestion, QuestionBank
from apps.exams.services import import_media
from apps.exams.services.import_media_docx import load_docx_manifest
from apps.exams.services.parsing import extract_text_from_upload
from apps.exams.services.parsing.docx_reader import docx_safety_check, normalize_image_bytes, read_docx
from apps.exams.services.visual_import_upload import prepare_question_upload
from apps.organizations.models import Membership, Organization
from core.constants import OrganizationType

User = get_user_model()
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _png(width=40, height=30, color="red") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, "PNG")
    return buffer.getvalue()


def _omath(inner: str):
    return etree.fromstring(f'<m:oMath xmlns:m="{M}">{inner}</m:oMath>')


_FRACTION = "<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>"


def build_docx(question_count=1, *, with_image=True, with_formula=True, option_image=False) -> bytes:
    """``question_count`` MCQ sualı: stem + (şəkil) + (düstur) + A–D, doğru B."""

    import docx

    document = docx.Document()
    for number in range(1, question_count + 1):
        paragraph = document.add_paragraph(f"{number}. Kəsri hesablayın {number}: ")
        if with_formula:
            paragraph._p.append(_omath(_FRACTION))
        if with_image:
            document.add_picture(io.BytesIO(_png(color=(number * 40 % 255, 20, 20))))
        document.add_paragraph(f"A) birinci {number}")
        document.add_paragraph(f"*B) ikinci {number}")
        if option_image:
            document.add_picture(io.BytesIO(_png(20, 20, "blue")))
        document.add_paragraph(f"C) üçüncü {number}")
        document.add_paragraph(f"D) dördüncü {number}")
        document.add_paragraph("")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class DocxReaderTests(TestCase):
    def test_reads_text_formula_and_image_with_anchor_markers(self):
        extract = read_docx(build_docx(1))
        lines = extract.text.splitlines()
        self.assertEqual(lines[0], "1. Kəsri hesablayın 1: \\(\\frac{a}{b}\\)")
        self.assertEqual(lines[1], "[[img:1]]")
        self.assertEqual(lines[2:], ["A) birinci 1", "*B) ikinci 1", "C) üçüncü 1", "D) dördüncü 1"])
        self.assertEqual(len(extract.images), 1)
        self.assertEqual(extract.images[0].content_type, "image/png")
        self.assertEqual(extract.formula_count, 1)
        self.assertEqual(extract.formula_fallbacks, [])
        self.assertEqual(extract.warnings, [])

    def test_option_image_marker_lands_in_option_line(self):
        extract = read_docx(build_docx(1, with_image=False, option_image=True))
        self.assertIn("[[img:1]]", extract.text)
        # Marker variant B-dən sonrakı paraqrafdadır → parser onu B-nin davamı sayır.
        lines = extract.text.splitlines()
        self.assertEqual(lines[lines.index("*B) ikinci 1") + 1], "[[img:1]]")

    def test_table_cells_become_separate_lines(self):
        import docx

        document = docx.Document()
        document.add_paragraph("1. Cədvəl sualı?")
        table = document.add_table(rows=2, cols=2)
        for cell, text in zip(table._cells, ["A) 1", "B) 2", "C) 3", "D) 4"]):
            cell.text = text
        out = io.BytesIO()
        document.save(out)
        extract = read_docx(out.getvalue())
        self.assertEqual(extract.text.splitlines(), ["1. Cədvəl sualı?", "A) 1", "B) 2", "C) 3", "D) 4"])

    def test_unsupported_formula_node_is_reported_not_dropped(self):
        import docx

        document = docx.Document()
        paragraph = document.add_paragraph("1. Sual ")
        paragraph._p.append(_omath("<m:foo><m:r><m:t>zz</m:t></m:r></m:foo>"))
        out = io.BytesIO()
        document.save(out)
        extract = read_docx(out.getvalue())
        self.assertIn("\\(\\text{zz}\\)", extract.text)
        self.assertEqual(extract.formula_fallbacks, ["\\text{zz}"])
        self.assertEqual(len(extract.warnings), 1)

    def test_image_is_downscaled_and_exif_free(self):
        big = io.BytesIO()
        image = Image.new("RGB", (3000, 1500), "green")
        exif = image.getexif()
        exif[0x010E] = "secret description"
        image.save(big, "JPEG", exif=exif.tobytes())
        data, content_type, width, height = normalize_image_bytes(big.getvalue())
        self.assertEqual((content_type, width, height), ("image/jpeg", 2000, 1000))
        with Image.open(io.BytesIO(data)) as reopened:
            self.assertFalse(reopened.getexif())

    def test_non_raster_and_oversized_images_are_skipped(self):
        self.assertIsNone(normalize_image_bytes(b"<svg xmlns='http://www.w3.org/2000/svg'/>"))
        self.assertIsNone(normalize_image_bytes(b""))
        self.assertIsNone(normalize_image_bytes(b"x" * (12 * 1024 * 1024 + 1)))

    def test_external_image_relationship_is_never_fetched(self):
        import docx

        document = docx.Document()
        document.add_paragraph("1. Xarici şəkil")
        document.add_picture(io.BytesIO(_png()))
        # Şəkil əlaqəsini xarici URL-ə çevir (TargetMode="External").
        part = document.part
        r_id = next(rid for rid, rel in part.rels.items() if "image" in rel.reltype)
        rel = part.rels[r_id]
        rel._target = "https://example.com/tracker.png"
        rel._is_external = True
        out = io.BytesIO()
        document.save(out)
        extract = read_docx(out.getvalue())
        self.assertEqual(extract.images, [])
        self.assertNotIn("[[img:", extract.text)
        self.assertEqual(len(extract.warnings), 1)


class DocxSafetyTests(TestCase):
    def _zip(self, names: dict[str, bytes]) -> bytes:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as archive:
            for name, data in names.items():
                archive.writestr(name, data)
        return out.getvalue()

    def test_rejects_non_zip_and_wrong_signature(self):
        with self.assertRaises(ValueError):
            docx_safety_check(b"%PDF-1.4 not a docx")
        with self.assertRaises(ValueError) as exc:
            docx_safety_check(b"PK\x03\x04fake-zip")
        self.assertIn("docx", str(exc.exception).lower())

    def test_rejects_macro_package_renamed_to_docx(self):
        data = self._zip(
            {
                "[Content_Types].xml": b"<Types/>",
                "word/document.xml": b"<w:document/>",
                "word/vbaProject.bin": b"\xd0\xcf\x11\xe0",
            }
        )
        with self.assertRaises(ValueError) as exc:
            docx_safety_check(data)
        self.assertEqual(str(exc.exception), pgettext("exams.service.parsing.error", "file_has_macros"))

    def test_rejects_zip_without_word_document(self):
        with self.assertRaises(ValueError):
            docx_safety_check(self._zip({"[Content_Types].xml": b"<Types/>", "other.txt": b"x"}))

    def test_zip_bomb_guard_is_applied(self):
        with patch("apps.exams.services.parsing.docx_reader.validate_zip_archive") as guard:
            from django.core.exceptions import ValidationError

            guard.side_effect = ValidationError("bomb")
            with self.assertRaises(ValueError) as exc:
                docx_safety_check(self._zip({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<w/>"}))
        self.assertIn("bomb", str(exc.exception))

    def test_extract_text_from_upload_accepts_docx_but_still_rejects_legacy_formats(self):
        uploaded = SimpleUploadedFile("q.docx", build_docx(1))
        text = extract_text_from_upload(uploaded)
        self.assertIn("\\(\\frac{a}{b}\\)", text)
        for name in ("old.doc", "rich.rtf", "macro.docm"):
            with self.assertRaises(ValueError):
                extract_text_from_upload(SimpleUploadedFile(name, b"whatever"))


@override_settings(MEDIA_URL="/media/")
class DocxBulkAddFlowTests(TestCase):
    """Bank toplu əlavə: DOCX yüklə → preview (nişanlar) → save → şəkil + LaTeX bazada."""

    def setUp(self):
        self.media_directory = tempfile.TemporaryDirectory()
        self.media_override = override_settings(MEDIA_ROOT=self.media_directory.name)
        self.media_override.enable()
        self.addCleanup(self.media_override.disable)
        self.addCleanup(self.media_directory.cleanup)

        self.teacher = User.objects.create_user(
            username="w3-docx-teacher", email="w3-docx@example.com", password="StrongPass123!"
        )
        self.organization = Organization.objects.create(
            name="W3 Docx Org",
            slug="w3-docx-org",
            org_type=OrganizationType.SCHOOL,
            owner=self.teacher,
            status="active",
            is_active=True,
        )
        profile = self.teacher.profile
        profile.organization = self.organization
        profile.organization_type = self.organization.org_type
        profile.role = ProfileRole.TEACHER
        profile.save(update_fields=["organization", "organization_type", "role", "updated_at"])
        Membership.objects.update_or_create(
            user=self.teacher,
            organization=self.organization,
            defaults={"role": self.organization.roles.get(name="teacher"), "is_primary": True, "is_active": True},
        )
        self.bank = QuestionBank.objects.create(
            name="W3 Docx Bank",
            created_by=self.teacher,
            organization=self.organization,
            default_question_type="test",
        )
        self.client.force_login(self.teacher)
        session = self.client.session
        session["active_organization"] = self.organization.slug
        session.save()
        self.url = reverse("exams:question_bank_bulk_add", kwargs={"bank_id": self.bank.id})

    def _preview(self, data: bytes, name="suallar.docx"):
        return self.client.post(
            self.url,
            {"action": "preview", "raw_text": "", "language": "az", "upload_file": SimpleUploadedFile(name, data)},
        )

    def test_preview_shows_formula_and_image_confidence_then_save_attaches_media(self):
        response = self._preview(build_docx(2, option_image=True))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertIn('name="math_token" value="', html)
        token = html.split('name="math_token" value="', 1)[1].split('"', 1)[0]
        self.assertEqual(len(token), 32)
        self.assertIn("q-badge--formula", html)
        self.assertIn("q-badge--image", html)
        self.assertIn('class="q-media-preview__img"', html)
        self.assertIn("\\(\\frac{a}{b}\\)", html)
        self.assertIn("data-ems-math", html)
        self.assertIn("vendor/katex/0.16.47/katex.min.js", html)
        # Mətn görünür qalır (PDF-dəki kimi sr-only deyil), marker mətndə yoxdur.
        self.assertNotIn("[[img:", html.split('id="questionList"', 1)[1])
        self.assertIn("filter-chip--formula", html)

        preview_url = reverse("exams:question_import_visual_preview", kwargs={"token": token, "source_index": 0})
        self.assertIn(preview_url, html)
        image_response = self.client.get(preview_url)
        self.assertEqual(image_response.status_code, 200)
        self.assertTrue(image_response.content.startswith(b"\x89PNG"))

        raw_text = response.context["raw_text"]
        self.assertIn("[[img:1]]", raw_text)
        save = self.client.post(
            self.url,
            {"action": "save", "raw_text": raw_text, "math_token": token, "language": "az", "selected_indices": "1,2"},
        )
        self.assertEqual(save.status_code, 302)

        questions = list(BankQuestion.objects.filter(bank=self.bank).order_by("id"))
        self.assertEqual(len(questions), 2)
        for index, question in enumerate(questions, start=1):
            self.assertEqual(question.text, f"Kəsri hesablayın {index}: \\(\\frac{{a}}{{b}}\\)")
            self.assertTrue(question.image, "stem şəkli bağlanmalıdır")
            self.assertFalse(question.image_replaces_text)
            self.assertTrue(question.image.storage.exists(question.image.name))
            options = {opt.label: opt for opt in question.options.all()}
            self.assertEqual(set(options), {"A", "B", "C", "D"})
            self.assertTrue(options["B"].is_correct)
            self.assertTrue(options["B"].image, "variant B-nin şəkli bağlanmalıdır")
            self.assertEqual(options["B"].text, f"ikinci {index}")
            self.assertFalse(options["A"].image)
        # Stash save-dən sonra təmizlənir.
        self.assertFalse(import_media.default_storage.exists(f"question_imports/{token}/manifest.json"))

    def test_oversized_docx_is_rejected_before_reading(self):
        big = SimpleUploadedFile("big.docx", build_docx(1))
        big.size = 45 * 1024 * 1024 + 1  # MAX_UPLOAD_BYTES (settings default 45 MB) + 1
        with patch("apps.exams.services.import_media_docx.read_docx") as reader:
            with self.assertRaises(ValueError):
                prepare_question_upload(big, owner_id=self.teacher.pk, organization_id=self.organization.pk)
        reader.assert_not_called()

    def test_docx_without_images_yields_text_only_and_no_token(self):
        text, token = prepare_question_upload(
            SimpleUploadedFile("q.docx", build_docx(1, with_image=False)),
            owner_id=self.teacher.pk,
            organization_id=self.organization.pk,
        )
        self.assertEqual(token, "")
        self.assertIn("\\(\\frac{a}{b}\\)", text)

    def test_stash_is_scoped_to_owner(self):
        _text, token = prepare_question_upload(
            SimpleUploadedFile("q.docx", build_docx(1)),
            owner_id=self.teacher.pk,
            organization_id=self.organization.pk,
        )
        manifest = load_docx_manifest(token)
        self.assertEqual(manifest["source"]["owner_id"], self.teacher.pk)
        from django.core.exceptions import PermissionDenied

        with self.assertRaises(PermissionDenied):
            import_media.bind_import_manifest(token, [], owner_id=self.teacher.pk + 1, organization_id=None)
        import_media.clear_stash(token)

    def test_unknown_marker_warns_instead_of_failing(self):
        response = self._preview(build_docx(1))
        token = response.content.decode("utf-8").split('name="math_token" value="', 1)[1].split('"', 1)[0]
        raw_text = response.context["raw_text"].replace("[[img:1]]", "[[img:9]]")
        response = self.client.post(
            self.url, {"action": "preview", "raw_text": raw_text, "math_token": token, "language": "az"}
        )
        self.assertEqual(response.status_code, 200)
        [question] = response.context["parsed"]
        self.assertIn("image_ref_unknown", [w["type"] for w in question["warnings"]])
        self.assertFalse(question.get("has_media"))
        self.assertTrue(question["meta"]["image_pending"])

    def test_save_query_budget_is_independent_of_question_count(self):
        def _run(count):
            response = self._preview(build_docx(count, option_image=True))
            token = response.content.decode("utf-8").split('name="math_token" value="', 1)[1].split('"', 1)[0]
            raw_text = response.context["raw_text"]
            BankQuestion.objects.filter(bank=self.bank).delete()
            with CaptureQueriesContext(connection) as captured:
                save = self.client.post(
                    self.url,
                    {
                        "action": "save",
                        "raw_text": raw_text,
                        "math_token": token,
                        "language": "az",
                        "selected_indices": ",".join(str(i) for i in range(1, count + 1)),
                    },
                )
            self.assertEqual(save.status_code, 302)
            self.assertEqual(BankQuestion.objects.filter(bank=self.bank, image__gt="").count(), count)
            return len(captured)

        one = _run(1)
        three = _run(3)
        self.assertEqual(one, three, f"sorğu sayı sual sayından asılıdır: 1→{one}, 3→{three}")
