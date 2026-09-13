"""2026-09-14 wave 2 (audit F-07, §27 «6 ixracda formula neytrallaşdırma») — exams XLSX ixracları.

* müəllim nəticə ixracı (``results/_export_builder.build_exam_results_xlsx_export``);
* sual bankı problem hesabatı (``question_bank/_reports._build_question_bank_report_xlsx``);
* imtahan mərkəzi statistika ixracı (``exam_center/statistics.exam_center_stats_export``).

Hər biri ``core.export_safety`` (``sheet_cell`` / ``sheet_append``) ilə yazır:
``=1+1`` mətn xanası ``'=1+1`` STRING olur (openpyxl ``data_type == 's'``), ədədlər
ədəd olaraq qalır.
"""

from __future__ import annotations

from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import RequestFactory, TestCase

from openpyxl import load_workbook

from apps.exams.models import Exam, ExamAttempt
from apps.organizations.models import Organization
from core.constants import OrganizationType

User = get_user_model()
PASSWORD = "StrongPass123!"
FORMULA = "=1+1"


def _sheet_values(content: bytes, sheet_index: int = 0):
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.worksheets[sheet_index]
    return sheet, [[cell.value for cell in row] for row in sheet.iter_rows()]


def _assert_no_formula_cells(test, sheet):
    for row in sheet.iter_rows():
        for cell in row:
            test.assertNotEqual(cell.data_type, "f", f"{cell.coordinate} formula kimi yazılıb: {cell.value!r}")


class _ExamWorld(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user("w2sec_teacher", "w2sec_t@test.az", PASSWORD)
        # Tələbənin adı formula-başlanğıclıdır — ixracda mətn kimi qalmalıdır.
        self.student = User.objects.create_user(
            "w2sec_student", "w2sec_s@test.az", PASSWORD, first_name=FORMULA, last_name=""
        )
        self.org = Organization.objects.create(
            name="W2SEC Org",
            org_type=OrganizationType.UNIVERSITY,
            owner=self.teacher,
            status="active",
            is_active=True,
        )
        self.exam = Exam.objects.create(
            # Qeyd: `/`, `?`, `*`, `[`, `]`, `:` vərəq adında qadağandır (openpyxl) — mətn
            # burada yalnız formula-başlanğıcı yoxlayır.
            title="+cmd calc A0",
            author=self.teacher,
            organization=self.org,
            exam_type="test",
            is_active=True,
        )
        self.attempt = ExamAttempt.objects.create(user=self.student, exam=self.exam, status="submitted")


class TeacherResultsExportFormulaTest(_ExamWorld):
    def test_student_name_cell_is_neutralised(self):
        from apps.exams.views.teacher.results._export_builder import build_exam_results_xlsx_export

        _filename, _content_type, content = build_exam_results_xlsx_export(self.exam, [self.attempt])
        sheet, values = _sheet_values(content)
        headers = values[0]
        name_col = headers.index("Ad Soyad")
        self.assertEqual(values[1][name_col], "'" + FORMULA)
        self.assertEqual(sheet.cell(row=2, column=name_col + 1).data_type, "s")
        # Sıra nömrəsi ədəd olaraq qalır.
        self.assertEqual(values[1][0], 1)
        self.assertEqual(sheet.cell(row=2, column=1).data_type, "n")
        _assert_no_formula_cells(self, sheet)


class QuestionBankReportFormulaTest(_ExamWorld):
    def test_question_text_and_options_are_neutralised(self):
        from apps.exams.views.teacher.question_bank._reports import _build_question_bank_report_xlsx

        parsed = [
            {
                "q_no": "1",
                "text": FORMULA,
                "options": {"A": "@SUM(A1)", "B": "-1", "C": "düz", "D": "", "E": ""},
                "correct": ["A"],
                "warnings": [{"type": "missing_option", "severity": "warning", "msg": "+cmd"}],
            }
        ]
        response = _build_question_bank_report_xlsx(
            exam=self.exam,
            raw_text="",
            parsed=parsed,
            test_level_warnings=[],
            category_counts={},
            warning_count=1,
            duplicate_count=0,
            error_count=0,
        )
        self.assertEqual(response.status_code, 200)
        summary_sheet, summary = _sheet_values(response.content, 0)
        problems_sheet, problems = _sheet_values(response.content, 1)
        # Xülasə: imtahan adı `+cmd…` → `'+cmd…`; sual sayı ədəd.
        self.assertIn(["İmtahan", "'" + self.exam.title], summary)
        self.assertIn(["Parse olunan sual sayı", 1], summary)
        flat_problems = [cell for row in problems for cell in row]
        self.assertIn("'" + FORMULA, flat_problems)
        self.assertIn("'@SUM(A1)", flat_problems)
        self.assertIn("'-1", flat_problems)
        self.assertIn("'+cmd", flat_problems)
        self.assertNotIn(FORMULA, flat_problems)
        # Sual sıra nömrəsi ədəd olaraq qalır.
        self.assertEqual(problems[1][0], 1)
        _assert_no_formula_cells(self, summary_sheet)
        _assert_no_formula_cells(self, problems_sheet)


class ExamCenterStatsExportFormulaTest(_ExamWorld):
    def test_student_and_exam_cells_are_neutralised(self):
        from apps.exams.views.exam_center import statistics as stats_module

        request = RequestFactory().get("/exams/center/stats/export/")
        request.user = self.teacher
        with mock.patch.object(stats_module, "_stats_org", return_value=self.org):
            response = stats_module.exam_center_stats_export(request)
        self.assertEqual(response.status_code, 200)
        sheet, values = _sheet_values(response.content)
        headers = values[0]
        student_col = headers.index("Tələbə")
        exam_col = headers.index("İmtahan")
        self.assertEqual(values[1][student_col], "'" + FORMULA)
        self.assertEqual(values[1][exam_col], "'" + self.exam.title)
        self.assertEqual(values[1][headers.index("İstifadəçi")], "w2sec_student")
        _assert_no_formula_cells(self, sheet)
