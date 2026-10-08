"""Sual idxalı — düzgün cavab markerləri və «A həmişə düzgündür» (müəllim rəyi S1, 2026-10-08).

Bug: müəllim 50 sual yüklədi, hamısında düzgün cavab A idi; 50-si də «Xətalı —
düzgün cavab işarəsi tapılmadı, müvəqqəti A seçildi» oldu. İndi:
  * işarəsizlik XƏTA yox, sarı xəbərdarlıqdır («Yoxlayın»);
  * «Düzgün cavab həmişə A variantıdır» seçimi mətndə ``*A)`` yazır, xəbərdarlıq çıxmır;
  * «Düzgün cavab: A», «Cavab: A) …», «+B)», «B) … +», sonda cavab açarı, DOCX-də
    qalın/vurğulu variant tanınır.
"""

import io

from django.test import SimpleTestCase
from django.urls import reverse

from apps.exams.services.parsing import parse_bulk_mcq
from apps.exams.services.parsing.answer_markers import (
    count_defaulted,
    mark_default_correct_a,
    split_answer_key,
)
from apps.exams.services.parsing.docx_reader import read_docx
from apps.exams.tests.test_question_submission import _Base


def _defaulted(question):
    return [w for w in question["warnings"] if w["type"] == "correct_defaulted"]


def _fifty_a_questions() -> str:
    """Müəllimin real faylına bənzər: 50 sual, işarə YOX, düzgün cavab həmişə A."""
    blocks = []
    for number in range(1, 51):
        blocks.append(
            f"{number}. Şəbəkə protokolları üzrə {number}-ci sual hansıdır?\n"
            f"A) Düzgün cavab {number}\n"
            f"B) Yanlış variant {number}-1\n"
            f"C) Yanlış variant {number}-2\n"
            f"D) Yanlış variant {number}-3\n"
            f"E) Yanlış variant {number}-4\n"
        )
    return "\n".join(blocks)


class MissingMarkerSeverityTests(SimpleTestCase):
    def test_missing_marker_is_a_warning_not_an_error(self):
        parsed = parse_bulk_mcq(_fifty_a_questions())
        self.assertEqual(len(parsed), 50)
        for question in parsed:
            self.assertEqual(question["correct"], ["A"])
            defaulted = _defaulted(question)
            self.assertEqual(len(defaulted), 1)
            self.assertEqual(defaulted[0]["severity"], "warning")
            self.assertFalse(any(w["severity"] == "error" for w in question["warnings"]))
        self.assertEqual(count_defaulted(parsed), 50)

    def test_default_a_option_marks_text_and_removes_warnings(self):
        text, marked = mark_default_correct_a(_fifty_a_questions())
        self.assertEqual(marked, 50)
        self.assertIn("*A) Düzgün cavab 1\n", text)
        self.assertIn("*A) Düzgün cavab 50", text)
        parsed = parse_bulk_mcq(text)
        self.assertEqual(len(parsed), 50)
        self.assertEqual(count_defaulted(parsed), 0)
        self.assertTrue(all(q["correct"] == ["A"] for q in parsed))
        # İdempotent: ikinci dəfə heç nə dəyişmir.
        self.assertEqual(mark_default_correct_a(text), (text, 0))

    def test_default_a_keeps_explicit_markers(self):
        raw = "1. Sual?\nA) bir\n*B) iki\nC) üç\nD) dörd\n\n2. Sual?\nA) x\nB) y\nC) z\nD) w\nCavab: C\n"
        text, marked = mark_default_correct_a(raw)
        self.assertEqual(marked, 0)
        self.assertEqual(text, raw)

    def test_default_a_marks_bullet_format_with_check(self):
        raw = "1. CI/CD nəyi avtomatlaşdırır?\n• Build və yerləşdirmə\n• Sənəd yazılışı\n• Dizayn\n• Dəstək\n"
        text, marked = mark_default_correct_a(raw)
        self.assertEqual(marked, 1)
        parsed = parse_bulk_mcq(text)
        self.assertEqual(parsed[0]["correct"], ["A"])
        self.assertEqual(_defaulted(parsed[0]), [])

    def test_end_question_format_is_untouched(self):
        raw = "Sual bir?\nBirinci\nİkinci\nÜçüncü\nDördüncü\nEND_QUESTION\n"
        self.assertEqual(mark_default_correct_a(raw), (raw, 0))


class AnswerLineMarkerTests(SimpleTestCase):
    def _single(self, raw):
        parsed = parse_bulk_mcq(raw)
        self.assertEqual(len(parsed), 1, parsed)
        return parsed[0]

    def test_duzgun_cavab_line(self):
        q = self._single("1. Paytaxt?\nA) Gəncə\nB) Bakı\nC) Şəki\nD) Quba\nDüzgün cavab: B\n")
        self.assertEqual(q["correct"], ["B"])
        self.assertEqual(_defaulted(q), [])

    def test_answer_line_with_label_and_option_text(self):
        q = self._single("1. İl?\nA) 1920\nB) 1918\nC) 1991\nD) 1990\nDüzgün cavab: B) 1918\n")
        self.assertEqual(q["correct"], ["B"])

    def test_answer_line_variants(self):
        for line, expected in (
            ("Cavab: (C)", ["C"]),
            ("cavab - d.", ["D"]),
            ("Doğru cavab — A", ["A"]),
            ("Correct answer: A, C", ["A", "C"]),
        ):
            with self.subTest(line=line):
                q = self._single(f"1. Sual?\nA) a1\nB) b1\nC) c1\nD) d1\n{line}\n")
                self.assertEqual(q["correct"], expected)

    def test_answer_word_inside_option_text_is_not_a_marker(self):
        q = self._single("1. Sual?\nA) Cavab: Bakı şəhəri\nB) b1\nC) c1\nD) d1\n")
        self.assertEqual(q["correct"], ["A"])
        self.assertEqual(q["options"]["A"], "Cavab: Bakı şəhəri")

    def test_plus_markers(self):
        q = self._single("1. Sual?\nA) a1\n+B) b1\nC) c1\nD) d1\n")
        self.assertEqual(q["correct"], ["B"])
        q = self._single("1. Sual?\nA) a1\nB) b1\nC) c1 +\nD) d1\n")
        self.assertEqual(q["correct"], ["C"])
        self.assertEqual(q["options"]["C"], "c1")
        q = self._single("1. Sual?\nA) 2+2\nB) 3+3\nC) c1 (+)\nD) d1\n")
        self.assertEqual(q["correct"], ["C"])
        self.assertEqual(q["options"]["A"], "2+2")


class AnswerKeySectionTests(SimpleTestCase):
    QUESTIONS = (
        "1. Birinci sual?\nA) a\nB) b\nC) c\nD) d\n\n"
        "2. İkinci sual?\nA) a\nB) b\nC) c\nD) d\n\n"
        "3. Üçüncü sual?\nA) a\nB) b\nC) c\nD) d\n\n"
    )

    def test_key_on_separate_lines(self):
        parsed = parse_bulk_mcq(self.QUESTIONS + "Düzgün cavablar:\n1. B\n2) C\n3-D\n")
        self.assertEqual([q["correct"] for q in parsed], [["B"], ["C"], ["D"]])
        self.assertEqual(count_defaulted(parsed), 0)
        self.assertEqual(len(parsed), 3)  # «1. B» yeni sual kimi oxunmur

    def test_key_inline(self):
        parsed = parse_bulk_mcq(self.QUESTIONS + "Cavablar: 1-B, 2-A 3-C\n")
        self.assertEqual([q["correct"] for q in parsed], [["B"], ["A"], ["C"]])

    def test_partial_key_leaves_others_flagged(self):
        parsed = parse_bulk_mcq(self.QUESTIONS + "Answer key: 1-C\n")
        self.assertEqual(parsed[0]["correct"], ["C"])
        self.assertEqual(len(_defaulted(parsed[1])), 1)

    def test_header_word_in_question_is_not_a_key(self):
        body, key = split_answer_key("1. Cavablar hansı formada verilir?\nA) a\nB) b\nC) c\nD) d\n")
        self.assertEqual(key, {})
        self.assertIn("Cavablar hansı", body)

    def test_default_a_respects_key(self):
        text, marked = mark_default_correct_a(self.QUESTIONS + "Cavablar:\n1-C\n")
        self.assertEqual(marked, 2)  # 2 və 3-cü suallar
        parsed = parse_bulk_mcq(text)
        self.assertEqual([q["correct"] for q in parsed], [["C"], ["A"], ["A"]])
        self.assertEqual(count_defaulted(parsed), 0)


def _docx(paragraphs) -> bytes:
    """[(mətn, {"bold"/"highlight"/"color"})] → DOCX baytları (python-docx)."""
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import RGBColor

    document = docx.Document()
    for text, style in paragraphs:
        run = document.add_paragraph().add_run(text)
        if style.get("bold"):
            run.bold = True
        if style.get("highlight"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if style.get("color"):
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DocxEmphasisTests(SimpleTestCase):
    def test_bold_option_becomes_marker(self):
        data = _docx(
            [
                ("1. Paytaxt hansıdır?", {"bold": True}),  # sual qalın — sayılmır
                ("A) Gəncə", {}),
                ("B) Bakı", {"bold": True}),
                ("C) Şəki", {}),
                ("D) Quba", {}),
                ("2. İl?", {}),
                ("A) 1918", {"highlight": True}),
                ("B) 1920", {}),
                ("C) 1991", {}),
                ("D) 1990", {}),
                ("3. Rəng?", {}),
                ("A) Qırmızı", {}),
                ("B) Yaşıl", {}),
                ("C) Göy", {"color": True}),
                ("D) Sarı", {}),
            ]
        )
        extract = read_docx(data)
        self.assertIn("*B) Bakı", extract.text)
        self.assertIn("*A) 1918", extract.text)
        self.assertIn("*C) Göy", extract.text)
        parsed = parse_bulk_mcq(extract.text)
        self.assertEqual([q["correct"] for q in parsed], [["B"], ["A"], ["C"]])
        self.assertEqual(count_defaulted(parsed), 0)

    def test_all_options_bold_is_not_a_marker(self):
        data = _docx([("1. Sual?", {})] + [(f"{label}) variant {label}", {"bold": True}) for label in "ABCD"])
        extract = read_docx(data)
        self.assertNotIn("*", extract.text)
        self.assertEqual(len(_defaulted(parse_bulk_mcq(extract.text)[0])), 1)

    def test_explicit_star_wins_over_formatting(self):
        data = _docx([("1. Sual?", {}), ("A) a", {"bold": True}), ("*B) b", {}), ("C) c", {}), ("D) d", {})])
        parsed = parse_bulk_mcq(read_docx(data).text)
        self.assertEqual(parsed[0]["correct"], ["B"])


class SubmissionWorkbenchDefaultATests(_Base):
    def _preview(self, **extra):
        client = self._client_for(self.teacher)
        data = {"action": "preview", "title": "A testi", "language": "az", "raw_text": _fifty_a_questions()}
        data.update(extra)
        response = client.post(reverse("exams:question_submission_create"), data)
        self.assertEqual(response.status_code, 200)
        return response.content.decode()

    def test_preview_offers_bulk_confirm_and_checkbox(self):
        html = self._preview()
        self.assertIn("data-wb-default-a", html)
        self.assertIn("data-wb-confirm-default-a", html)
        self.assertIn("Hamısını təsdiqlə — A düzgündür (50)", html)
        self.assertIn("workbench_default_a.js", html)
        self.assertNotIn("Xətalı sual: 50", html)

    def test_preview_with_option_marks_a_and_clears_warnings(self):
        html = self._preview(default_correct_a="1")
        self.assertIn("*A) Düzgün cavab 1", html)
        self.assertNotIn("data-wb-confirm-default-a", html)
        self.assertRegex(html, r'id="wbDefaultCorrectA"[^>]*checked')
