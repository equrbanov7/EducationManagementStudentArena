"""İmtahanın öz nəticə səhifəsinin Excel ixracı — «Dəyərləndirilmiş nəticələr» ilə EYNİ sütunlar (sahib 2026-10-27).

* başlıqlar ``results_export.HEADERS``-dir (hər iki səhifə);
* zal + kompüter: cəhdin qeydli kompüteri, yoxdursa final biletinin zalı + yeri;
* «Yoxlayan»: avtomatik test → «Avtomatik (sistem)», bal dəyişikliyi edən insan → onun adı;
* sətir quran funksiyanın sorğu sayı cəhd sayından asılı deyil.
"""

from __future__ import annotations

from io import BytesIO

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from openpyxl import load_workbook

from apps.accounts.tests.test_review_results_export_2026_10_27 import AUTO, _build_export_world
from apps.exams.models import ExamAttempt, ExamGradeEvent
from apps.exams.services.results_export import HEADERS, exam_attempt_rows
from apps.exams.services.results_export_meta import checked_by_map, location_map


class ExamResultsUnifiedExportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.world = _build_export_world("erx27")
        cls.exam = cls.world["exams"]["test"]

    def _login(self):
        self.client.force_login(self.world["teacher"])
        session = self.client.session
        session["active_organization"] = self.world["org"].slug
        session.save()

    def test_exam_page_export_uses_unified_columns(self):
        self._login()
        response = self.client.get(reverse("exams:export_exam_results_xlsx", args=[self.exam.slug]))
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(tuple(rows[0]), HEADERS)
        body = [dict(zip(rows[0], row)) for row in rows[1:]]
        self.assertEqual(len(body), 3)
        by_username = {row["İstifadəçi adı"]: row for row in body}
        attempts = self.world["test_attempts"]
        self.assertEqual(by_username[attempts[0].user.username]["Kompüter"], "PC-7")
        self.assertEqual(by_username[attempts[1].user.username]["Zal"], self.world["ticket_hall"].name)
        self.assertEqual({row["Yoxlayan"] for row in body}, {AUTO})
        self.assertEqual({row["Fənn"] for row in body}, {self.world["subject"].name})
        self.assertEqual({row["Kateqoriya"] for row in body}, {self.exam.get_exam_type_extended_display()})
        self.assertEqual({row["Cəhd №"] for row in body}, {1})

    def test_meta_maps(self):
        attempts = list(ExamAttempt.objects.filter(exam=self.exam).select_related("exam", "user").order_by("user_id"))
        locations = location_map(attempts)
        self.assertEqual(locations[attempts[0].pk], (self.world["hall"].name, "PC-7"))
        self.assertEqual(locations[attempts[1].pk], (self.world["ticket_hall"].name, "PC-3"))
        self.assertEqual(locations[attempts[2].pk], ("", ""))

        reviewer = self.world["reviewer"]
        ExamGradeEvent.objects.create(attempt=attempts[2], grader=reviewer, old_score=1, new_score=2, max_points=1)
        labels = checked_by_map(attempts)
        self.assertEqual(labels[attempts[0].pk], AUTO)
        self.assertEqual(labels[attempts[2].pk], "Rəy Çi")

    def test_row_builder_query_count_is_constant(self):
        def count(limit):
            attempts = list(
                ExamAttempt.objects.filter(exam__organization=self.world["org"])
                .select_related("exam", "user", "exam__subject", "exam__course")
                .order_by("id")[:limit]
            )
            with CaptureQueriesContext(connection) as ctx:
                rows = exam_attempt_rows(attempts)
            self.assertEqual(len(rows), len(attempts))
            return len(ctx.captured_queries)

        self.assertEqual(count(2), count(6))
