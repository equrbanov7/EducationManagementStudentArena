"""Müəllim nəticələri — qrup sütunu, ana/alt qrup süzgəci, əlavə süzgəc/sıralama (sahib 2026-10-25).

Kök səbəb: cədvəldə qrup sütunu yox idi; tələbənin qeydi ALT qrupdadır («234 K-1»), imtahan
isə ANA qrupa («234 K az») təyin olunub — ana qrup seçiləndə süzgəc 0 qaytarırdı.
"""

from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import ProfileRole
from apps.appeals.constants import APPEAL_STATUS_PENDING
from apps.appeals.models import Appeal
from apps.exams.domain.unit_scope_filters import unit_filter_value
from apps.exams.models import Exam, ExamAttempt, ExamRoom
from apps.exams.tests.test_exam_center_policy import PASSWORD, _assign_user_to_org
from apps.exams.tests.test_w5_unit_leftovers import _login
from apps.organizations.models import Organization, OrgUnit
from apps.registrar.models import Curriculum, Program, StudentAcademicRecord
from core.constants import OrganizationType, OrgUnitType
from core.rls import bypass_rls

User = get_user_model()


class ResultsGroupFixture(TestCase):
    @classmethod
    def setUpTestData(cls):
        owner = User.objects.create_user("rg_owner", "rg_owner@test.az", PASSWORD)
        cls.org = Organization.objects.create(
            name="RG Uni", org_type=OrganizationType.UNIVERSITY, owner=owner, status="active", is_active=True
        )
        cls.teacher = User.objects.create_user("rg_teacher", "rg_teacher@test.az", PASSWORD)
        _assign_user_to_org(cls.teacher, cls.org, ProfileRole.TEACHER, "teacher")
        names = {
            "s1": "Aysel",
            "s2": "Bəhruz",
            "s3": "Cəmilə",
            "s4": "Dəniz",
            "s5": "Elnur",
            "s6": "Fidan",
        }
        cls.students = {}
        for key, first in names.items():
            user = User.objects.create_user(f"rg_{key}", f"rg_{key}@test.az", PASSWORD, first_name=first)
            _assign_user_to_org(user, cls.org, ProfileRole.STUDENT, "student")
            cls.students[key] = user
        with bypass_rls():
            spec = OrgUnit.objects.create(
                organization=cls.org, name="Kompüter elmləri", slug="rg-spec", unit_type=OrgUnitType.SPECIALTY
            )
            cls.parent = OrgUnit.objects.create(
                organization=cls.org, parent=spec, name="234 K az", slug="rg-234k", unit_type=OrgUnitType.GROUP
            )
            cls.k1 = OrgUnit.objects.create(
                organization=cls.org,
                parent=spec,
                name="234 K-1",
                slug="rg-234k1",
                unit_type=OrgUnitType.GROUP,
                settings={"parent_group": str(cls.parent.pk)},
            )
            cls.k2 = OrgUnit.objects.create(  # ana qrup yalnız ad şablonu ilə tanınır
                organization=cls.org, parent=spec, name="234 K-2", slug="rg-234k2", unit_type=OrgUnitType.GROUP
            )
            cls.ing = OrgUnit.objects.create(
                organization=cls.org, parent=spec, name="234 İT ing", slug="rg-234it", unit_type=OrgUnitType.GROUP
            )
            program = Program.objects.create(organization=cls.org, code="KE", name="KE", absence_limit_percent=25)
            curriculum = Curriculum.objects.create(organization=cls.org, program=program, admission_year=2024)
            placement = {"s1": cls.k1, "s2": cls.k1, "s3": cls.k2, "s4": cls.ing, "s6": cls.k1}
            for key, unit in placement.items():
                StudentAcademicRecord.objects.create(
                    organization=cls.org,
                    student=cls.students[key],
                    program=program,
                    curriculum=curriculum,
                    group=unit,
                    admission_year=2024,
                    # s6 sonradan xaric olunub — köhnə qrupu yenə görünür.
                    is_active=key != "s6",
                    status="expelled" if key == "s6" else "enrolled",
                )
        cls.room = ExamRoom.objects.create(organization=cls.org, name="Zal 101", code="Z101")

    def setUp(self):
        now = timezone.now()
        self.exam = Exam.objects.create(
            title="RG Test",
            author=self.teacher,
            organization=self.org,
            exam_type="test",
            exam_type_extended="final",
            is_active=True,
            is_public=False,
            start_datetime=now - timedelta(days=3),
            end_datetime=now + timedelta(days=2),
        )
        self.exam.allowed_units.add(self.parent, self.ing)
        yesterday = now - timedelta(days=1)
        # (tələbə, düzgün say, gün, zal) — legacy test balı = düzgün say.
        plan = {"s1": (7, now, True), "s2": (3, now, False), "s3": (9, yesterday, False)}
        plan.update({"s4": (5, yesterday, False), "s5": (1, now, False), "s6": (6, now, False)})
        self.attempts = {}
        for key, (correct, started, in_room) in plan.items():
            attempt = ExamAttempt.objects.create(
                user=self.students[key],
                exam=self.exam,
                status="submitted",
                attempt_number=1,
                finished_at=started,
                correct_count=correct,
                wrong_count=10 - correct,
                room=self.room if in_room else None,
            )
            ExamAttempt.objects.filter(pk=attempt.pk).update(started_at=started)
            self.attempts[key] = attempt
        self.client = _login(self.teacher, self.org)
        self.url = reverse("exams:teacher_exam_results", args=[self.exam.slug])

    def _users(self, response):
        return {row["attempt"].user.username.removeprefix("rg_") for row in response.context["attempts_data"]}

    def _get(self, params=None):
        response = self.client.get(self.url, {"per_page": 100, **(params or {})})
        self.assertEqual(response.status_code, 200)
        return response


class GroupColumnAndFilterTests(ResultsGroupFixture):
    def test_every_row_has_group_with_parent(self):
        response = self._get()
        rows = {
            row["attempt"].user.username.removeprefix("rg_"): row["group"] for row in response.context["attempts_data"]
        }
        self.assertEqual(rows["s1"].name, "234 K-1")
        self.assertEqual(rows["s1"].parent_name, "234 K az")
        self.assertEqual(rows["s3"].parent_name, "234 K az")  # ad şablonu ilə
        self.assertEqual(rows["s4"].name, "234 İT ing")
        self.assertEqual(rows["s4"].parent_name, "")
        self.assertFalse(rows["s6"].is_active_record)
        self.assertIsNone(rows["s5"])
        self.assertContains(response, 'class="etr-group__name">234 K-1<')

    def test_options_have_counts_and_parent_rolls_up_subgroups(self):
        options = {opt.name: opt for opt in self._get().context["available_groups"]}
        self.assertEqual(options["234 K az"].count, 4)  # s1, s2, s6 (K-1) + s3 (K-2)
        self.assertEqual(set(options["234 K az"].subgroup_names), {"234 K-1", "234 K-2"})
        self.assertEqual(options["234 K-1"].count, 3)
        self.assertEqual(options["234 İT ing"].count, 1)
        self.assertEqual(options["Qrupsuz"].count, 1)

    def test_parent_group_filter_includes_subgroup_students(self):
        response = self._get({"group": unit_filter_value(self.parent.pk)})
        self.assertEqual(self._users(response), {"s1", "s2", "s3", "s6"})
        self.assertEqual(self._users(self._get({"group": unit_filter_value(self.k1.pk)})), {"s1", "s2", "s6"})

    def test_multi_select_and_no_group(self):
        response = self.client.get(
            self.url, {"group": [unit_filter_value(self.k2.pk), unit_filter_value(self.ing.pk)], "per_page": 100}
        )
        self.assertEqual(self._users(response), {"s3", "s4"})
        self.assertEqual(len(response.context["group_filter_values"]), 2)
        self.assertEqual(self._users(self._get({"group": "none"})), {"s5"})
        # Yad / zibil dəyər atılır.
        self.assertEqual(len(self._get({"group": "unit:zibil"}).context["attempts_data"]), 6)

    def test_grant_select_skips_no_group_option(self):
        grant = {opt.id for opt in self._get().context["grant_groups"]}
        self.assertNotIn("none", grant)
        self.assertIn(unit_filter_value(self.parent.pk), grant)


class ExtraFiltersAndSortTests(ResultsGroupFixture):
    def test_day_room_and_score_filters(self):
        response = self._get()
        self.assertEqual(len(response.context["day_options"]), 2)
        day = (timezone.localtime(timezone.now()) - timedelta(days=1)).strftime("%d.%m.%Y")
        self.assertEqual(self._users(self._get({"day": day})), {"s3", "s4"})
        self.assertEqual(self._users(self._get({"room": str(self.room.pk)})), {"s1"})
        self.assertEqual(self._users(self._get({"room": "none"})), {"s2", "s3", "s4", "s5", "s6"})
        self.assertEqual(self._users(self._get({"score_min": "5", "score_max": "7"})), {"s1", "s4", "s6"})

    def test_sort_by_score_and_group(self):
        response = self._get({"sort_by": "score", "sort_dir": "desc"})
        order = [row["attempt"].user.username.removeprefix("rg_") for row in response.context["attempts_data"]]
        self.assertEqual(order, ["s3", "s1", "s6", "s4", "s2", "s5"])
        response = self._get({"sort_by": "group", "sort_dir": "asc"})
        groups = [row["group"].name if row["group"] else "" for row in response.context["attempts_data"]]
        self.assertEqual(groups[-1], "")  # qrupsuz sonda (NULL)
        # Qrup adlarının nisbi sırası DB collation-undan asılıdır («İ» vs «K»): yalnız eyni qrupun ardıcıl
        # gəldiyini və qrupların say tərkibini yoxla.
        named = [name for name in groups if name]
        runs = [name for index, name in enumerate(named) if index == 0 or named[index - 1] != name]
        self.assertEqual(len(runs), len(set(runs)), named)
        self.assertEqual(named.count("234 K-1"), 3)

    def test_appeal_filter_and_badge(self):
        Appeal.objects.create(
            attempt=self.attempts["s2"],
            exam=self.exam,
            student=self.students["s2"],
            organization=self.org,
            status=APPEAL_STATUS_PENDING,
        )
        response = self._get({"appeal": "open"})
        self.assertEqual(self._users(response), {"s2"})
        self.assertTrue(response.context["attempts_data"][0]["appeal"])
        self.assertEqual(len(self._get({"appeal": "none"}).context["attempts_data"]), 5)

    def test_search_by_exact_fin(self):
        profile = self.students["s4"].profile
        profile.fin = "7AB12CD"
        profile.save(update_fields=["fin"])
        self.assertEqual(self._users(self._get({"q": "7ab12cd"})), {"s4"})

    def test_summary_shows_filtered_of_total_and_url_keeps_state(self):
        response = self._get({"group": unit_filter_value(self.k1.pk), "sort_by": "score", "sort_dir": "asc"})
        self.assertEqual(response.context["review_stats"]["total"], 3)
        self.assertEqual(response.context["overall_total"], 6)
        self.assertIn("group=unit", response.context["pagination_query"])
        self.assertIn("sort_by=score", response.context["export_query"])
        self.assertIsNone(self._get().context["overall_total"])

    def test_query_count_independent_of_row_count(self):
        self._get()  # sessiya yazısı və s. isinmə
        with CaptureQueriesContext(connection) as small:
            self._get({"group": unit_filter_value(self.ing.pk)})  # 1 sətir
        with CaptureQueriesContext(connection) as large:
            self._get({"group": unit_filter_value(self.parent.pk)})  # 4 sətir
        self.assertEqual(len(small.captured_queries), len(large.captured_queries))


class ExportFollowsFiltersTests(ResultsGroupFixture):
    def test_export_has_group_parent_room_appeal_columns_and_respects_filter_sort(self):
        from openpyxl import load_workbook

        url = reverse("exams:export_exam_results_xlsx", args=[self.exam.slug])
        response = self.client.get(
            url, {"group": unit_filter_value(self.parent.pk), "sort_by": "score", "sort_dir": "desc"}
        )
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        header = rows[0]
        self.assertEqual(header[1:4], ("Qruplar", "Ana qrup", "Ad Soyad"))
        self.assertEqual(header[-2:], ("Zal", "Apellyasiya"))
        body = rows[1:]
        self.assertEqual([row[4] for row in body], ["rg_s3", "rg_s1", "rg_s6", "rg_s2"])
        self.assertEqual(body[1][1:3], ("234 K-1", "234 K az"))
        self.assertEqual(body[1][-2], "Zal 101")
