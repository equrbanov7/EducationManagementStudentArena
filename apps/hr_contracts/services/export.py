"""Excel ixracı — HR siyahısının cari filtr dəsti (bir vərəq, başlıq sətri qalın + sabitlənmiş).

Hücrələr ``core.export_safety.sheet_append`` ilə yazılır (formula injection neytrallaşdırılır).
"""

from __future__ import annotations

from io import BytesIO

from django.utils.translation import pgettext

from core.export_safety import sheet_append

from ..constants import CTX, ContractState, ContractType
from . import queries
from .notify import person_name


def _headers() -> list[str]:
    return [
        pgettext(CTX, "İşçi"),
        pgettext(CTX, "İstifadəçi adı"),
        pgettext(CTX, "Müqavilə №"),
        pgettext(CTX, "Növ"),
        pgettext(CTX, "Vəzifə"),
        pgettext(CTX, "Bölmə"),
        pgettext(CTX, "Başlama"),
        pgettext(CTX, "Bitmə"),
        pgettext(CTX, "Qalan gün"),
        pgettext(CTX, "Sınaq müddəti"),
        pgettext(CTX, "Həftəlik saat"),
        pgettext(CTX, "Ştat"),
        pgettext(CTX, "Vəziyyət"),
        pgettext(CTX, "Göndərilib"),
        pgettext(CTX, "Təsdiqlənib"),
    ]


def _when(value) -> str:
    if not value:
        return ""
    from django.utils import timezone

    return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")


def xlsx_bytes(contracts, *, on, window) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = pgettext(CTX, "Müqavilələr")[:31]
    headers = _headers()
    sheet_append(sheet, headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for contract in contracts:
        state = queries.state_of(contract, on=on, window=window)
        sheet_append(
            sheet,
            [
                person_name(contract.employee),
                contract.employee.username,
                contract.number,
                str(ContractType(contract.contract_type).label),
                contract.position,
                contract.unit.name if contract.unit_id else "",
                contract.start_date,
                contract.end_date,
                queries.days_left(contract, on=on),
                contract.probation_end,
                float(contract.weekly_hours) if contract.weekly_hours is not None else None,
                float(contract.staff_rate) if contract.staff_rate is not None else None,
                str(ContractState(state).label),
                _when(contract.sent_at),
                _when(contract.acknowledged_at),
            ],
        )
    for index, _header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = 18
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if hasattr(cell.value, "isoformat") and not isinstance(cell.value, str):
                cell.number_format = "DD.MM.YYYY"
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["xlsx_bytes"]
