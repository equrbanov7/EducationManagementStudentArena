"""2026-09-14 wave 2 (audit F-07, §27 «6 ixracda formula neytrallaşdırma») — dərs yükü XLSX.

``workload.views.teacher_api.my_export`` ``core.export_safety.sheet_append`` ilə yazır:
fənn/qrup adı ``=1+1`` → ``'=1+1`` (string), saat sütunu və CƏMİ ədəd olaraq qalır.
"""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import RequestFactory, TestCase

from openpyxl import load_workbook

from apps.workload.views import teacher_api

User = get_user_model()


class WorkloadExportFormulaTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("w2sec_workload", "w2sec_workload@test.az", "StrongPass123!")

    def test_subject_and_group_cells_are_neutralised(self):
        rows = [
            {
                "academic_year": "2025/2026",
                "season_label": "Payız",
                "subject": "=1+1",
                "groups": "@SUM(A1), -QR-1",
                "activity_label": "+Mühazirə",
                "hours": 30,
                "education_form": "əyani",
                "degree_level": "bakalavr",
                "is_hourly_paid": False,
            }
        ]
        request = RequestFactory().get("/workload/mene/ixrac/")
        request.user = self.user
        request.organization = SimpleNamespace(pk=1)
        actor = SimpleNamespace(has=lambda perm: True)
        with (
            mock.patch.object(teacher_api, "actor_for", return_value=actor),
            mock.patch.object(teacher_api, "teacher_workload_rows", return_value=rows),
        ):
            response = teacher_api.my_export(request)
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(values[0][2], "Fənn")
        self.assertEqual(values[1][2], "'=1+1")
        self.assertEqual(values[1][3], "'@SUM(A1), -QR-1")
        self.assertEqual(values[1][4], "'+Mühazirə")
        self.assertEqual(values[1][5], 30)
        self.assertEqual(sheet.cell(row=2, column=6).data_type, "n")
        # CƏMİ sətri: etiket mətn, cəm ədəd.
        self.assertEqual(values[2][4], "CƏMİ")
        self.assertEqual(values[2][5], 30)
        for row in sheet.iter_rows():
            for cell in row:
                self.assertNotEqual(cell.data_type, "f", cell.coordinate)
