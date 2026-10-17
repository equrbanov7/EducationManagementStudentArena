"""Sahib 2026-10-17: magistratura cədvəli → «dərs yükündə olmayanları da əlavə et»
(``assign_task_rows_from_timetable --create-missing-rows``).

Magistratura hüceyrəsi «Fənn (N saat) [dərəcə] Ad Soyad»: müəllim TAM ADLA (hər iki sıra) tapılır; heç bir
sətrə uyğun gəlməyən dərs üçün müəllimin kafedrasında tapşırıq + sətir + kataloq fənni yaranır, müəllim
mühazirəyə yazılır, açılış + qrupun tələbələri qeydiyyatdan keçir. Dry-run heç nə yazmır, təkrar idempotentdir,
mövcud sətir varsa yenisi yaradılmır, təsdiqlənmiş sənədə sətir düzəliş axını ilə əlavə olunur."""

from __future__ import annotations

import csv
import datetime
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone

import docx

from apps.audit.models import AuditLog
from apps.notifications.models import InAppNotification
from apps.organizations.models import AcademicPeriod, Membership, OrgUnit
from apps.registrar.models import (
    AcademicStatus,
    CourseOffering,
    Curriculum,
    Enrollment,
    Program,
    StudentAcademicRecord,
    Subject,
)
from apps.workload.models import TeacherAssignment, TeachingTask, TeachingTaskRow, WorkloadAmendment
from apps.workload.services.scoping import WorkloadDenied
from core.constants import OrgUnitType, RoleScopeType

from .factories import TEACHER_PERMS, activate_member, make_org

User = get_user_model()
HEADER = ["QRUPLAR", "QRUPLAR", "", "213/6BM ing (1)", "", "631/5 K (8)"]
HOURS = ["Saatlar", "Saatlar", "aud", "Fənlər", "aud", "Fənlər"]
SLOTS = [("Bazar ertəsi", "18:40-20:00"), ("Bazar ertəsi", "20:10-21:30"), ("Çərşənbə", "18:40-20:00")]


class TimetableMissingRowsTests(TestCase):
    def setUp(self):
        today = timezone.localdate()
        self.org = make_org("qku-mag")
        faculty = OrgUnit.objects.create(organization=self.org, name="SİF", unit_type=OrgUnitType.FACULTY)
        self.ir_chair = OrgUnit.objects.create(
            organization=self.org, name="Beynəlxalq münasibətlər", unit_type=OrgUnitType.CHAIR, parent=faculty
        )
        self.it_chair = OrgUnit.objects.create(
            organization=self.org, name="İnformasiya texnologiyaları", unit_type=OrgUnitType.CHAIR, parent=faculty
        )
        specialty = OrgUnit.objects.create(
            organization=self.org, name="Beynəlxalq münasibətlər (mag)", unit_type=OrgUnitType.SPECIALTY,
            parent=faculty,
        )  # fmt: skip
        self.bm = OrgUnit.objects.create(
            organization=self.org, name="213/6 BM ing", unit_type=OrgUnitType.GROUP, parent=specialty
        )
        self.k = OrgUnit.objects.create(organization=self.org, name="631/5K", unit_type=OrgUnitType.GROUP)
        self.period = AcademicPeriod.objects.create(
            organization=self.org, name="Payız", period_type="semester", academic_year="2026/2027",
            start_date=today - datetime.timedelta(days=20), end_date=today + datetime.timedelta(days=100),
            is_current=True,
        )  # fmt: skip
        self.admin = User.objects.create_superuser("superadmin", "sa@x.test", "pw")
        activate_member(self.org, self.admin, "rector", permissions=["*"], level=100)
        self.users = {}
        for username, first, last, chair in (
            ("rail.s", "Rail", "Səfiyev", self.ir_chair),
            ("nail.q", "Nail", "Qasımzadə", self.ir_chair),
            ("asmer.n", "Əsmər", "Nəbiyeva", self.it_chair),
        ):
            self._teacher(username, first, last, chair)
        self.student = User.objects.create_user("mag.student", "ms@x.test", "pw")
        activate_member(self.org, self.student, "student", scope_unit=self.bm)
        program = Program.objects.create(organization=self.org, code="MAG-BM", name="Beynəlxalq münasibətlər")
        curriculum = Curriculum.objects.create(organization=self.org, program=program, admission_year=2026)
        StudentAcademicRecord.objects.create(
            organization=self.org, student=self.student, program=program, curriculum=curriculum, group=self.bm,
            admission_year=2026, status=AcademicStatus.ENROLLED,
        )  # fmt: skip
        self.tmp = TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def _teacher(self, username, first, last, chair):
        user = User.objects.create_user(username, f"{username}@x.test", "pw", first_name=first, last_name=last)
        activate_member(
            self.org, user, "teacher", permissions=TEACHER_PERMS, scope_unit=chair, scope_type=RoleScopeType.COURSE
        )
        self.users[username] = user
        return user

    def _docx(self, rows):
        """``rows``: [(BM hüceyrəsi, K hüceyrəsi), …] — gün/saat ``SLOTS``-dan."""
        document = docx.Document()
        table = document.add_table(rows=2 + len(rows), cols=6)
        for col, text in enumerate(HEADER):
            table.cell(0, col).text = text
        for col, text in enumerate(HOURS):
            table.cell(1, col).text = text
        for offset, ((day, time), (bm, k)) in enumerate(zip(SLOTS, rows), start=2):
            for col, value in enumerate((day, time, "204", bm, "11", k)):
                table.cell(offset, col).text = value
        for row, text in ((0, "QRUPLAR"), (1, "Saatlar")):
            table.cell(row, 0).merge(table.cell(row, 1)).text = text
        path = Path(self.tmp.name) / "magistratura.docx"
        document.save(str(path))
        return path

    def _run(self, path, *extra):
        out = StringIO()
        call_command(
            "assign_task_rows_from_timetable", "--file", str(path), "--org", self.org.slug,
            "--period", "Payız 2026/2027", "--year", "2026/2027", "--actor", "superadmin",
            "--report", self.tmp.name, *extra, stdout=out,
        )  # fmt: skip
        return out.getvalue()

    def _report(self):
        with (Path(self.tmp.name) / "timetable_task_assign.csv").open(encoding="utf-8") as handle:
            return list(csv.DictReader(handle))

    def _notified(self, username):
        return InAppNotification.objects.filter(
            recipient=self.users[username], metadata__event="workload_assigned"
        ).count()

    def _diplomatiya_draft(self, **hours):
        subject = Subject.objects.create(organization=self.org, code="QKU-9", name="Diplomatiya", ects=5)
        task = TeachingTask.objects.create(
            organization=self.org, chair=self.ir_chair, academic_year="2026/2027", created_by=self.admin
        )
        row = TeachingTaskRow.objects.create(
            organization=self.org, task=task, season="fall", period=self.period, subject=subject,
            subject_text="Diplomatiya", total_hours=sum(hours.values()), **hours,
        )  # fmt: skip
        row.groups.set([self.bm])
        return task, row

    def _counts(self):
        return (
            TeachingTask.objects.filter(organization=self.org).count(),
            TeachingTaskRow.objects.filter(organization=self.org).count(),
            Subject.objects.filter(organization=self.org).count(),
            CourseOffering.objects.filter(organization=self.org).count(),
        )

    def test_dry_run_writes_nothing_then_apply_creates_task_row_subject_offering_and_is_idempotent(self):
        # Eyni fənn + qrup iki hüceyrədə (45 / 15 saat) → BİR sətir, saat maksimum; tam ad, dərəcə nişanı.
        path = self._docx(
            [
                ("Tədqiqat metodları (45 saat) t.ü.f.d. Rail Səfiyev", ""),
                ("Tədqiqat metodları (15 saat) Rail Səfiyev", ""),
            ]
        )
        before = self._counts()
        out = self._run(path, "--create-missing-rows")
        self.assertIn("DRY-RUN", out)
        self.assertNotIn("Səfiyev", out)
        for line in ("row_created=1", "subject_created=1", "task_created=1", "row_hours_differ=1"):
            self.assertIn(line, out)
        self.assertEqual(self._counts(), before)
        self.assertEqual({row["status"] for row in self._report()}, {"row_planned"})
        self.assertEqual({row["chair"] for row in self._report()}, {"Beynəlxalq münasibətlər"})

        out = self._run(path, "--create-missing-rows", "--apply")
        self.assertIn("row_created=1", out)
        self.assertIn("result_assigned=2", out)
        task = TeachingTask.objects.get(organization=self.org, chair=self.ir_chair)
        self.assertEqual(task.status, "distributed")
        row = task.rows.get()
        self.assertEqual((row.lecture_total, row.seminar_total, row.lab_total), (45, 0, 0))
        self.assertEqual(row.note, "saat bölgüsü cədvəldən — kafedra dəqiqləşdirsin")
        self.assertEqual(row.degree_level, "master")
        self.assertEqual(list(row.groups.all()), [self.bm])
        self.assertEqual(row.subject.name, "Tədqiqat metodları")
        self.assertEqual(row.subject.chair_unit, self.ir_chair)
        assignment = TeacherAssignment.objects.get(row=row, activity="lecture")
        self.assertEqual((assignment.teacher, assignment.hours), (self.users["rail.s"], 45))
        offering = CourseOffering.objects.get(organization=self.org, subject=row.subject, group=self.bm)
        self.assertEqual(offering.instructor, self.users["rail.s"])
        self.assertTrue(Enrollment.objects.filter(offering=offering, student=self.student).exists())

        counts = self._counts()
        out = self._run(path, "--create-missing-rows", "--apply")
        self.assertNotIn("row_created", out)
        self.assertIn("cell_already_assigned=2", out)
        self.assertEqual(self._counts(), counts)

    def test_student_without_active_membership_does_not_abort_the_import(self):
        # Prod 2026-10-17: qrupda aktiv tələbə üzvlüyü olmayan qeyd registrar_guard_active_member-də
        # dayanırdı və bütün tətbiq geri qayıdırdı — indi o tələbə sayılıb ötürülür.
        ghost = User.objects.create_user("mag.ghost", "mg@x.test", "pw")
        activate_member(self.org, ghost, "student", scope_unit=self.bm)
        record = StudentAcademicRecord.objects.filter(student=self.student).first()
        StudentAcademicRecord.objects.create(
            organization=self.org, student=ghost, program=record.program, curriculum=record.curriculum,
            group=self.bm, admission_year=2026, status=AcademicStatus.ENROLLED,
        )  # fmt: skip
        Membership.objects.filter(organization=self.org, user=ghost).update(is_active=False)
        path = self._docx([("Tədqiqat metodları (45 saat) Rail Səfiyev", "")])
        out = self._run(path, "--create-missing-rows", "--apply")
        self.assertIn("row_created=1", out)
        self.assertIn("enrollment_not_member=1", out)
        offering = CourseOffering.objects.get(organization=self.org, group=self.bm, subject__name="Tədqiqat metodları")
        self.assertTrue(Enrollment.objects.filter(offering=offering, student=self.student).exists())
        self.assertFalse(Enrollment.objects.filter(offering=offering, student=ghost).exists())

    def test_distributed_task_gets_new_row_through_amendment(self):
        self._run(self._docx([("Tədqiqat metodları (45 saat) Rail Səfiyev", "")]), "--create-missing-rows", "--apply")
        task = TeachingTask.objects.get(organization=self.org, chair=self.ir_chair)
        self.assertEqual(task.status, "distributed")
        # Soyad əvvəl yazılıb («Qasımzadə Nail»); eyni kafedranın təsdiqlənmiş sənədinə yeni sətir.
        path = self._docx([("İnvestisiya portfelinin idarə edilməsi (45 saat) Qasımzadə Nail", "")])
        out = self._run(path, "--create-missing-rows", "--apply")
        self.assertIn("row_created=1", out)
        self.assertNotIn("task_created", out)
        task.refresh_from_db()
        self.assertEqual(task.status, "distributed")
        row = task.rows.get(subject__name="İnvestisiya portfelinin idarə edilməsi")
        self.assertTrue(WorkloadAmendment.objects.filter(task=task, target_id=row.pk).exists())
        self.assertEqual(
            set(TeacherAssignment.objects.filter(row=row).values_list("teacher__username", flat=True)), {"nail.q"}
        )
        # Əvvəl bölünmüş sənədin yenidən təsdiqi — kafedraya təkrar «Dərs yükü təyin edildi» getmir, yalnız yeni
        # təyin olunan müəllimə bir dəfə (sahib 2026-10-17); dəyişikliksiz təkrar iş heç nə göndərmir.
        self.assertEqual(self._notified("rail.s"), 1)
        self.assertEqual(self._notified("nail.q"), 1)
        self._run(path, "--create-missing-rows", "--apply")
        self.assertEqual((self._notified("rail.s"), self._notified("nail.q")), (1, 1))

    def test_existing_row_is_used_and_subject_owner_chair_wins(self):
        subject = Subject.objects.create(
            organization=self.org, code="QKU-7", name="Müasir proqramlaşdırma dilləri", ects=5,
            chair_unit=self.it_chair,
        )  # fmt: skip
        path = self._docx([("", "Müasir proqramlaşdırma dilləri (45 saat) Əsmər Nəbiyeva")])
        self._run(path, "--create-missing-rows", "--apply")
        task = TeachingTask.objects.get(organization=self.org, chair=self.it_chair)
        row = task.rows.get()
        self.assertEqual((row.subject, list(row.groups.all())), (subject, [self.k]))  # «631/5 K» ↔ «631/5K»
        self.assertEqual(row.degree_level, "master")
        # Sətir indi var → ikinci cədvəl hüceyrəsi YENİ sətir yaratmır (adi axın).
        path = self._docx([("", "Müasir proqramlaşdırma dilləri (60 saat) Əsmər Nəbiyeva")])
        out = self._run(path, "--create-missing-rows", "--apply")
        self.assertNotIn("row_created", out)
        self.assertEqual(TeachingTaskRow.objects.filter(organization=self.org).count(), 1)
        self.assertEqual(self._report()[0]["status"], "already_assigned")

    def test_without_flag_nothing_is_created(self):
        out = self._run(self._docx([("Tədqiqat metodları (45 saat) Rail Səfiyev", "")]), "--apply")
        self.assertIn("cell_group_no_task_rows=1", out)
        self.assertFalse(TeachingTaskRow.objects.filter(organization=self.org).exists())

    def test_ambiguous_full_name_and_missing_teacher_are_reported_not_written(self):
        self._teacher("rail.s2", "Rail", "Səfiyev", self.it_chair)
        path = self._docx([("Tədqiqat metodları (45 saat) Səfiyev Rail", "Sitogenetika (45 saat)")])
        out = self._run(path, "--create-missing-rows", "--apply")
        self.assertIn("cell_teacher_ambiguous=1", out)
        self.assertIn("cell_teacher_missing=1", out)
        self.assertFalse(TeachingTaskRow.objects.filter(organization=self.org).exists())
        self.assertFalse(TeachingTask.objects.filter(organization=self.org).exists())

    def test_existing_task_row_gets_full_name_teacher_on_lecture_by_default(self):
        subject = Subject.objects.create(organization=self.org, code="QKU-9", name="Diplomatiya", ects=5)
        task = TeachingTask.objects.create(
            organization=self.org, chair=self.ir_chair, academic_year="2026/2027", created_by=self.admin
        )
        row = TeachingTaskRow.objects.create(
            organization=self.org, task=task, season="fall", period=self.period, subject=subject,
            subject_text="Diplomatiya", lecture_total=30, seminar_total=15, total_hours=45,
        )  # fmt: skip
        row.groups.set([self.bm])
        out = self._run(self._docx([("Diplomatiya (45 saat) dos. Səfiyev Rail", "")]), "--create-missing-rows")
        self.assertIn("result_would_assign=1", out)
        self.assertNotIn("row_created", out)
        report = self._report()[0]
        self.assertEqual(
            (report["status"], report["activity"], report["detail"]), ("planned", "lecture", "kind_default_lecture")
        )
        self.assertEqual(report["username"], "rail.s")

    def test_draft_task_with_new_and_existing_rows_is_approved_and_confirmed_once(self):
        # 2026-10-17 rəyi (1): eyni kafedranın tapşırığında həm yeni sətir, həm mövcud sətrin bölgüsü —
        # əvvəl köhnəlmiş «draft» statusu ilə İKİ dəfə təsdiqlənir, bildirişlər iki dəfə gedirdi.
        task, row = self._diplomatiya_draft(lecture_total=30)
        path = self._docx(
            [("Diplomatiya (30 saat) Rail Səfiyev", ""), ("Tədqiqat metodları (45 saat) Rail Səfiyev", "")]
        )
        out = self._run(path, "--create-missing-rows", "--apply")
        for line in ("row_created=1", "task_confirmed=1", "result_assigned=2"):
            self.assertIn(line, out)
        self.assertNotIn("reassigned", out)
        task.refresh_from_db()
        self.assertEqual(task.status, "distributed")
        reasons = list(AuditLog.objects.filter(resource_id=str(task.pk)).values_list("reason", flat=True))
        self.assertEqual(reasons.count("workload.task_approved_by_official_workbook_import"), 1)
        self.assertEqual(reasons.count("workload.distribution_confirmed"), 1)
        self.assertEqual(self._notified("rail.s"), 1)
        self.assertEqual(self._teachers_of(row), {"rail.s"})
        self.assertEqual(self._teachers_of(task.rows.get(subject__name="Tədqiqat metodları")), {"rail.s"})

    def _teachers_of(self, row):
        return set(TeacherAssignment.objects.filter(row=row).values_list("teacher__username", flat=True))

    def test_two_spellings_of_one_new_subject_make_one_row_in_dry_run_and_apply(self):
        # Rəy (2): «Tədqiqat metodları» / «Tədqiqatın metodları» (≥ 0.9 oxşar) eyni fənnə düşür → BİR sətir;
        # əvvəl dry-run row_created=2 deyirdi, tətbiq isə ``duplicate_row`` ilə bütün kafedranı geri qaytarırdı.
        path = self._docx(
            [
                ("Tədqiqat metodları (45 saat) Rail Səfiyev", ""),
                ("Tədqiqatın metodları (30 saat) Rail Səfiyev", ""),
            ]
        )
        for extra in ((), ("--apply",)):
            out = self._run(path, "--create-missing-rows", *extra)
            for line in ("row_created=1", "subject_created=1", "row_hours_differ=1"):
                self.assertIn(line, out)
            self.assertNotIn("task_blocked", out)
        self.assertEqual({r["status"] for r in self._report()}, {"row_created"})
        self.assertEqual({r["detail"] for r in self._report()}, {"hours_max:45/30"})
        row = TeachingTaskRow.objects.get(organization=self.org)
        self.assertEqual(
            (row.subject.name, row.lecture_total, list(row.groups.all())), ("Tədqiqat metodları", 45, [self.bm])
        )
        self.assertEqual(Subject.objects.filter(organization=self.org).count(), 1)

    def test_rolled_back_chair_restores_every_counter(self):
        # Rəy (3): kafedranın savepoint-i geri qayıdanda yalnız ``subject_created`` bərpa olunurdu —
        # ``task_created`` / ``row_created`` / ``assigned`` … log-da qalırdı (prod log-u yalnız saylardır).
        self._diplomatiya_draft(lecture_total=30)
        path = self._docx(
            [("Diplomatiya (30 saat) Rail Səfiyev", ""), ("Tədqiqat metodları (45 saat) Rail Səfiyev", "")]
        )
        denied = WorkloadDenied("workload.distribution_incomplete", "test")
        with mock.patch("apps.workload.services.task_workbook_import.confirm_distribution", side_effect=denied):
            out = self._run(path, "--create-missing-rows", "--apply")
        counters = dict(line.strip().split("=", 1) for line in out.splitlines() if line.startswith("  "))
        self.assertEqual(
            counters,
            {"cell_planned": "1", "cell_task_blocked": "1", "chair_rolled_back": "1", "task_blocked": "1",
             "result_task_blocked": "1"},
        )  # fmt: skip
        self.assertEqual(TeachingTaskRow.objects.filter(organization=self.org).count(), 1)
        self.assertEqual(Subject.objects.filter(organization=self.org).count(), 1)
        self.assertEqual(TeachingTask.objects.get(organization=self.org).status, "draft")
        self.assertFalse(TeacherAssignment.objects.filter(row__task__organization=self.org).exists())
        self.assertEqual({r["status"] for r in self._report()}, {"planned", "task_blocked"})
