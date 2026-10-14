"""Sahib 2026-10-12: «Forma 1» dərs yükü (Word) → kafedranın mövcud tapşırıq sətirlərinə
müəllim bölgüsü: sütunlar başlıqdan tapılır, tam adla müəllim, saathesabı bölməsi,
«+» / --choice seçimi, qrup açarları («235EKO AZ» = «235 EKO»), hesab yaradılmır,
fərqli müəllim yalnız --reassign ilə (açılışın müəllimi də dəyişir)."""

from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase

from docx import Document
from openpyxl import Workbook

from apps.organizations.models import Membership, OrgUnit
from apps.registrar.models import CourseOffering, Subject
from apps.workload.models import TeacherAssignment, TeachingTask
from apps.workload.services.task_form1_parsing import _is_total, group_key, group_tokens, hours_value, name_words
from core.constants import OrgUnitType, RoleScopeType

from .factories import TEACHER_PERMS, activate_member, make_org

User = get_user_model()
_HEAD = ["Müəllimin S.A.A., elmi dərəcəsi", "FAKÜLTƏ, QRUP, KURS", "FƏNNİN ADI", "Tələbələrin sayı"]


class Form1HelperTests(SimpleTestCase):
    def test_helpers(self):
        self.assertEqual(hours_value("15*2"), 30)
        self.assertEqual(hours_value("9+8+10"), 27)
        self.assertEqual(hours_value(""), 0)
        self.assertEqual(group_key("234 BİO AZ"), group_key("234 BİO"))
        self.assertEqual(group_key("235BİO İNG"), group_key("235 BİO ing"))
        self.assertEqual(group_tokens("233BİO SBİO İNG"), ["233 BİO İNG", "233 SBİO İNG"])
        self.assertEqual(
            name_words(
                "1.Kafedra müdiri, P.ü.f.d., dos.\nHacıyeva Rəna Cavadxan q.\n(yarımştat)\n1997-ci ildən işləyir"
            ),
            ("haciyeva", "rena", "cavadxan"),
        )
        self.assertEqual(name_words("b.ü.f.dMahmud Hümbətov Fərzəli Ştat"), ("mahmud", "humbetov", "ferzeli"))

    def test_umumi_subjects_are_not_totals(self):
        # 2026-10-15: «Ümumi kimya» / «Ümumi biologiya» fənləri yekun sətri kimi atılırdı.
        for subject in ("Ümumi kimya", "Ümumi biologiya BT/QM", "ATMF II blok\n2. Ümumi mikrobiologiya"):
            self.assertFalse(_is_total("236 QM", subject), subject)
        for total in ("Ümumi:", "Ştat ümumi:", "Saat hesabı ümumi:", "CƏMİ", "UNİVERSİTET ÜZRƏ CƏMİ"):
            self.assertTrue(_is_total("", total), total)
        self.assertEqual(hours_value("30x3"), 90)


class Form1ImportTests(TestCase):
    def setUp(self):
        self.org = make_org("qku-f1")
        faculty = OrgUnit.objects.create(organization=self.org, name="YTİM", unit_type=OrgUnitType.FACULTY)
        self.chair = OrgUnit.objects.create(
            organization=self.org, name="Proqramlaşdırma və informasiya təhlükəsizliyi",
            unit_type=OrgUnitType.CHAIR, parent=faculty,
        )  # fmt: skip
        specialty = OrgUnit.objects.create(
            organization=self.org, name="Ekologiya", slug="eko", unit_type=OrgUnitType.SPECIALTY, parent=faculty
        )
        self.group = OrgUnit.objects.create(
            organization=self.org, name="235 EKO", unit_type=OrgUnitType.GROUP, parent=specialty
        )
        for code, name in (("QKU-1", "Veb texnologiyalar"), ("QKU-2", "Alqoritmlər"), ("QKU-3", "Qraflar")):
            Subject.objects.create(organization=self.org, code=code, name=name, ects=5)
        self.admin = User.objects.create_superuser("superadmin", "sa@x.test", "pw")
        activate_member(self.org, self.admin, "rector", permissions=["*"], level=100)
        self.teachers = {}
        for username, first, last in (("sekine.h", "Səkinə", "Həsənova"), ("minaye.a", "Minayə", "Əliyeva")):
            user = User.objects.create_user(username, f"{username}@x.test", "pw", first_name=first, last_name=last)
            activate_member(
                self.org, user, "teacher", permissions=TEACHER_PERMS, scope_unit=self.chair,
                scope_type=RoleScopeType.COURSE,
            )  # fmt: skip
            self.teachers[username] = user
        self.tmp = TemporaryDirectory()
        self._tapsiriq()
        self.task = TeachingTask.objects.get(organization=self.org, chair=self.chair)

    def tearDown(self):
        self.tmp.cleanup()

    def _tapsiriq(self):
        """Tapşırıq sətirləri TAPŞIRIQ kitabçasından (müəllim sütunu boş) — qaralama sənəd."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Proqramlaşdırma"
        for _ in range(10):
            ws.append([None])
        for season, subject in (("PAYIZ", "Veb texnologiyalar"), ("PAYIZ", "Alqoritmlər"), ("YAZ", "Qraflar")):
            ws.append([season, "235 EKO", subject, "Ekologiya", "20", "30", "30", None, None, "60"])
        path = Path(self.tmp.name) / "tapsiriq.xlsx"
        wb.save(path)
        call_command(
            "import_teaching_task_workbook", "--file", str(path), "--org", self.org.slug, "--actor", "superadmin",
            "--sheet", "Proqramlaşdırma", "--no-create-teachers", "--apply", stdout=StringIO(),
        )  # fmt: skip

    def _docx(self, rows):
        document = Document()
        table = document.add_table(rows=3, cols=10)
        for col, text in enumerate(_HEAD + ["BİRİNCİ SEMESTR"] * 3 + ["İKİNCİ SEMESTR"] * 3):
            table.rows[0].cells[col].text = text
        for col, text in enumerate(_HEAD + ["Mühazirə", "Məşğələ (seminar)", "Laboratoriya"] * 2):
            table.rows[1].cells[col].text = text
        for col in range(10):
            table.rows[2].cells[col].text = str(col + 1)
        for values in rows:
            cells = table.add_row().cells
            for col, value in enumerate(values):
                cells[col].text = value
        path = Path(self.tmp.name) / "forma1.docx"
        document.save(path)
        return path

    def _run(self, path, *extra):
        out = StringIO()
        call_command(
            "import_teaching_load_form1", "--file", str(path), "--chair", "Proqramlaşdırma və informasiya "
            "təhlükəsizliyi kafedrası", "--org", self.org.slug, "--actor", "superadmin", "--report",
            self.tmp.name, *extra, stdout=out,
        )  # fmt: skip
        return out.getvalue()

    def _report(self):
        with (Path(self.tmp.name) / "forma1_import.csv").open(encoding="utf-8") as handle:
            return list(csv.DictReader(handle))

    def _rows(self, sekine="1.Müəllim Həsənova Səkinə Ramiz q. (Ştat)"):
        return [
            [sekine, "235EKO AZ", "Veb texnologiyalar", "20", "30", "30", "", "", "", ""],
            [sekine, "235 EKO", "ATMF I blok:\n1. Alqoritmlər +\n2. C++ proqramlaşdırma", "20", "30", "", "", "", "", ""],
            [sekine, "", "Ştat ümumi:", "", "", "", "", "", "", ""],
            [sekine, "235 EKO", "Alqoritmlər", "20", "", "30", "", "", "", ""],
            ["2. Naməlum Şəxs Əli o. (saat hesabı)", "235 EKO", "ATMF II blok:\n1. Qraflar\n2. Alqoritmlər", "20", "",
             "", "", "30", "30", ""],
            ["2. Naməlum Şəxs Əli o. (saat hesabı)", "CƏMİ", "CƏMİ", "", "", "", "", "30", "30", ""],
        ]  # fmt: skip

    def test_assigns_from_form1_without_creating_users(self):
        path = self._docx(self._rows())
        users_before = User.objects.count()
        out = self._run(path)
        self.assertIn("DRY-RUN", out)
        self.assertFalse(TeacherAssignment.objects.filter(row__task=self.task, teacher__isnull=False).exists())
        self.assertNotIn("Səkinə", out)
        out = self._run(path, "--apply", "--choice", "7=1")
        self.assertEqual(User.objects.count(), users_before)
        self.assertIn("teacher_not_found=1", out)
        veb = self.task.rows.get(subject__name="Veb texnologiyalar")
        self.assertEqual(
            set(TeacherAssignment.objects.filter(row=veb).values_list("teacher__username", flat=True)), {"sekine.h"}
        )
        alq = self.task.rows.get(subject__name="Alqoritmlər")
        lecture = TeacherAssignment.objects.get(row=alq, activity="lecture")
        seminar = TeacherAssignment.objects.get(row=alq, activity="seminar")
        self.assertEqual((lecture.teacher, lecture.is_hourly_paid), (self.teachers["sekine.h"], False))
        # «Ştat ümumi:» yekunundan sonrakı sətir saathesabıdır.
        self.assertEqual((seminar.teacher, seminar.is_hourly_paid), (self.teachers["sekine.h"], True))
        report = {row["line"]: row for row in self._report()}
        self.assertEqual(report["4"]["chosen_by"], "plus")
        self.assertEqual(report["7"]["chosen_by"], "override")
        self.assertEqual(report["7"]["teacher_action"], "not_found")
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "distributed")
        offering = CourseOffering.objects.get(organization=self.org, subject__name="Veb texnologiyalar")
        self.assertEqual(offering.instructor, self.teachers["sekine.h"])

    def test_reassign_from_form1(self):
        self._run(self._docx(self._rows()), "--apply", "--choice", "7=1")
        path = self._docx(self._rows(sekine="Əliyeva Minayə Kamal q. Ştat"))
        out = self._run(path, "--apply", "--choice", "7=1")
        self.assertIn("teacher_differs=", out)
        offering = CourseOffering.objects.get(organization=self.org, subject__name="Veb texnologiyalar")
        self.assertEqual(offering.instructor, self.teachers["sekine.h"])
        out = self._run(path, "--apply", "--reassign", "--choice", "7=1")
        self.assertIn("reassigned=", out)
        offering.refresh_from_db()
        self.assertEqual(offering.instructor, self.teachers["minaye.a"])
        veb = self.task.rows.get(subject__name="Veb texnologiyalar")
        self.assertEqual(
            set(TeacherAssignment.objects.filter(row=veb).values_list("teacher__username", flat=True)), {"minaye.a"}
        )

    def test_allow_outside_chair_accepts_a_unique_account_of_another_unit(self):
        """Sahib 2026-10-12: başqa kafedranın (və ya dekan müavini kimi qeyri-müəllim rolunun) TƏK hesabı bağlanır."""
        other_chair = OrgUnit.objects.create(
            organization=self.org, name="Ekologiya kafedrası", unit_type=OrgUnitType.CHAIR
        )
        outsider = User.objects.create_user("lamiye.a", "l@x.test", "pw", first_name="Lamiyə", last_name="Abasova")
        activate_member(
            self.org, outsider, "teacher", permissions=TEACHER_PERMS, scope_unit=other_chair,
            scope_type=RoleScopeType.COURSE,
        )  # fmt: skip
        vice = User.objects.create_user("aynure.q", "a@x.test", "pw", first_name="Aynurə", last_name="Qocuyeva")
        activate_member(self.org, vice, "vice_dean", permissions=["unit.view"], level=75)
        rows = [
            ["1. Abasova Lamiyə Vidadi q. (Ştat)", "235 EKO", "Veb texnologiyalar", "20", "30", "30", "", "", "", ""],
            ["2. Qocuyeva Aynurə Füzuli q.", "235 EKO", "Alqoritmlər", "20", "30", "30", "", "", "", ""],
        ]
        path = self._docx(rows)
        self._run(path)
        actions = {row["line"]: row["teacher_action"] for row in self._report()}
        self.assertEqual(set(actions.values()), {"outside_chair", "not_found"})
        self._run(path, "--allow-outside-chair")
        report = {row["line"]: row for row in self._report()}
        self.assertEqual({r["teacher_action"] for r in report.values()}, {"matched_outside_chair"})
        self.assertEqual({r["username"] for r in report.values()}, {"lamiye.a", "aynure.q"})
        # İki eyni adlı hesab varsa qəbul edilmir.
        twin = User.objects.create_user("lamiye.b", "l2@x.test", "pw", first_name="Lamiyə", last_name="Abasova")
        activate_member(self.org, twin, "teacher", permissions=TEACHER_PERMS, scope_unit=other_chair)
        self._run(path, "--allow-outside-chair")
        self.assertEqual(
            {r["line"]: r["teacher_action"] for r in self._report()}[report and min(report)], "outside_chair"
        )

    def test_outside_chair_teacher_is_assigned_with_a_secondary_chair_membership(self):
        other_chair = OrgUnit.objects.create(
            organization=self.org, name="Ekologiya kafedrası", unit_type=OrgUnitType.CHAIR
        )
        outsider = User.objects.create_user("lamiye.a", "l@x.test", "pw", first_name="Lamiyə", last_name="Abasova")
        activate_member(
            self.org, outsider, "teacher", permissions=TEACHER_PERMS, scope_unit=other_chair,
            scope_type=RoleScopeType.COURSE,
        )  # fmt: skip
        rows = [
            ["1. Abasova Lamiyə Vidadi q. (Ştat)", "235 EKO", "Veb texnologiyalar", "20", "30", "30", "", "", "", ""]
        ]
        out = self._run(self._docx(rows), "--apply", "--allow-outside-chair")
        self.assertIn("chair_membership_added=1", out)
        veb = self.task.rows.get(subject__name="Veb texnologiyalar")
        self.assertEqual(
            set(TeacherAssignment.objects.filter(row=veb).values_list("teacher__username", flat=True)), {"lamiye.a"}
        )
        memberships = Membership.objects.filter(user=outsider, organization=self.org, is_active=True)
        self.assertEqual({m.scope_unit_id for m in memberships}, {other_chair.pk, self.chair.pk})
        self.assertFalse(memberships.get(scope_unit=self.chair).is_primary)
