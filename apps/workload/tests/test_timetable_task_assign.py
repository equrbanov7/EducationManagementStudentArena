"""Sahib 2026-10-12 / 10-15: Forma 1 göndərməyən kafedranın boş tapşırıq sətrinə dərs cədvəlindəki
müəllim (``assign_task_rows_from_timetable``): boş fəaliyyət dolur + açılış yaranır, real müəllimli
fəaliyyətə toxunulmur, «Vakant» yeri düzəliş axını ilə dolur, qeyri-müəyyən müəllim yazılmır,
başqa kafedranın müəllimi yalnız bayraqla (əlavə kafedra üzvlüyü), dry-run heç nə yazmır, təkrar idempotent."""

from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase

import docx
from openpyxl import Workbook

from apps.organizations.models import Membership, OrgUnit
from apps.registrar.models import CourseOffering, Subject
from apps.workload.models import TeacherAssignment, TeachingTask
from core.constants import OrgUnitType, RoleScopeType

from .factories import TEACHER_PERMS, activate_member, make_org

User = get_user_model()
HEADER = ["Qruplar", "Qruplar", "", "III KURS 235 EKO 20"]
HOURS = ["Saatlar", "Saatlar", "aud", "Fənlər"]
SLOTS = [
    ("Bazar ertəsi", "8:30-10:00"),
    ("Bazar ertəsi", "10:10-11:40"),
    ("Çərşənbə axşamı", "8:30-10:00"),
    ("Çərşənbə axşamı", "10:10-11:40"),
    ("Cümə", "8:30-10:00"),
]


class TimetableTaskAssignTests(TestCase):
    def setUp(self):
        self.org = make_org("qku-tt")
        faculty = OrgUnit.objects.create(organization=self.org, name="YTİM", unit_type=OrgUnitType.FACULTY)
        self.chair = OrgUnit.objects.create(
            organization=self.org, name="Mexanika və riyaziyyat", unit_type=OrgUnitType.CHAIR, parent=faculty
        )
        self.other_chair = OrgUnit.objects.create(
            organization=self.org, name="Ekologiya kafedrası", unit_type=OrgUnitType.CHAIR, parent=faculty
        )
        specialty = OrgUnit.objects.create(
            organization=self.org, name="Ekologiya", slug="eko", unit_type=OrgUnitType.SPECIALTY, parent=faculty
        )
        OrgUnit.objects.create(organization=self.org, name="235 EKO", unit_type=OrgUnitType.GROUP, parent=specialty)
        for code, name in (("QKU-1", "Veb texnologiyalar"), ("QKU-2", "Alqoritmlər"), ("QKU-3", "Qraflar")):
            Subject.objects.create(organization=self.org, code=code, name=name, ects=5)
        self.admin = User.objects.create_superuser("superadmin", "sa@x.test", "pw")
        activate_member(self.org, self.admin, "rector", permissions=["*"], level=100)
        self.users = {}
        for username, first, last, chair in (
            ("sekine.h", "Səkinə", "Həsənova", self.chair),
            ("minaye.a", "Minayə", "Əliyeva", self.chair),
            ("rauf.m", "Rauf", "Məmmədov", self.chair),
            ("resad.m", "Rəşad", "Məmmədov", self.chair),
            ("lamiye.a", "Lamiyə", "Abasova", self.other_chair),
        ):
            user = User.objects.create_user(username, f"{username}@x.test", "pw", first_name=first, last_name=last)
            activate_member(
                self.org, user, "teacher", permissions=TEACHER_PERMS, scope_unit=chair,
                scope_type=RoleScopeType.COURSE,
            )  # fmt: skip
            self.users[username] = user
        self.tmp = TemporaryDirectory()
        self._tapsiriq()
        self.task = TeachingTask.objects.get(organization=self.org, chair=self.chair)
        self.veb = self.task.rows.get(subject__name="Veb texnologiyalar")
        self.alq = self.task.rows.get(subject__name="Alqoritmlər")

    def tearDown(self):
        self.tmp.cleanup()

    def _tapsiriq(self):
        """Forma 1-siz kafedra: tapşırıq sətirləri müəllimsiz (TAPŞIRIQ kitabçasının müəllim sütunu boş)."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Mexanika"
        for _ in range(10):
            ws.append([None])
        for season, subject in (("PAYIZ", "Veb texnologiyalar"), ("PAYIZ", "Alqoritmlər"), ("YAZ", "Qraflar")):
            ws.append([season, "235 EKO", subject, "Ekologiya", "20", "30", "30", None, None, "60"])
        path = Path(self.tmp.name) / "tapsiriq.xlsx"
        wb.save(path)
        call_command(
            "import_teaching_task_workbook", "--file", str(path), "--org", self.org.slug, "--actor", "superadmin",
            "--sheet", "Mexanika", "--no-create-teachers", "--apply",
            stdout=StringIO(),
        )  # fmt: skip

    def _docx(self, cells):
        document = docx.Document()
        table = document.add_table(rows=2 + len(cells), cols=4)
        for col, text in enumerate(HEADER):
            table.cell(0, col).text = text
        for col, text in enumerate(HOURS):
            table.cell(1, col).text = text
        for offset, ((day, time), text) in enumerate(zip(SLOTS, cells), start=2):
            for col, value in enumerate((day, time, str(10 + offset), text)):
                table.cell(offset, col).text = value
        for row, text in ((0, "Qruplar"), (1, "Saatlar")):
            table.cell(row, 0).merge(table.cell(row, 1)).text = text
        path = Path(self.tmp.name) / "cedvel.docx"
        document.save(str(path))
        return path

    def _run(self, path, *extra):
        out = StringIO()
        call_command(
            "assign_task_rows_from_timetable", "--file", str(path), "--org", self.org.slug,
            "--period", "Payız 2026/2027", "--year", "2026/2027", "--actor", "superadmin",
            "--report", self.tmp.name, *extra, stdout=out,
        )  # fmt: skip
        return out.getvalue()

    def _report(self):
        with (Path(self.tmp.name) / "timetable_task_assign.csv").open(encoding="utf-8") as handle:
            return list(csv.DictReader(handle))

    def _teachers(self, row, activity):
        return set(
            TeacherAssignment.objects.filter(row=row, activity=activity).values_list("teacher__username", flat=True)
        )

    def _real_assignments(self):
        return TeacherAssignment.objects.filter(row__task=self.task, teacher__isnull=False)

    def test_vacant_row_assigned_dry_run_first_then_idempotent(self):
        path = self._docx(["Veb texnologiyalar müh.\nHəsənova S.", "Veb texnologiyalar sem. Həsənova S."])
        before = self._real_assignments().count()
        out = self._run(path)
        self.assertIn("DRY-RUN", out)
        self.assertNotIn("Həsənova", out)
        self.assertIn("result_would_assign=2", out)
        self.assertEqual(self._real_assignments().count(), before)
        self.assertFalse(CourseOffering.objects.filter(organization=self.org).exists())
        self.assertEqual({row["status"] for row in self._report()}, {"planned"})

        out = self._run(path, "--apply")
        self.assertIn("result_assigned=2", out)
        self.assertEqual(self._teachers(self.veb, "lecture"), {"sekine.h"})
        self.assertEqual(self._teachers(self.veb, "seminar"), {"sekine.h"})
        lecture = TeacherAssignment.objects.get(row=self.veb, activity="lecture")
        self.assertEqual((lecture.hours, lecture.is_hourly_paid), (self.veb.lecture_total, False))
        offering = CourseOffering.objects.get(organization=self.org, subject__name="Veb texnologiyalar")
        self.assertEqual(offering.instructor, self.users["sekine.h"])
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "distributed")
        self.assertEqual({row["username"] for row in self._report()}, {"sekine.h"})

        count = self._real_assignments().count()
        out = self._run(path, "--apply")
        self.assertIn("cell_already_assigned=2", out)
        self.assertNotIn("result_assigned", out)
        self.assertEqual(self._real_assignments().count(), count)

    def test_assigned_activity_untouched_and_vacant_activity_filled(self):
        self._run(self._docx(["Veb texnologiyalar müh.\nHəsənova S."]), "--apply")
        # Təsdiqdən sonra seminar «Vakant»dır; mühazirənin müəllimi var. Sətirdə lab saatı yoxdur →
        # cədvəlin «lab.» hüceyrəsi seminara düşür (Forma 1 qaydası).
        self.assertEqual(self._teachers(self.veb, "seminar"), {None})
        path = self._docx(["Veb texnologiyalar müh.\nƏliyeva M.", "Veb texnologiyalar lab.\nƏliyeva M."])
        out = self._run(path, "--apply")
        self.assertIn("cell_already_assigned=1", out)
        self.assertIn("result_reassigned=1", out)
        self.assertEqual(self._teachers(self.veb, "lecture"), {"sekine.h"})
        self.assertEqual(self._teachers(self.veb, "seminar"), {"minaye.a"})
        self.assertEqual(self._report()[1]["activity"], "seminar")
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "distributed")
        offering = CourseOffering.objects.get(organization=self.org, subject__name="Veb texnologiyalar")
        self.assertEqual(offering.instructor, self.users["sekine.h"])

    def test_ambiguous_teacher_is_reported_not_written(self):
        out = self._run(self._docx(["Alqoritmlər müh.\nMəmmədov R."]), "--apply")
        self.assertIn("cell_teacher_ambiguous=1", out)
        self.assertEqual(self._report()[0]["status"], "teacher_ambiguous")
        self.assertFalse(self._real_assignments().exists())

    def test_conflicting_teachers_for_one_activity_are_skipped(self):
        path = self._docx(["Alqoritmlər sem.\nHəsənova S.", "Alqoritmlər sem.\nƏliyeva M."])
        out = self._run(path, "--apply")
        self.assertIn("cell_conflict=2", out)
        self.assertFalse(self._real_assignments().exists())

    def test_outside_chair_teacher_needs_flag_and_gets_secondary_membership(self):
        path = self._docx(["Alqoritmlər sem.\nAbasova L."])
        self._run(path, "--apply")
        self.assertEqual(self._report()[0]["status"], "teacher_outside_chair")
        self.assertFalse(self._real_assignments().exists())
        out = self._run(path, "--apply", "--allow-outside-chair")
        self.assertIn("chair_membership_added=1", out)
        self.assertEqual(self._teachers(self.alq, "seminar"), {"lamiye.a"})
        memberships = Membership.objects.filter(user=self.users["lamiye.a"], organization=self.org, is_active=True)
        self.assertEqual({m.scope_unit_id for m in memberships}, {self.other_chair.pk, self.chair.pk})
        self.assertFalse(memberships.get(scope_unit=self.chair).is_primary)

    def test_explicit_teacher_map_resolves_a_different_surname_spelling(self):
        # 2026-10-16: «Əmiraslanlı F.» ↔ «Əmiraslanova F.» — operatorun açıq xəritəsi.
        path = self._docx(["Alqoritmlər sem.\nAbasovalı L."])
        self._run(path, "--apply")
        self.assertEqual(self._report()[0]["status"], "teacher_not_found")
        out = self._run(path, "--apply", "--teacher", "Abasovalı L.=lamiye.a")
        self.assertIn("chair_membership_added=1", out)
        self.assertEqual(self._teachers(self.alq, "seminar"), {"lamiye.a"})
        self.assertEqual(self._report()[0]["teacher_action"], "teacher_map")

    def test_unmatched_subject_and_spring_rows_are_ignored(self):
        path = self._docx(["Qraflar müh.\nHəsənova S.", "Astronomiya müh.\nHəsənova S.", "TƏCRÜBƏ"])
        out = self._run(path, "--apply")
        self.assertIn("cell_subject_not_in_group=2", out)
        self.assertIn("cell_special_practice=1", out)
        self.assertFalse(self._real_assignments().exists())

    def test_distinct_subgroup_teachers_split_the_activity_evenly(self):
        cells = ["Alqoritmlər sem.\nHəsənova S. 235 EKO-1", "Alqoritmlər sem.\nƏliyeva M. 235 EKO-2"]
        self._run(self._docx(cells), "--apply")
        self.assertFalse(self._real_assignments().exists())  # subgroup_count=1 → conflict
        self.assertEqual({row["detail"] for row in self._report()}, {"teachers_differ"})
        self.alq.subgroup_count = 2
        self.alq.save(update_fields=["subgroup_count"])
        out = self._run(self._docx(cells), "--apply")
        self.assertIn("result_assigned=2", out)
        hours = dict(
            TeacherAssignment.objects.filter(row=self.alq, activity="seminar").values_list("teacher__username", "hours")
        )
        self.assertEqual(hours, {"sekine.h": self.alq.seminar_total // 2, "minaye.a": self.alq.seminar_total // 2})
