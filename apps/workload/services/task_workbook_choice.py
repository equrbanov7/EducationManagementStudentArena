"""Seçmə blokunda (ATMF) SEÇİLƏN fənnin müəyyənləşdirilməsi — DB-siz, saf funksiyalar.

Sahibin qaydası (2026-10-12): paylanmış TAPŞIRIQ kitabçasında blokun hər bəndi ayrıca
sətirdədir; seçilən bənd QARA/QALIN (bold) yazılır. Köhnə fayllarda seçim bəndin sonundakı
ayrıca «+» işarəsi ilə idi. Üstünlük sırası:

1. ``bold``   — yeganə qalın bənd (blokun başlığı «ATMF III blok:» da qalındır — sayılmır);
2. ``plus``   — yeganə bənddə AYRICA «+» işarəsi (söz daxilindəki «C++» işarə DEYİL);
3. ``single`` — blokda yalnız bir dolu bənd və ya «ATMF: X» forması;
4. ``override`` — operatorun açıq seçimi (``--choice LINE=N``) — YALNIZ işarəsiz xanaya;
5. ``none``   — işarə yoxdur → sətir yaradılmır/təyin olunmur, hesabatda «seçim işarəsi yoxdur».

Qalınlıq sətir səviyyəsində hesablanır: hərf-rəqəm simvollarının yarıdan çoxu qalındırsa
bənd qalın sayılır (Excel tez-tez bənd nömrəsini «2» ayrı, mətni ayrı formatlayır).
"""

from __future__ import annotations

import re
from dataclasses import dataclass

CHOSEN_BOLD = "bold"
CHOSEN_PLUS = "plus"
CHOSEN_SINGLE = "single"
CHOSEN_OVERRIDE = "override"
CHOSEN_NONE = "none"

_OPTION_RE = re.compile(r"^\s*(\d+)\s*[.)]?\s*(.*)$")
_LETTERS = "a-zəöüğışçA-ZƏÖÜĞIŞÇİ"
#: Blok başlığı: «ATMF III blok:», «ATMF5. BLOK», «ATMF2BLOK», «BLOK8», «SF 3. BLOK …»,
#: «1. SEÇMƏ FƏNN BLOKU», «II seçmə fənn bloku», «ÜFS-II blok», «ATMF: I blok:».
_BLOCK_HEADER_RE = re.compile(
    r"^\s*(?:\d+\s*[.)]?\s*)?(?:ATMF|ÜFS|UFS|SF|S\.F)?[\s.:\-]*(?:\d+|[IVXLC]+)?[\s.:\-]*"
    rf"(?:seçmə\s+fənn\s+)?blok(?:u)?(?![{_LETTERS}])\s*\d*\s*[.:]*\s*",
    re.IGNORECASE,
)
#: Başlıqsız prefiks: «ATMF: X», «ATMF Struktur biologiya» → X.
_PREFIX_RE = re.compile(rf"^\s*(?:ATMF|SF|S\.F)(?![{_LETTERS}])[\s.:]*", re.IGNORECASE)
#: «C+», «C++» — proqramlaşdırma dili adı, seçim işarəsi deyil.
_LANGUAGE_PLUS_RE = re.compile(r"(?:^|\s)[A-Za-z]\++$")


@dataclass(frozen=True)
class Option:
    number: int
    text: str
    bold: bool
    plus: bool


class ChoiceError(ValueError):
    """Operatorun ``--choice`` dəyəri xanaya uyğun deyil."""


def rich_lines(value) -> list[tuple[str, list]]:
    """Xana dəyəri → ``[(sətir, [simvol qalındırmı, ...]), ...]`` (sətir kənar boşluqsuz).

    Qəbul edir: düz mətn; openpyxl ``CellRichText`` (``TextBlock.font.b``); ``(mətn, qalın)``
    cütlərinin siyahısı (Word xanası — :func:`docx_cell_chunks`).
    """
    chars: list[tuple[str, bool]] = []
    if value is None:
        return []
    if isinstance(value, str):
        chars = [(ch, False) for ch in value]
    else:
        for block in value:
            if isinstance(block, str):
                chars += [(ch, False) for ch in block]
            elif isinstance(block, tuple):
                chars += [(ch, bool(block[1])) for ch in str(block[0])]
            else:
                bold = bool(getattr(getattr(block, "font", None), "b", False))
                chars += [(ch, bold) for ch in str(getattr(block, "text", ""))]
    lines, current = [], []
    for ch, bold in chars + [("\n", False)]:
        if ch != "\n":
            current.append((ch, bold))
            continue
        while current and current[0][0].isspace():
            current.pop(0)
        while current and current[-1][0].isspace():
            current.pop()
        if current:
            lines.append(("".join(c for c, _ in current), [b for _, b in current]))
        current = []
    return lines


def _is_bold(text: str, flags: list, start: int = 0) -> bool:
    alnum = [flag for ch, flag in zip(text[start:], flags[start:]) if ch.isalnum()]
    return bool(alnum) and sum(alnum) * 2 > len(alnum)


def docx_cell_chunks(cell) -> list[tuple[str, bool]]:
    """python-docx xanası → ``[(mətn, qalın), ...]``; abzaslar arasında «\\n»."""
    chunks: list[tuple[str, bool]] = []
    for index, paragraph in enumerate(cell.paragraphs):
        if index:
            chunks.append(("\n", False))
        chunks += [(run.text, run.bold is True) for run in paragraph.runs]
    return chunks


def _has_plus(text: str) -> bool:
    stripped = text.rstrip()
    if re.search(r"(?:^|\s)\+(?=\s|$)", stripped):
        return True
    return stripped.endswith("+") and not _LANGUAGE_PLUS_RE.search(stripped)


def _clean(text: str) -> str:
    text = re.sub(r"(?:^|\s)\+(?=\s|$)", " ", text)
    if text.rstrip().endswith("+") and not _LANGUAGE_PLUS_RE.search(text.rstrip()):
        text = text.rstrip().rstrip("+")
    return re.sub(r"\s+", " ", text).strip(" .;")


def _header_tail(text: str):
    """Blok başlığıdırsa başlıqdan sonrakı mətn («» ola bilər), deyilsə ``None``."""
    match = _BLOCK_HEADER_RE.match(text)
    if match:
        return match.end()
    match = _PREFIX_RE.match(text)
    return match.end() if match else None


def block_options(lines: list[tuple[str, list]]) -> tuple[list[Option], str]:
    """→ (bəndlər, başlıqdan sonrakı mətn). Başlıq sətri bənd deyil. Nömrəli bənd görünənə
    qədər nömrəsiz sətirlər ayrıca bənddir (sıra nömrəsi ilə: «SF 3. BLOK İnsan genetikası /
    Heyvan genetikası»), nömrəli bənddən sonra isə əvvəlkinin davamıdır («3.Tibbi» + «biologiya»)."""
    options: list[Option] = []
    header_tail = ""
    header_seen = numbered = False
    for text, flags in lines:
        tail_at = None if (options or header_seen) else _header_tail(text)
        if tail_at is not None:
            header_seen = True
            header_tail = text[tail_at:].strip()
            if header_tail:
                bold = _is_bold(text, flags, tail_at)
                options.append(Option(1, _clean(header_tail), bold, _has_plus(header_tail)))
            continue
        match = _OPTION_RE.match(text)
        if match:
            body = match.group(2)
            if body.strip(" .+"):  # boş bənd («2.») ötürülür
                bold = _is_bold(text, flags, match.start(2))
                options.append(Option(int(match.group(1)), _clean(body), bold, _has_plus(body)))
                numbered = True
        elif options and numbered:
            last = options[-1]
            options[-1] = Option(
                last.number,
                _clean(f"{last.text} {text}"),
                last.bold or _is_bold(text, flags),
                last.plus or _has_plus(text),
            )
        else:
            options.append(Option(len(options) + 1, _clean(text), _is_bold(text, flags), _has_plus(text)))
    return options, header_tail


def choose_block_subject(value, override: str = "") -> tuple[str, str]:
    """Blok xanası → ``(seçilən fənn, necə seçildi)``; bax modul başlığı."""
    options, header_tail = block_options(rich_lines(value))
    bold = [o for o in options if o.bold]
    if len(bold) == 1:
        return bold[0].text, CHOSEN_BOLD
    pool = bold or options
    plus = [o for o in pool if o.plus]
    if len(plus) == 1:
        return plus[0].text, CHOSEN_PLUS
    if len(options) == 1:
        return options[0].text, CHOSEN_SINGLE
    if override:
        return apply_override(options, override), CHOSEN_OVERRIDE
    return "", CHOSEN_NONE


def apply_override(options: list[Option], override: str) -> str:
    """``--choice`` dəyəri (bənd nömrəsi və ya fənn adı) xananın bəndləri ilə yoxlanılır."""
    from .task_workbook_parsing import slug_name

    override = override.strip()
    if override.isdigit():
        for option in options:
            if option.number == int(override):
                return option.text
        raise ChoiceError(f"bənd {override} xanada yoxdur (bəndlər: {[o.number for o in options]})")
    wanted = slug_name(override)
    for option in options:
        if slug_name(option.text) == wanted:
            return option.text
    raise ChoiceError("fənn adı xananın bəndlərindən heç birinə uyğun deyil")


def parse_choice_overrides(values) -> dict[int, str]:
    """``["144=1", "145=Tətbiqi statistika…"]`` və ya ``"144=1,145=1"`` → ``{144: "1", 145: "…"}``."""
    result: dict[int, str] = {}
    for value in values or []:
        parts = [value] if not re.fullmatch(r"\s*\d+\s*=\s*\d+\s*(,\s*\d+\s*=\s*\d+\s*)+", value) else value.split(",")
        for part in parts:
            line, sep, choice = part.partition("=")
            if not sep or not line.strip().isdigit() or not choice.strip():
                raise ChoiceError(f"--choice formatı LINE=N və ya LINE=Fənn adı olmalıdır: {part!r}")
            result[int(line)] = choice.strip()
    return result


def rich_subject_map(path, sheet_titles=None) -> dict[str, dict[int, object]]:
    """Kitabçanı ``rich_text=True`` ilə İKİNCİ dəfə açır → ``{vərəq: {sətir: C xanası}}``.

    ``read_only`` rejimi zəngin mətni oxumur, ona görə yalnız fənn sütunu (C) üçün tam
    yükləmə; rəqəmlər əvvəlki kimi ``data_only`` keçidindən gəlir.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, rich_text=True, data_only=True)
    try:
        result = {}
        for ws in workbook.worksheets:
            if sheet_titles is not None and ws.title not in sheet_titles:
                continue
            result[ws.title] = {
                cell.row: cell.value for (cell,) in ws.iter_rows(min_col=3, max_col=3) if cell.value is not None
            }
        return result
    finally:
        workbook.close()


__all__ = [
    "CHOSEN_BOLD",
    "CHOSEN_NONE",
    "CHOSEN_OVERRIDE",
    "CHOSEN_PLUS",
    "CHOSEN_SINGLE",
    "ChoiceError",
    "Option",
    "apply_override",
    "block_options",
    "choose_block_subject",
    "docx_cell_chunks",
    "parse_choice_overrides",
    "rich_lines",
    "rich_subject_map",
]
