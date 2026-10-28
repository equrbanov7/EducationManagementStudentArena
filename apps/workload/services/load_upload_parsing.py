"""Tədris şöbəsinin «Dərs yükü» faylı — Excel şablonu + oxu (DB-siz) (sahib 2026-10-23).

Sahib: «Tədris şöbəsi üçün bir Excel faylı olsun — oradan kafedranı seçsin, hansı qrup, mühazirə neçə saat,
seminar, laboratoriya, kredit, keçəcək müəllim — hamısı olsun, atdığım (Forma 1) kimi.»

İki format qəbul olunur, ikisi də :class:`LoadEntry` siyahısı verir (bir sətir × semestr = bir yazı):

* **Excel şablonu** (``build_template`` — «Şablonu yüklə» düyməsi): bir sətir = müəllim + qrup(lar) + fənn +
  kredit + tələbə sayı + saathesabı + I/II semestrin Mühazirə / Məşğələ (seminar) / Laboratoriya saatları.
  Sütunlar BAŞLIQ ETİKETİNDƏN tapılır (sıra dəyişsə də oxunur); müəllim xanası boşdursa yuxarıdakı müəllim
  davam edir (Forma 1-dəki kimi birləşdirilmiş xana); «CƏMİ» sətirləri atılır; saat «15*2», «9+8» kimi də ola bilər.
* **Forma 1 (Word)** — kafedraların indiki sənədi; mövcud oxuyucu :func:`.task_form1_parsing.parse_form1`.

Kafedra faylda DEYİL, səhifədə seçilir. Seçmə blokunun (ATMF) seçimi Excel-də yazılmır — sətirdə seçilmiş
fənnin ÖZÜ yazılır; bir neçə bəndli xana ``no_choice_marker`` kimi ötürülür (Forma 1-dəki qayda).
"""

from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass
from datetime import date, time, timedelta

from .task_form1_parsing import Form1Entry, group_key, group_tokens, hours_value, is_hourly_teacher, name_words
from .task_workbook_choice import CHOSEN_NONE
from .task_workbook_parsing import is_block, slug_name

MAX_ROWS = 2000
MAX_BYTES = 10 * 1024 * 1024

SHEET_TITLE = "Dərs yükü"
HELP_TITLE = "Təlimat"
#: Şablonun sütunları (sıra = Forma 1-in məntiqi sırası). Oxu etiketə görədir, sıraya görə yox.
TEMPLATE_HEADERS = (
    "Müəllim (Soyad Ad Ata adı, elmi dərəcə)",
    "Qrup(lar)",
    "Fənnin adı",
    "Kredit",
    "Tələbə sayı",
    "Saathesabı (bəli / boş)",
    "I semestr: Mühazirə",
    "I semestr: Məşğələ (seminar)",
    "I semestr: Laboratoriya",
    "II semestr: Mühazirə",
    "II semestr: Məşğələ (seminar)",
    "II semestr: Laboratoriya",
)
_ACTIVITIES = (("muhazire", "lecture"), ("mesgele", "seminar"), ("seminar", "seminar"), ("laborator", "lab"))
_YES = {"beli", "he", "yes", "1", "x", "+", "sh", "saathesabi"}


class LoadFileError(Exception):
    def __init__(self, code: str, message: str, line: int | None = None):
        super().__init__(code, message, line)
        self.code = code
        self.message = message
        self.line = line


@dataclass
class LoadEntry(Form1Entry):
    credits: str = ""
    students: str = ""

    def to_dict(self) -> dict:
        data = asdict(self)
        data["name_words"] = list(self.name_words)
        data["group_keys"] = sorted(self.group_keys)
        return data

    @classmethod
    def from_dict(cls, data: dict) -> LoadEntry:
        data = dict(data)
        data["name_words"] = tuple(data.get("name_words") or ())
        data["group_keys"] = frozenset(data.get("group_keys") or ())
        return cls(**data)


# ── Excel şablonu ─────────────────────────────────────────────────────────────


def build_template(chair_name: str = "") -> bytes:
    """Boş şablon (+ «Təlimat» vərəqi nümunə ilə) — ``.xlsx`` baytları."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from core.export_safety import sheet_append

    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TITLE
    sheet_append(sheet, TEMPLATE_HEADERS)
    head_fill = PatternFill("solid", fgColor="DBEAFE")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = (38, 22, 40, 8, 10, 12, 11, 11, 11, 11, 11, 11)
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.row_dimensions[1].height = 48
    sheet.freeze_panes = "A2"
    hourly = DataValidation(type="list", formula1='"bəli"', allow_blank=True)
    sheet.add_data_validation(hourly)
    hourly.add(f"F2:F{MAX_ROWS + 1}")

    help_sheet = book.create_sheet(HELP_TITLE)
    lines = [
        "Dərs yükü faylı — necə doldurulur",
        f"Kafedra: {chair_name}" if chair_name else "Kafedra faylda yazılmır — sistemdə yükləyəndə seçilir.",
        "",
        "1. Hər sətir: bir müəllimin bir fənni (bir və ya bir neçə qrupla). Müəllim xanası boş qalsa, yuxarıdakı "
        "müəllim davam edir.",
        "2. Müəllim: «Soyad Ad Ata adı» — elmi dərəcə/ad (b.ü.f.d., dos.) yazıla bilər, sistem onları atır. Yazı "
        "səhvi (bir hərf, ə/a, q/g) tanınır. Sistemdə olmayan müəllimə FİN-siz müəllim hesabı avtomatik açılır.",
        "3. Qrup(lar): «235 BİO İNG», «709/5 SBİO» — bir neçə qrup vergüllə və ya boşluqla. Qrup sistemdə "
        "olmalıdır (tələbələr oradan götürülür).",
        "4. Fənnin adı: seçmə blokunda (ATMF) yalnız SEÇİLMİŞ fənni yazın. Sistemdə olmayan fənn kataloqa əlavə olunur.",
        "5. Saatlar: hər semestr üçün Mühazirə / Məşğələ (seminar) / Laboratoriya. «15*2», «9+8» də olar.",
        "6. Saathesabı: müəllim bu fənni saathesabı ilə keçirsə «bəli».",
        "7. «CƏMİ» sətirləri yazmağa ehtiyac yoxdur — yazılsa atılır.",
        "8. Eyni faylı ikinci dəfə yükləmək heç nəyi təkrarlamır — yalnız yeni və ya fərqli sətirlər götürülür.",
        "",
        "Nümunə (bu vərəq oxunmur):",
    ]
    for line in lines:
        sheet_append(help_sheet, [line])
    help_sheet["A1"].font = Font(bold=True, size=13)
    sheet_append(help_sheet, TEMPLATE_HEADERS)
    for cell in help_sheet[help_sheet.max_row]:
        cell.font = Font(bold=True)
    sheet_append(
        help_sheet, ["b.ü.f.d. Məmmədov Ayaz Məmməd oğlu", "235 BİO İNG", "Genetika", 5, 23, "", 45, 30, "", "", "", ""]
    )
    sheet_append(help_sheet, ["", "233 BİO, 233 SBİO", "Bioinformatika", 4, 37, "", "", "", "", 30, "", 30])
    help_sheet.column_dimensions["A"].width = 110
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _semester(label: str) -> str:
    if label.startswith("iisem") or "ikinci" in label or label.startswith("yaz") or "iisemestr" in label:
        return "spring"
    if label.startswith("isem") or "birinci" in label or "payiz" in label:
        return "fall"
    return ""


def _columns(rows: list) -> tuple[int, dict]:
    """Başlıq sətri (ilk 10 sətirdə «Müəllim») + sütun xəritəsi; iki sətirli başlıqda semestr üst sətirdəndir."""
    for index, row in enumerate(rows[:10]):
        labels = [slug_name(value) for value in row]
        if not any(label.startswith("muellim") for label in labels):
            continue
        above = [slug_name(value) for value in rows[index - 1]] if index else []
        columns: dict = {"activities": {}}
        last_semester = ""
        for col, label in enumerate(labels):
            for key, needle in (
                ("teacher", "muellim"),
                ("groups", "qrup"),
                ("subject", "fenn"),
                ("credits", "kredit"),
                ("students", "telebe"),
                ("hourly", "saathesab"),
            ):
                if label.startswith(needle) and key not in columns:
                    columns[key] = col
            top = above[col] if col < len(above) else ""
            if top:
                last_semester = _semester(top) or last_semester
            semester = _semester(label) or (last_semester if top or last_semester else "")
            for needle, activity in _ACTIVITIES:
                if needle in label and semester:
                    columns["activities"].setdefault((semester, activity), col)
        missing = [name for name in ("teacher", "groups", "subject") if name not in columns]
        if missing or not columns["activities"]:
            raise LoadFileError(
                "columns_missing",
                "Faylda «Müəllim», «Qrup», «Fənnin adı» və semestr saatlarının sütunları tapılmadı — şablonu yükləyin.",
            )
        return index, columns
    raise LoadFileError("header_missing", "Başlıq sətri («Müəllim …») tapılmadı — sistemdəki şablonu işlədin.")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _is_total(*texts) -> bool:
    for text in texts:
        words = [slug_name(word) for word in re.findall(r"\w+", str(text or ""))]
        if words and (words[0] == "cemi" or words[-1] == "cemi" or "universitetuzre" in "".join(words)):
            return True
    return False


_DATE_OR_TIME = re.compile(
    r"^\s*(?:\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2}(?:\d{2})?)(?:\D|$)|\b\d{1,2}:\d{1,2}(?::\d{1,2})?"
)


def _technical_header(row) -> bool:
    """Excel's repeated Column1…Column12 labels are never a teacher/course."""
    return (
        len(row) >= len(TEMPLATE_HEADERS)
        and all(
            _text(value).casefold() == f"column{index}"
            for index, value in enumerate(row[: len(TEMPLATE_HEADERS)], start=1)
        )
        and not any(_text(value) for value in row[len(TEMPLATE_HEADERS) :])
    )


def _row_hours(raw_row, columns, line):
    result = {}
    for key, col in columns["activities"].items():
        if col >= len(raw_row):
            continue
        value = raw_row[col]
        if isinstance(value, (date, time, timedelta)) or (isinstance(value, str) and _DATE_OR_TIME.search(value)):
            raise LoadFileError(
                "invalid_hours",
                f"Sətir {line}: saat xanasında tarix və ya vaxt var. Saatı rəqəm kimi düzəldib faylı yenidən yükləyin.",
                line=line,
            )
        result[key] = hours_value(value)
    return result


def parse_workbook(stream) -> tuple[list[LoadEntry], list[dict]]:
    """Excel şablonu → (yazılar, ötürülən sətirlər)."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — korlanmış / Excel olmayan fayl
        raise LoadFileError("file_unreadable", "Fayl Excel (.xlsx) kimi oxunmadı.") from exc
    sheet = book[SHEET_TITLE] if SHEET_TITLE in book.sheetnames else book.worksheets[0]
    raw_rows = list(sheet.iter_rows(values_only=True))
    rows = [[_text(value) for value in row] for row in raw_rows]
    book.close()
    header, cols = _columns(rows)
    if len(rows) - header - 1 > MAX_ROWS:
        raise LoadFileError("too_many_rows", f"Faylda {MAX_ROWS}-dən çox sətir var — kafedra üzrə bölün.")
    entries, skipped = [], []
    teacher_text = ""
    for offset, row in enumerate(rows[header + 1 :], start=header + 2):

        def cell(name, row=row):
            col = cols.get(name)
            return row[col] if col is not None and col < len(row) else ""

        if not any(row):
            continue
        if _technical_header(raw_rows[offset - 1]):
            skipped.append({"line": offset, "reason": "technical_header", "subject": ""})
            continue
        row_hours = _row_hours(raw_rows[offset - 1], cols, offset)
        teacher_text = cell("teacher") or teacher_text
        groups_text, subject_text = cell("groups"), cell("subject")
        if _is_total(groups_text, subject_text):
            continue
        if not subject_text:
            if any(row[col] for col in cols["activities"].values() if col < len(row)):
                skipped.append({"line": offset, "reason": "subject_missing", "subject": ""})
            continue
        if not teacher_text:
            skipped.append({"line": offset, "reason": "teacher_missing", "subject": subject_text})
            continue
        block = is_block(subject_text)
        base = {
            "line": offset,
            "teacher_text": teacher_text,
            "name_words": name_words(teacher_text),
            "hourly": slug_name(cell("hourly")) in _YES or is_hourly_teacher(teacher_text),
            "groups_text": groups_text,
            "group_keys": frozenset(group_key(token) for token in group_tokens(groups_text)),
            "subject_text": subject_text,
            "chosen_subject": "",
            "chosen_by": CHOSEN_NONE if block else "",
            "credits": cell("credits")[:20],
            "students": cell("students")[:100],
        }
        produced = False
        for season in ("fall", "spring"):
            hours = {activity: value for (sem, activity), value in row_hours.items() if sem == season and value}
            if hours:
                entries.append(LoadEntry(season=season, hours=hours, **base))
                produced = True
        if not produced:
            skipped.append({"line": offset, "reason": "no_contact_hours", "subject": subject_text})
        elif block:
            skipped.append({"line": offset, "reason": "no_choice_marker", "subject": subject_text})
    return entries, skipped


def parse_form1_docx(stream) -> tuple[list[LoadEntry], list[dict]]:
    """Forma 1 (Word) → eyni :class:`LoadEntry` siyahısı (kredit/tələbə sayı Forma 1-də ayrıca oxunmur)."""
    from docx import Document

    from .task_form1_parsing import parse_form1

    try:
        document = Document(stream)
    except Exception as exc:  # noqa: BLE001
        raise LoadFileError("file_unreadable", "Fayl Word (.docx) kimi oxunmadı.") from exc
    if not document.tables:
        raise LoadFileError("header_missing", "Word faylında Forma 1 cədvəli tapılmadı.")
    try:
        found, skipped = parse_form1(document)
    except ValueError as exc:
        raise LoadFileError("columns_missing", str(exc)) from exc
    fields = Form1Entry.__dataclass_fields__
    return [LoadEntry(**{name: getattr(entry, name) for name in fields}) for entry in found], skipped


def parse_upload(upload) -> tuple[list[LoadEntry], list[dict], str]:
    """Yüklənən fayl (``.xlsx`` / ``.docx``) → (yazılar, ötürülənlər, format)."""
    name = (getattr(upload, "name", "") or "").lower()
    size = getattr(upload, "size", 0) or 0
    if size > MAX_BYTES:
        raise LoadFileError("file_too_large", "Fayl 10 MB-dan böyükdür.")
    data = io.BytesIO(upload.read())
    if name.endswith((".xlsx", ".xlsm")):
        entries, skipped = parse_workbook(data)
        return entries, skipped, "xlsx"
    if name.endswith(".docx"):
        entries, skipped = parse_form1_docx(data)
        return entries, skipped, "docx"
    raise LoadFileError("file_type", "Yalnız Excel şablonu (.xlsx) və ya Forma 1 (.docx) qəbul olunur.")


def safe_filename(chair_name: str) -> str:
    base = slug_name(chair_name)[:40] or "kafedra"
    return re.sub(r"[^a-z0-9]+", "-", f"ders-yuku-{base}").strip("-") + ".xlsx"


__all__ = [
    "MAX_ROWS",
    "TEMPLATE_HEADERS",
    "LoadEntry",
    "LoadFileError",
    "build_template",
    "parse_form1_docx",
    "parse_upload",
    "parse_workbook",
    "safe_filename",
]
