"""Dərs cədvəli .docx faylının (məktəbin «əsas dərs cədvəli») TƏMİZ oxunması — DB-siz.

Sahib 2026-10-12: «cədvəli oxu; orada dərsi olan hər kəsi sistemdə seç və yüklə — hamının
cədvəlinə düşsün». Bu modul yalnız Word cədvəllərini hüceyrə qeydlərinə çevirir; qrup/fənn/müəllim
tutuşdurması ``schedule_docx_match``-da, yazı ``schedule_docx_import``-dadır.

SƏNƏDİN QURULUŞU (8 cədvəl, hər biri bir kursun bir neçə qrupu)
----------------------------------------------------------------
* «Qruplar» sətri — hər qrupun başlığı («IV KURS 233 İ ing 8» = kurs + qrup kodu + tələbə sayı);
  başlığı BOŞ sütunlar «aud» (otaq) sütunlarıdır və SAĞINDAKI qrupa aiddir. Başlıq hüceyrəsi bir
  neçə sütunu tuta bilər (üfüqi birləşmə) — qrupun fənn sütunları onlardır.
* «Saatlar» sətri — adətən «aud | Fənlər»; bəzən qrup sütunlarının altında ALT QRUP adları
  («234 K-1 20», «234 K-2 17») — onda hər sütun bir alt qrupdur.
* Məlumat sətirləri: gün («Bazar ertəsi» …) + saat («8:30-\\n10:00») + hər qrup üçün (otaq, fənn).

BİRLƏŞMƏLƏR: python-docx birləşmiş hüceyrənin mətnini hər sütun/sətirdə TƏKRARLAYIR; eyni hüceyrə
``cell._tc`` kimliyi ilə tanınır və BİR dəfə sayılır (üfüqi: bir qrupun iki sütunu; şaquli: üst/alt
həftə cütünün hər iki sətri).

ÜST/ALT HƏFTƏ: eyni gün + eyni saat ARDICIL iki sətirdə gəlirsə bu, bölünmüş hüceyrədir —
yuxarı sətir «üst həftə» (``WeekType.ODD``), aşağı sətir «alt həftə» (``EVEN``). Hər iki sətirdə
EYNİ dərs (eyni hüceyrə və ya eyni mətn + otaq) → hər həftə (``ALL``). Mətnin içində açıq «üst» /
«alt» sözü varsa o, mövqedən üstündür (qeyd: ``week_marker``). Tək sətir → ``ALL``.

FƏNN HÜCEYRƏSİ: «<fənn> <növ>.\\n<Soyad İ.>»; növ: «müh» = mühazirə, «sem» = seminar, «lab»/«l» =
laboratoriya; birləşmiş növ («müh/sem», «l+sem», «lab/sem») ilk/«lab» seçilir və qeyd olunur. Növsüz
hüceyrə ``kind=""`` qalır — tutuşdurma onu fənnin saat bölgüsündən çıxarmağa çalışır. «TƏCRÜBƏ»
(təcrübə) və «LEVEL» (ingilis dili səviyyə qrupları) dərs slotu DEYİL — ``special`` ilə işarələnir.
"""

from __future__ import annotations

import datetime
import re
import unicodedata
from dataclasses import dataclass, field

KIND_LECTURE = "lecture"
KIND_SEMINAR = "seminar"
KIND_LAB = "lab"

WEEK_ALL = "all"
WEEK_ODD = "odd"
WEEK_EVEN = "even"

SPECIAL_PRACTICE = "practice"
SPECIAL_LEVEL = "level"
SPECIAL_JUNK = "junk"

_WEEKDAYS = (
    ("bazar ertesi", 1),
    ("cersenbe axsami", 2),
    ("cersenbe", 3),
    ("cume axsami", 4),
    ("cume", 5),
    ("senbe", 6),
    ("bazar", 7),
)

_AZ_ASCII = str.maketrans(
    {"ə": "e", "Ə": "e", "ö": "o", "Ö": "o", "ü": "u", "Ü": "u", "ğ": "g", "Ğ": "g", "ş": "s", "Ş": "s"}
    | {"ç": "c", "Ç": "c", "ı": "i", "I": "i", "İ": "i"}
)

_TIME_RE = re.compile(r"(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})")
_ROOM_RE = re.compile(r"^(\(?\d{1,4}\)?\s?\d{0,4}|akt zal[ıi](?: [a-zA-Z])?)$", re.IGNORECASE)
_WORD = r"[0-9A-Za-zƏəÖöÜüĞğŞşÇçİıI]"
#: Növ işarəsi — tam söz kimi («mühəndis» UYĞUN GƏLMİR); «(müh.)», «müh/sem», «l+sem», «lab/sem».
_KIND_RE = re.compile(
    rf"(?<!{_WORD})\(?(müh|muh|sem|lab|l(?=\s*[/+]))(?:\s*[/+]\s*(müh|muh|sem|lab))?\)?(?:\s*\.)?\)?(?!{_WORD})",
    re.IGNORECASE,
)
_WEEK_MARKER_RE = re.compile(rf"(?<!{_WORD})(üst|ust|alt)(?:\s+h[əe]ft[əe])?(?!{_WORD})", re.IGNORECASE)
_INITIAL_RE = re.compile(r"^[A-ZƏÖÜĞŞÇİI]\.?$")
_SURNAME_RE = re.compile(r"^[A-ZƏÖÜĞŞÇİ][a-zəöüğşçı-]+$")
_FRAGMENT_RE = re.compile(r"^[a-zəöüğşçı]{1,3}$")
_GROUP_RE = re.compile(r"^(\d{3,4})\s*([^\W\d_]+?)\s*(?:-(\d)(?!\d))?\s*(az|ing|rus|en|ru)?\s*[-–]?\s*([\d+\s]*)$")
_SECTOR_WORDS = (
    (re.compile(r"\b(?:İNG|İng|ING|Ing)\b"), "ing"),
    (re.compile(r"\b(?:RUS|Rus)\b"), "rus"),
    (re.compile(r"\b(?:AZ|Az)\b"), "az"),
)


def fold(value) -> str:
    """AZ hərfləri ASCII, kiçik, tək boşluq — müqayisə açarı («Çərşənbə axşamı» → «cersenbe axsami»)."""
    text = unicodedata.normalize("NFC", str(value or "")).translate(_AZ_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    return " ".join(text.split())


def squash(value) -> str:
    """Bütün boşluqlar tək boşluğa (sətir sonları daxil)."""
    return " ".join(str(value or "").replace("\t", " ").split())


def is_room(text: str) -> bool:
    return bool(_ROOM_RE.match(squash(text)))


def parse_weekday(text) -> int:
    key = fold(text)
    for name, number in _WEEKDAYS:
        if key == name:
            return number
    return 0


def parse_time(text):
    match = _TIME_RE.search(re.sub(r"\s+", "", str(text or "")))
    if not match:
        return None, None
    h1, m1, h2, m2 = (int(part) for part in match.groups())
    try:
        start, end = datetime.time(h1, m1), datetime.time(h2, m2)
    except ValueError:
        return None, None
    return (start, end) if start < end else (None, None)


@dataclass(frozen=True)
class GroupLabel:
    """«233 İTex ing 12» → digits=233, code=«İTex», sector=ing, sub=«», size=«12»."""

    raw: str
    digits: str
    code: str
    sector: str
    sub: str = ""

    @property
    def display(self) -> str:
        sub = f"-{self.sub}" if self.sub else ""
        return f"{self.digits} {self.code}{sub} {self.sector}".strip()


def parse_group_label(text) -> GroupLabel | None:
    """Başlıq/qrup adı → :class:`GroupLabel` («IV KURS» prefiksi və tələbə sayı atılır)."""
    raw = squash(text)
    body = re.sub(r"^[IVX]+\s*KURS\s*", "", raw, flags=re.IGNORECASE).strip()
    for pattern, word in _SECTOR_WORDS:
        body = pattern.sub(word, body)
    match = _GROUP_RE.match(body)
    if not match:
        return None
    digits, code, sub, sector, _size = match.groups()
    sector = {"en": "ing", "ru": "rus"}.get((sector or "az").lower(), (sector or "az").lower())
    return GroupLabel(raw=raw, digits=digits, code=code, sector=sector, sub=sub or "")


@dataclass
class SubjectText:
    raw: str
    subject: str = ""
    kind: str = ""
    kind_note: str = ""
    teacher: str = ""
    week_marker: str = ""
    special: str = ""


def _clean_teacher(text: str) -> str:
    """«Əzizov  a S.» → «Əzizova S.»; «.Məmmədova L» → «Məmmədova L.»; «Yusifov A..» → «Yusifov A.»."""
    tokens = [tok for tok in squash(text).strip(" .,").split(" ") if tok and tok != "."]
    if not tokens:
        return ""
    initial = ""
    if len(tokens) > 1 and _INITIAL_RE.match(tokens[-1]):
        initial = tokens.pop()[0]
    surname = tokens[0].lstrip(".")
    for tok in tokens[1:]:
        if _FRAGMENT_RE.match(tok):  # «Əzizov a» — sözün içində yazı xətası boşluğu
            surname += tok
        elif _INITIAL_RE.match(tok) and not initial:
            initial = tok[0]
        else:
            return ""  # ad kimi görünmür
    if not _SURNAME_RE.match(surname):
        return ""
    return f"{surname} {initial}." if initial else surname


def _split_trailing_teacher(text: str) -> tuple[str, str]:
    """Növsüz mətn: son sətir (və ya sətrin sonu) «Soyad İ.» formasındadırsa onu ayır."""
    lines = [squash(line) for line in str(text or "").splitlines() if squash(line)]
    if len(lines) > 1:
        teacher = _clean_teacher(lines[-1])
        if teacher:
            return squash(" ".join(lines[:-1])), teacher
    flat = squash(" ".join(lines))
    match = re.search(r"\s([A-ZƏÖÜĞŞÇİ][a-zəöüğşçı]+\s+[A-ZƏÖÜĞŞÇİ]\.?)\s*$", flat)
    if match:
        teacher = _clean_teacher(match.group(1))
        if teacher:
            return flat[: match.start()].strip(), teacher
    return flat, ""


def _kind_of(first: str, second: str) -> tuple[str, str]:
    names = [fold(part) for part in (first, second) if part]
    mapped = []
    for name in names:
        if name in ("muh",):
            mapped.append(KIND_LECTURE)
        elif name == "sem":
            mapped.append(KIND_SEMINAR)
        else:  # lab / l
            mapped.append(KIND_LAB)
    if len(mapped) == 1:
        return mapped[0], ""
    if KIND_LAB in mapped:
        return KIND_LAB, "combined"
    return mapped[0], "combined"


def parse_subject_text(text) -> SubjectText:
    """Fənn hüceyrəsinin mətni → fənn adı + növ + müəllim nişanı + həftə işarəsi."""
    raw = str(text or "").replace("\t", " ")
    flat = squash(raw)
    result = SubjectText(raw=flat)
    if not flat:
        return result
    folded = fold(flat)
    if folded.replace(" ", "") in ("tecrube", "tecrub", "tecrubesi"):
        result.special = SPECIAL_PRACTICE
        return result
    if folded == "level":
        result.special = SPECIAL_LEVEL
        return result
    if len(folded.replace(" ", "")) < 3:
        result.special = SPECIAL_JUNK
        return result
    markers = list(_WEEK_MARKER_RE.finditer(raw))
    if markers:
        result.week_marker = WEEK_ODD if fold(markers[-1].group(1)) == "ust" else WEEK_EVEN
        raw = _WEEK_MARKER_RE.sub(" ", raw)
    kinds = [m for m in _KIND_RE.finditer(raw) if m.start() > 0]
    if kinds:
        match = kinds[-1]
        result.kind, result.kind_note = _kind_of(match.group(1), match.group(2))
        result.subject = squash(raw[: match.start()]).strip(" .,(")
        result.teacher = _clean_teacher(raw[match.end() :])
        if not result.teacher and squash(raw[match.end() :]):
            result.kind_note = (result.kind_note + " teacher_unparsed").strip()
        return result
    result.subject, result.teacher = _split_trailing_teacher(raw)
    result.subject = result.subject.strip(" .,")
    return result


@dataclass
class Entry:
    """Bir qrupun (alt qrupun) bir slotu — idxalın vahidi."""

    table: int
    row: int
    group: GroupLabel
    weekday: int
    start: datetime.time
    end: datetime.time
    week_type: str
    room: str
    text: SubjectText
    notes: list[str] = field(default_factory=list)


@dataclass
class _Group:
    label: GroupLabel
    aud_cols: list[int]
    subject_cols: list[int]
    sub_labels: dict[int, GroupLabel] = field(default_factory=dict)


@dataclass
class ParseResult:
    entries: list[Entry] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)
    groups: list[str] = field(default_factory=list)
    tables: int = 0


def _header_rows(rows) -> tuple[int, int]:
    header = next((i for i, row in enumerate(rows) if fold(row[0].text) == "qruplar"), -1)
    hours = next((i for i, row in enumerate(rows) if i > header and fold(row[0].text) == "saatlar"), -1)
    return header, hours


def _groups(rows, header: int, hours: int, table_index: int, result: ParseResult) -> list[_Group]:
    cells = rows[header]
    groups, pending, col = [], [], 2
    while col < len(cells):
        tc = cells[col]._tc
        span = [col]
        while col + 1 < len(cells) and cells[col + 1]._tc is tc:
            col += 1
            span.append(col)
        text = squash(cells[span[0]].text)
        if not text:
            pending.extend(span)
        else:
            label = parse_group_label(text)
            if label is None:
                result.skipped.append({"table": table_index, "row": header, "reason": "group_header_unparsed"})
            else:
                groups.append(_Group(label=label, aud_cols=pending, subject_cols=span))
            pending = []
        col += 1
    if hours >= 0:
        sub_cells = rows[hours]
        for group in groups:
            labels = {}
            for col in group.subject_cols:
                text = squash(sub_cells[col].text) if col < len(sub_cells) else ""
                sub = parse_group_label(text) if text and fold(text) not in ("fenler", "aud") else None
                if sub is not None and sub.sub:
                    labels[col] = GroupLabel(
                        raw=sub.raw, digits=sub.digits, code=sub.code, sector=group.label.sector, sub=sub.sub
                    )
            if len(labels) == len(group.subject_cols) and len(labels) > 1:
                group.sub_labels = labels
    return groups


def _room(row_cells, group: _Group) -> str:
    for col in group.aud_cols:
        if col < len(row_cells):
            text = squash(row_cells[col].text)
            if text:
                return text[:64]
    return ""


def _subject_cells(row_cells, group: _Group, foreign_aud_tcs: set) -> list:
    seen, out = set(), []
    for col in group.subject_cols:
        if col >= len(row_cells):
            continue
        cell = row_cells[col]
        if cell._tc in seen or cell._tc in foreign_aud_tcs:
            continue
        seen.add(cell._tc)
        if squash(cell.text) and not is_room(cell.text):
            out.append(cell)
    return out


def _units(row_cells, group: _Group, foreign_aud_tcs: set) -> list[tuple[GroupLabel, object]]:
    """Bir sətirdə qrupun (və ya alt qruplarının) fənn hüceyrələri: [(label, cell|None), …]."""
    if group.sub_labels:
        cells = [(group.sub_labels[col], row_cells[col]) for col in group.subject_cols if col < len(row_cells)]
        texts = {fold(cell.text) for _label, cell in cells}
        if len(texts) == 1:  # bütün alt qruplarda EYNİ dərs → qrupun özünün dərsi
            cell = cells[0][1]
            return [(group.label, cell if squash(cell.text) else None)]
        return [(label, cell if squash(cell.text) else None) for label, cell in cells]
    cells = _subject_cells(row_cells, group, foreign_aud_tcs)
    if not cells:
        return [(group.label, None)]
    return [(group.label, cell) for cell in cells]


def _key(cell, room: str) -> tuple:
    text = parse_subject_text(cell.text)
    return (fold(text.subject), text.kind, fold(text.teacher), fold(room))


def _week_entries(table_index, slot_rows, group, foreign_aud_tcs, weekday, start, end) -> list[Entry]:
    """Bir saat zolağının (1 və ya 2 sətir) bir qrup üçün qeydləri — üst/alt həftə qaydası ilə."""
    per_row = []
    for row_index, row_cells in slot_rows:
        room = _room(row_cells, group)
        per_row.append((row_index, room, _units(row_cells, group, foreign_aud_tcs)))
    entries = []
    labels = []
    for _row_index, _room_text, units in per_row:
        for label, _cell in units:
            if label not in labels:
                labels.append(label)
    for label in labels:
        found = []
        for position, (row_index, room, units) in enumerate(per_row):
            for unit_label, cell in units:
                if unit_label == label and cell is not None:
                    found.append((position, row_index, room, cell))
        if not found:
            continue
        if len(per_row) == 1:
            plan = [(WEEK_ALL, item) for item in found]
        elif len(per_row) > 2:
            plan = [(WEEK_ALL, item) for item in found]
        else:
            top = [item for item in found if item[0] == 0]
            bottom = [item for item in found if item[0] == 1]
            same = len(top) == len(bottom) == 1 and (
                top[0][3]._tc is bottom[0][3]._tc or _key(top[0][3], top[0][2]) == _key(bottom[0][3], bottom[0][2])
            )
            if same:
                plan = [(WEEK_ALL, top[0])]
            else:
                plan = [(WEEK_ODD, item) for item in top] + [(WEEK_EVEN, item) for item in bottom]
        for week_type, (_position, row_index, room, cell) in plan:
            text = parse_subject_text(cell.text)
            notes = []
            if len(per_row) > 2:
                notes.append("more_than_two_rows")
            if text.week_marker and text.week_marker != week_type:
                notes.append(f"week_marker:{week_type}->{text.week_marker}")
                week_type = text.week_marker
            entries.append(
                Entry(
                    table=table_index,
                    row=row_index,
                    group=label,
                    weekday=weekday,
                    start=start,
                    end=end,
                    week_type=week_type,
                    room=room,
                    text=text,
                    notes=notes,
                )
            )
    return entries


def parse_tables(tables) -> ParseResult:
    """``docx.Document(...).tables`` → :class:`ParseResult` (DB-siz, test üçün sintetik sənəd də olar)."""
    result = ParseResult(tables=len(tables))
    for table_index, table in enumerate(tables):
        rows = [row.cells for row in table.rows]
        header, hours = _header_rows(rows)
        if header < 0:
            result.skipped.append({"table": table_index, "row": 0, "reason": "no_group_header"})
            continue
        groups = _groups(rows, header, hours, table_index, result)
        result.groups.extend(group.label.display for group in groups)
        bands: list[tuple[tuple, list]] = []
        weekday = 0
        for row_index in range(max(header, hours) + 1, len(rows)):
            row_cells = rows[row_index]
            if len(row_cells) < 3:
                continue
            weekday = parse_weekday(row_cells[0].text) or weekday
            start, end = parse_time(row_cells[1].text)
            has_content = any(squash(cell.text) for cell in row_cells[2:])
            if start is None:
                if has_content:
                    result.skipped.append({"table": table_index, "row": row_index, "reason": "no_time"})
                continue
            if not weekday:
                result.skipped.append({"table": table_index, "row": row_index, "reason": "no_weekday"})
                continue
            key = (weekday, start, end)
            if bands and bands[-1][0] == key:
                bands[-1][1].append((row_index, row_cells))
            else:
                bands.append((key, [(row_index, row_cells)]))
        for (band_weekday, start, end), slot_rows in bands:
            for group in groups:
                foreign = {
                    cells[col]._tc
                    for _ri, cells in slot_rows
                    for other in groups
                    if other is not group
                    for col in other.aud_cols
                    if col < len(cells)
                }
                result.entries.extend(_week_entries(table_index, slot_rows, group, foreign, band_weekday, start, end))
    return result


def parse_docx(path) -> ParseResult:
    import docx  # python-docx — yalnız idxal zamanı

    return parse_tables(docx.Document(str(path)).tables)


__all__ = [
    "Entry",
    "GroupLabel",
    "ParseResult",
    "SubjectText",
    "fold",
    "is_room",
    "parse_docx",
    "parse_group_label",
    "parse_subject_text",
    "parse_tables",
    "parse_time",
    "parse_weekday",
    "squash",
]
