"""Yazılı imtahan ballarının FAYLDAN (XLSX / CSV) köçürülməsi — İmtahan Mərkəzi.

Sahibin tələbi (2026-09-12): «tələbələrin balları sistemə yüklənsin». Axın:

1. **Şablon** — seçilmiş açılışın siyahısı ilə DOLDURULMUŞ fayl endirilir
   (Tələbə № · FİN · Ad Soyad · Qrup · Cari bal · Bal); operator yalnız «Bal»
   sütununu doldurur.
2. **Quru icra** — fayl oxunur, hər sətir siyahıya uyğunlaşdırılır və
   yoxlanılır (naməlum tələbə, dublikat, diapazon, qeydiyyatsız, dəyişiklik =
   səbəb tələb edir). HEÇ NƏ YAZILMIR.
3. **Tətbiq** — EYNİ plan qurucusu, sonra sətirlər YALNIZ
   ``exam_score_entry.save_roster_scores`` → ``record_exam_score`` ilə yazılır
   (tək yazı yolu; hər sətir üçün bir ``ExamScoreEntry``, partiya üçün bir
   ``ExamScoreSheet`` xülasəsi).

Fayl serverdə SAXLANILMIR — tətbiq eyni faylı yenidən göndərir. ``openpyxl``
``requirements/base.txt``-dədir (3.1.5); yoxdursa CSV yolu qalır.

MODUL SƏRHƏDİ: ``apps.accounts.services.intake`` oxuyucusu TƏKRAR YAZILMIR,
çünki registrar ``accounts``-u import edə bilməz (``scripts/module_deps.py``).

2026-09-14 (W2 `w2paper`): fayl oxuyucusu ``exam_score_import_reader``-dədir
(re-eksport); «S1..Sn» sual sütunları dəstəklənir — sətirdə S xanası doludursa
bal onların cəmidir (``question_scores`` plana və servisə ötürülür). Şablona
S1..S<n> sütunları da əlavə olunur (``question_count``).

2026-09-14 (W6 `w6paper`): şablon vərəqin ÖZ şəbəkəsi (``question_count`` /
``question_max``) və imtahan növü ilə gəlir — «İmtahan növü» sütunu (XLSX-də
siyahı seçimi) vərəq səviyyəli növün fayldakı izidir; ``build_plan`` seçilmiş
növlə uyğunsuzluğu sətir xətası kimi göstərir. Quru icra və tətbiq EYNİ
şəbəkəni alır (view POST-dan ötürür) — «gördüyün nəticə = alacağın nəticə».
"""

from __future__ import annotations

import csv
import io
import re
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.utils.translation import pgettext

from . import exam_score_entry as service
from . import exam_score_import_reader as reader
from . import exam_score_questions as questions
from .exam_score_import_reader import (  # noqa: F401 — re-eksport (çağıranlar `importer.read_rows` / sabitlər)
    ALLOWED_SUFFIXES,
    AZ_TRANSLIT,
    MAX_ROWS,
    MAX_UPLOAD_BYTES,
    SHEET_NAME,
    ImportFileError,
    normalize_header,
    normalize_name,
    question_count_in,
)
from .exam_score_import_safety import export_text
from .models.exam_score_entry import ExamScoreSheetKind

_CTX = "registrar.exam_score_import"
#: Vərəq / bölmə mətnləri (mövcud msgid-lər: «İmtahan növü», növ xətası) — kataloq təkrarlanmır.
_CTX_ENTRY = "registrar.exam_score_entry"

#: Şablonun sabit (S-dən əvvəlki) sütun sayı: Tələbə № · FİN · Ad Soyad · Qrup ·
#: İmtahan növü · Cari bal · Bal. «Bal» 7-ci (G), S1 8-ci (H) sütundur.
_FIXED_COLUMN_COUNT = 7
_KIND_COLUMN_INDEX = 5
_SCORE_COLUMN_INDEX = 7

#: Sətir vəziyyətləri (quru icra + tətbiq eyni açarları işlədir).
STATUS_NEW = "new"  # cari bal yoxdur → sərbəst yazılacaq
STATUS_CHANGE = "change"  # cari bal var və fərqlidir → səbəb + qeyd + sənəd
STATUS_UNCHANGED = "unchanged"  # eyni bal → idempotent, toxunulmur
STATUS_SKIP = "skip"  # boş bal → toxunulmur
STATUS_ERROR = "error"  # bloklayan xəta


def read_rows(uploaded_file) -> list:
    """Yüklənmiş faylı sətir siyahısına çevirir — ``exam_score_import_reader`` üzərindən, BU modulun limitləri ilə."""
    return reader.read_rows(uploaded_file, max_rows=MAX_ROWS, max_upload_bytes=MAX_UPLOAD_BYTES)


# ── Şablon ───────────────────────────────────────────────────────────────────


def template_columns(exam_score_max, question_count=None, question_max=None) -> list:
    """Şablonun sütun başlıqları (oxuyucu bunları tanıyır; sıra sərbəstdir).

    2026-09-14: «Bal» sütunundan sonra S1..S<n> sual sütunları (``question_count``
    verilməsə defolt 5; ``0`` → sual sütunu yoxdur — tək bal rejimi).
    """
    count = questions.QUESTION_COUNT_DEFAULT if question_count is None else int(question_count)
    per_question = questions.QUESTION_MAX_DEFAULT if question_max is None else int(question_max)
    # W6 (2026-09-14): «İmtahan növü» sütunu «Qrup»dan sonra — oxuyucu onu
    # ``exam_kind`` açarı ilə tanıyır (başlıq sətri parser ilə uyğundur).
    return [
        pgettext(_CTX, "Tələbə №"),
        pgettext(_CTX, "FİN"),
        pgettext(_CTX, "Ad Soyad"),
        pgettext(_CTX, "Qrup"),
        pgettext(_CTX_ENTRY, "İmtahan növü"),
        pgettext(_CTX, "Cari bal"),
        "%s (0–%s)" % (pgettext(_CTX, "Bal"), int(exam_score_max)),
    ] + ["%s (0–%s)" % (label, per_question) for label in questions.question_labels(count)]


def exam_kind_label(exam_kind) -> str:
    """``written`` → «Yazılı» (cari dildə); boş / yad dəyər → yazılı (vərəq defoltu)."""
    value = (exam_kind or "").strip() if isinstance(exam_kind, str) else exam_kind
    if value not in ExamScoreSheetKind.values:
        value = ExamScoreSheetKind.WRITTEN
    return str(ExamScoreSheetKind(value).label)


def _kind_aliases() -> dict:
    """``{normallaşdırılmış mətn: növ dəyəri}`` — fayl xanası dəyər, AZ/EN etiket və ya cari dil etiketi ola bilər."""
    fixed = {
        ExamScoreSheetKind.WRITTEN: ("written", "yazili", "yazılı"),
        ExamScoreSheetKind.PRACTICAL: ("practical", "praktiki", "praktik"),
    }
    aliases = {}
    for kind in ExamScoreSheetKind:
        for text in fixed[kind] + (kind.value, str(kind.label)):
            aliases[normalize_header(text)] = kind.value
            aliases[str(text).strip().casefold()] = kind.value
    aliases.pop("", None)
    return aliases


def _kind_mismatch_message(cell_text, exam_kind):
    """Sətrin «İmtahan növü» xanası seçilmiş növə uyğun deyilsə mesaj, əks halda ``None``.

    Boş xana = «vərəqin növü» (köhnə şablonlar, əl ilə hazırlanmış fayl) — xəta
    deyil. Uyğunsuzluq səssiz ötürülmür: operator ya vərəq məlumatlarında növü
    dəyişir, ya xananı düzəldir (bir protokol = bir növ).
    """
    text = str(cell_text or "").strip()
    if not text or not exam_kind:
        return None
    aliases = _kind_aliases()
    resolved = aliases.get(normalize_header(text)) or aliases.get(text.casefold())
    if resolved == exam_kind:
        return None
    return pgettext(_CTX, "Sətirdəki imtahan növü («%(cell)s») vərəqin növü (%(kind)s) ilə uyğun gəlmir.") % {
        "cell": text[:40],
        "kind": exam_kind_label(exam_kind),
    }


def _template_rows(roster, question_count, kind_label="") -> list:
    rows = []
    for row in roster["rows"]:
        student = row["student"]
        rows.append(
            [
                student.username,
                getattr(student.profile, "fin", "") or "",
                student.get_full_name() or student.username,
                service.offering_label(roster["offering"]),
                kind_label,
                _score_text(row.get("exam_score")),
                "",
            ]
            + [""] * int(question_count)
        )
    return rows


def _score_text(value) -> str:
    if value is None:
        return ""
    return str(int(Decimal(value)))


def build_template(*, roster, fmt="xlsx", question_count=None, question_max=None, exam_kind=None):
    """``(bytes, content_type, filename)`` — siyahı ilə doldurulmuş şablon (S1..Sn sütunları ilə).

    W6 (2026-09-14): ``exam_kind`` (``written`` / ``practical``) «İmtahan növü»
    sütununa cari dildə etiket kimi yazılır; XLSX-də həmin sütun siyahı seçimidir.
    """
    offering = roster["offering"]
    count = questions.QUESTION_COUNT_DEFAULT if question_count is None else int(question_count)
    per_question = questions.QUESTION_MAX_DEFAULT if question_max is None else int(question_max)
    kind_label = exam_kind_label(exam_kind)
    headers = template_columns(roster["exam_score_max"], count, per_question)
    rows = [[export_text(value) for value in row] for row in _template_rows(roster, count, kind_label)]
    stem = "imtahan_ballari_%s_%s" % (offering.subject.code or "fenn", service.offering_label(offering))
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", stem.translate(AZ_TRANSLIT)).strip("_") or "imtahan_ballari"
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(rows)
        return ("﻿" + buffer.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", f"{stem}.csv"
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.worksheet.datavalidation import DataValidation
    except Exception:  # pragma: no cover — paket olmayan mühit
        return build_template(
            roster=roster, fmt="csv", question_count=count, question_max=per_question, exam_kind=exam_kind
        )
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column, width in zip("ABCDEFG", (16, 12, 32, 14, 14, 10, 12)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    kind_column, score_column = get_column_letter(_KIND_COLUMN_INDEX), get_column_letter(_SCORE_COLUMN_INDEX)
    if rows:
        # «İmtahan növü» — seçim siyahısı (sahibin növləri: yazılı / praktiki).
        kind_validation = DataValidation(
            type="list",
            formula1='"%s"' % ",".join(str(kind.label) for kind in ExamScoreSheetKind),
            allow_blank=True,
        )
        kind_validation.error = pgettext(_CTX_ENTRY, "İmtahan növü yazılı və ya praktiki olmalıdır.")
        sheet.add_data_validation(kind_validation)
        kind_validation.add(f"{kind_column}2:{kind_column}{len(rows) + 1}")
        validation = DataValidation(
            type="whole",
            operator="between",
            formula1="0",
            formula2=str(int(roster["exam_score_max"])),
            allow_blank=True,
        )
        validation.error = pgettext(_CTX, "Bal 0 ilə maksimum arasında tam ədəd olmalıdır.")
        sheet.add_data_validation(validation)
        validation.add(f"{score_column}2:{score_column}{len(rows) + 1}")
        if count:
            # S1..Sn — hər sual 0..question_max (sahib: «hər sualdan max 10»).
            per_question_validation = DataValidation(
                type="whole", operator="between", formula1="0", formula2=str(per_question), allow_blank=True
            )
            per_question_validation.error = pgettext(_CTX, "Bal 0 ilə maksimum arasında tam ədəd olmalıdır.")
            sheet.add_data_validation(per_question_validation)
            first = get_column_letter(_FIXED_COLUMN_COUNT + 1)
            last = get_column_letter(_FIXED_COLUMN_COUNT + count)
            per_question_validation.add(f"{first}2:{last}{len(rows) + 1}")
    output = io.BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"{stem}.xlsx",
    )


# ── Quru icra (plan) ─────────────────────────────────────────────────────────


def _roster_index(roster) -> dict:
    """Siyahını üç açarla indekslə: istifadəçi adı / institusional id, FİN, normal ad."""
    by_key, by_fin, by_name = {}, {}, {}

    def add_identifier(index, value, row):
        if value in index and index[value] is not row:
            index[value] = None  # ambiguous aliases must never select a student
        else:
            index[value] = row

    for row in roster["rows"]:
        student = row["student"]
        add_identifier(by_key, student.username.lower(), row)
        add_identifier(by_key, export_text(student.username).lower(), row)
        institutional = getattr(student.profile, "institutional_identifier", None)
        if institutional:
            add_identifier(by_key, str(institutional).strip().lower(), row)
        fin = getattr(student.profile, "fin", "")
        if fin:
            add_identifier(by_fin, fin.strip().upper(), row)
        name = normalize_name(student.get_full_name())
        if name:
            by_name.setdefault(name, []).append(row)
    return {"key": by_key, "fin": by_fin, "name": by_name}


def _resolve_student(record, index):
    """``(roster_row | None, xəta kodu, xəbərdarlıq)`` — fayl sətrini siyahıya uyğunlaşdır."""
    key = (record.get("key") or "").strip().lower()
    fin = (record.get("fin") or "").strip().upper()
    by_key = index["key"].get(key) if key else None
    by_fin = index["fin"].get(fin) if fin else None
    if key and fin and (by_key is None or by_fin is None or by_key is not by_fin):
        return None, "mismatch", ""
    if by_key is not None:
        return by_key, "", ""
    if by_fin is not None:
        return by_fin, "", ""
    if key or fin:
        return None, "unknown", ""
    name = normalize_name(record.get("full_name"))
    candidates = index["name"].get(name, []) if name else []
    if len(candidates) == 1:
        return candidates[0], "", pgettext(_CTX, "ada görə uyğunlaşdırıldı")
    if len(candidates) > 1:
        return None, "ambiguous", ""
    return None, "unknown", ""


_ERROR_MESSAGES = {
    "mismatch": lambda: pgettext(_CTX, "Tələbə № və FİN fərqli tələbələrə aiddir."),
    "unknown": lambda: pgettext(_CTX, "Tələbə bu qrupun siyahısında tapılmadı (qeydiyyatı yoxdur)."),
    "ambiguous": lambda: pgettext(_CTX, "Eyni adlı bir neçə tələbə var — Tələbə № və ya FİN yazın."),
    "duplicate": lambda: pgettext(_CTX, "Bu tələbə faylda təkrarlanır."),
}


def _question_list(record) -> list:
    """Sətrin S1..Sn xanaları → sıralı xam siyahı (``[S1, …, Sn]``; aradakı boşluq boş sətirdir)."""
    cells = record.get("questions") or {}
    if not cells:
        return []
    return [cells.get(number, "") for number in range(1, max(cells) + 1)]


def build_plan(*, roster, rows, question_count=None, question_max=None, exam_kind=None) -> list:
    """Quru icra: hər fayl sətri üçün vəziyyət + mesaj. HEÇ NƏ YAZMIR.

    ``roster`` — ``exam_score_entry.roster_for_offering`` nəticəsi (cari ballar
    oradan gəlir, əlavə sorğu yoxdur).

    2026-09-14 (W2 `w2paper`): sətirdə S1..Sn xanası doludursa bal ONLARIN
    CƏMİDİR (``question_scores`` plana düşür); şəbəkə verilməsə fayldakı S
    sütunlarının sayı və defolt tavan (10) götürülür — tətbiqdə servis vərəqin
    öz şəbəkəsi ilə yenidən yoxlayır (uyğunsuzluq sətir xətası kimi görünür).

    2026-09-14 (W6 `w6paper`): view quru icraya da POST şəbəkəsini ötürür — ön
    baxış = tətbiq. ``exam_kind`` verilərsə sətrin «İmtahan növü» xanası onunla
    tutuşdurulur (boş xana sərbəstdir; fərqli növ → sətir xətası).
    """
    index = _roster_index(roster)
    cap = roster["exam_score_max"]
    file_question_count = max((max(record.get("questions") or {0: ""}) for record in rows), default=0)
    grid = {
        "question_count": int(question_count) if question_count is not None else file_question_count,
        "question_max": int(question_max) if question_max is not None else questions.QUESTION_MAX_DEFAULT,
    }
    seen: set = set()
    plan = []
    for record in rows:
        item = {
            "row": record["_row"],
            "key": record.get("key", ""),
            "fin": record.get("fin", ""),
            "full_name": record.get("full_name", ""),
            "raw_score": record.get("score", ""),
            "enrollment_id": "",
            "student": "",
            "username": "",
            "current": None,
            "score": None,
            "question_scores": None,
            "status": STATUS_ERROR,
            "message": "",
            "warning": "",
        }
        plan.append(item)
        roster_row, code, warning = _resolve_student(record, index)
        if roster_row is None:
            item["message"] = _ERROR_MESSAGES[code]()
            continue
        student = roster_row["student"]
        item.update(
            enrollment_id=str(roster_row["enrollment"].id),
            student=student.get_full_name() or student.username,
            username=student.username,
            current=_score_text(roster_row.get("exam_score")) or None,
            warning=warning,
        )
        if item["enrollment_id"] in seen:
            item["message"] = _ERROR_MESSAGES["duplicate"]()
            continue
        seen.add(item["enrollment_id"])
        kind_problem = _kind_mismatch_message(record.get("exam_kind"), exam_kind)
        if kind_problem:
            item["message"] = kind_problem
            continue
        raw = str(record.get("score") or "").strip().replace(",", ".")
        question_list = _question_list(record)
        try:
            if not questions.is_blank_list(question_list):
                cleaned_questions, cleaned = questions.clean_question_scores(question_list, cap=cap, **grid)
                item["question_scores"] = cleaned_questions
            else:
                cleaned = service._clean_score(raw, cap)
        except ValidationError as exc:
            item["message"] = " ".join(exc.messages)
            continue
        if cleaned is None:
            item["status"] = STATUS_SKIP
            item["message"] = pgettext(_CTX, "Bal boşdur — toxunulmur.")
            continue
        item["score"] = str(int(cleaned))
        current = roster_row.get("exam_score")
        same_questions = item["question_scores"] is None or questions.same_question_scores(
            roster_row.get("question_scores") or None, item["question_scores"]
        )
        if current is None:
            item["status"] = STATUS_NEW
        elif Decimal(current) == cleaned and same_questions:
            item["status"] = STATUS_UNCHANGED
            item["message"] = pgettext(_CTX, "Eyni bal artıq yazılıb.")
        else:
            item["status"] = STATUS_CHANGE
            item["message"] = pgettext(_CTX, "Yazılmış bal dəyişir — səbəb, qeyd və sənəd tələb olunur.")
    return plan


def summarize(plan, *, roster=None) -> dict:
    """Vəziyyət sayğacları (+ faylda OLMAYAN siyahı tələbələri — «missing»)."""
    counts = {STATUS_NEW: 0, STATUS_CHANGE: 0, STATUS_UNCHANGED: 0, STATUS_SKIP: 0, STATUS_ERROR: 0}
    for item in plan:
        counts[item["status"]] = counts.get(item["status"], 0) + 1
    summary = {"total": len(plan), **counts}
    summary["writes"] = counts[STATUS_NEW] + counts[STATUS_CHANGE]
    if roster is not None:
        matched = {item["enrollment_id"] for item in plan if item["enrollment_id"]}
        summary["missing"] = sum(1 for row in roster["rows"] if str(row["enrollment"].id) not in matched)
    return summary


def needs_justification(plan) -> bool:
    return any(item["status"] == STATUS_CHANGE for item in plan)


def rows_for_service(plan, *, reason="", note="") -> list:
    """Plandan YALNIZ yazılacaq sətirləri servis formatına çevir."""
    return [
        {
            "enrollment_id": item["enrollment_id"],
            "score": item["score"],
            "question_scores": item.get("question_scores"),
            "reason": reason,
            "note": note,
        }
        for item in plan
        if item["status"] in (STATUS_NEW, STATUS_CHANGE)
    ]


def apply_plan(*, offering, plan, by_user, request=None, sheet=None, reason="", note=""):
    """Planı tətbiq et — yazı YALNIZ servis qatından keçir (``save_roster_scores``).

    Nəticə: ``save_roster_scores`` lüğəti + hər plan sətrinin son vəziyyəti
    (``written`` / ``failed`` işarəsi ilə).
    """
    result = service.save_roster_scores(
        offering=offering,
        rows=rows_for_service(plan, reason=reason, note=note),
        by_user=by_user,
        request=request,
        sheet=sheet,
    )
    failed = result.get("failed_by_enrollment") or {}
    for item in plan:
        if item["status"] not in (STATUS_NEW, STATUS_CHANGE):
            continue
        problem = failed.get(item["enrollment_id"])
        if problem:
            item["status"] = STATUS_ERROR
            item["message"] = problem
        elif item["enrollment_id"] in result["written_ids"]:
            item["message"] = pgettext(_CTX, "Yazıldı.")
            item["written"] = True
        else:
            item["status"] = STATUS_UNCHANGED
            item["message"] = pgettext(_CTX, "Eyni bal artıq yazılıb.")
    result["total"] = len(plan)
    result["failed"] = sum(item["status"] == STATUS_ERROR for item in plan)
    result["skipped"] = len(plan) - result["written"] - result["failed"]
    return result


def enrollment_ids_for(plan) -> set:
    """Plandakı uyğunlaşmış qeydiyyatlar (testlər / yoxlama üçün)."""
    return {item["enrollment_id"] for item in plan if item["enrollment_id"]}


__all__ = [
    "ALLOWED_SUFFIXES",
    "MAX_ROWS",
    "MAX_UPLOAD_BYTES",
    "STATUS_CHANGE",
    "STATUS_ERROR",
    "STATUS_NEW",
    "STATUS_SKIP",
    "STATUS_UNCHANGED",
    "ImportFileError",
    "apply_plan",
    "build_plan",
    "build_template",
    "exam_kind_label",
    "needs_justification",
    "question_count_in",
    "read_rows",
    "rows_for_service",
    "summarize",
    "template_columns",
]
