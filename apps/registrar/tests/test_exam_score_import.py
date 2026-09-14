"""Academic import regressions: identities, input limits, preview and writes."""

import io
from decimal import Decimal
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from apps.accounts.tests import test_exam_score_entry_section as fixtures
from apps.registrar import exam_score_entry as service
from apps.registrar import exam_score_import as importer
from apps.registrar.models import ExamScoreEntry, FinalGrade


@override_settings(UNIVERSITY_MODE=True)
class ScoreImportTest(TestCase):
    setUpTestData = classmethod(fixtures.ExamScoreEntrySectionTest.setUpTestData.__func__)
    _client = fixtures.ExamScoreEntrySectionTest._client

    def _csv(self, score="32", username=None):
        return SimpleUploadedFile(
            "CODEX_TEST_scores.csv", f"username,score\n{username or self.student.username},{score}\n".encode()
        )

    def test_preview_does_not_create_grade_or_entry(self):
        before = (FinalGrade.objects.count(), ExamScoreEntry.objects.count())
        response = self._client(self.center).post(
            reverse("accounts:exam_score_import_preview"),
            {
                "offering_id": self.offering.pk,
                "file": self._csv(),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["summary"]["new"], 1)
        self.assertEqual(before, (FinalGrade.objects.count(), ExamScoreEntry.objects.count()))

    def test_apply_is_idempotent_and_correction_requires_evidence(self):
        client = self._client(self.center)
        url = reverse("accounts:exam_score_import_apply")
        for _ in range(2):
            response = client.post(url, {"offering_id": self.offering.pk, "file": self._csv()})
            self.assertEqual(response.status_code, 200)
        self.assertEqual(ExamScoreEntry.objects.count(), 1)
        response = client.post(url, {"offering_id": self.offering.pk, "file": self._csv("33")})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(FinalGrade.objects.get(enrollment=self.enrollment).exam_score, Decimal(32))

    def test_student_and_teacher_cannot_import(self):
        for user in (self.student, self.teacher):
            for action in ("preview", "apply"):
                response = self._client(user).post(
                    reverse(f"accounts:exam_score_import_{action}"),
                    {
                        "offering_id": self.offering.pk,
                        "file": self._csv(),
                    },
                )
                self.assertEqual(response.status_code, 403)

    def test_foreign_tenant_cannot_read_template(self):
        from apps.organizations.models import Membership, Organization

        org = Organization.objects.create(
            name="CODEX_TEST_foreign",
            slug="codex-foreign",
            owner=self.teacher,
            org_type="university",
            status="active",
            is_active=True,
        )
        Membership.objects.filter(user=self.center).delete()
        Membership.objects.create(
            user=self.center,
            organization=org,
            role=org.roles.get(name="exam_center_head"),
            is_active=True,
            is_primary=True,
        )
        response = self._client(self.center).get(
            reverse("accounts:exam_score_import_template"),
            {
                "offering": self.offering.pk,
            },
        )
        self.assertIn(response.status_code, (403, 404))

    def test_unknown_fin_cannot_override_known_username(self):
        roster = service.roster_for_offering(offering=self.offering)
        plan = importer.build_plan(
            roster=roster,
            rows=[
                {
                    "_row": 2,
                    "key": self.student.username,
                    "fin": "WRONG01",
                    "score": "32",
                }
            ],
        )
        self.assertEqual(plan[0]["status"], importer.STATUS_ERROR)

    def test_duplicate_and_unknown_rows_rejected(self):
        roster = service.roster_for_offering(offering=self.offering)
        rows = [
            {"_row": i, "key": key, "score": "32"}
            for i, key in enumerate([self.student.username, self.student.username, "CODEX_TEST_unknown"], 2)
        ]
        plan = importer.build_plan(roster=roster, rows=rows)
        self.assertEqual([row["status"] for row in plan], ["new", "error", "error"])

    def test_nonfinite_scores_are_validation_errors(self):
        for value in ("NaN", "sNaN", "Infinity", "-Infinity"):
            with self.assertRaises(ValidationError):
                service._clean_score(value, 50)

    def test_row_overflow_is_rejected_instead_of_truncated(self):
        with patch.object(importer, "MAX_ROWS", 1):
            with self.assertRaises(importer.ImportFileError) as caught:
                importer.read_rows(SimpleUploadedFile("CODEX_TEST.csv", b"username,score\na,1\nb,2\n"))
        self.assertEqual(caught.exception.code, "import_too_many_rows")

    def test_malformed_xlsx_rejected(self):
        with self.assertRaises(importer.ImportFileError):
            importer.read_rows(SimpleUploadedFile("CODEX_TEST.xlsx", b"not a zip"))

    def test_xlsx_expansion_limit_rejected(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["username", "score"])
        workbook.active.append([self.student.username, 32])
        buffer = io.BytesIO()
        workbook.save(buffer)
        with patch("apps.registrar.exam_score_import_safety.MAX_EXPANDED_BYTES", 10):
            with self.assertRaises(importer.ImportFileError):
                importer.read_rows(SimpleUploadedFile("CODEX_TEST.xlsx", buffer.getvalue()))

    def test_export_cells_cannot_be_formulas(self):
        from openpyxl import load_workbook

        roster = service.roster_for_offering(offering=self.offering)
        roster["rows"][0]["student"].first_name = '=HYPERLINK("https://example.invalid")'
        data, _, _ = importer.build_template(roster=roster)
        workbook = load_workbook(io.BytesIO(data))
        self.assertNotEqual(workbook.active["C2"].data_type, "f")
        workbook.close()

    def test_zero_score_is_preserved(self):
        rows = importer.read_rows(self._csv("0"))
        plan = importer.build_plan(roster=service.roster_for_offering(offering=self.offering), rows=rows)
        self.assertEqual(plan[0]["score"], "0")

    def test_invalid_offering_is_rejected_without_server_error(self):
        client = self._client(self.center)
        for action in ("preview", "apply"):
            response = client.post(
                reverse(f"accounts:exam_score_import_{action}"),
                {
                    "offering_id": "invalid-uuid",
                    "file": self._csv(),
                },
            )
            self.assertEqual(response.status_code, 404)

    def test_score_sheet_media_requires_matching_structure_scope(self):
        from apps.organizations.models import Membership, Role
        from apps.registrar import exam_score_sheets
        from core.constants import RoleScopeType
        from core.media_policies import check_exam_score_sheet_access

        sheet = exam_score_sheets.create_sheet(offering=self.offering, by_user=self.center)
        path = f"exam_score_sheets/{self.org.pk}/CODEX_TEST.pdf"
        sheet.evidence = path
        sheet.save(update_fields=["evidence"])
        self.assertTrue(check_exam_score_sheet_access(self.center, path))
        self.assertTrue(check_exam_score_sheet_access(self.teacher, path))
        self.assertFalse(check_exam_score_sheet_access(self.student, path))
        restricted = Role.objects.create(
            organization=self.org,
            name="CODEX_TEST_scoped",
            level=80,
            scope_type=RoleScopeType.UNIT,
            permissions=["final_score.entry"],
        )
        Membership.objects.filter(user=self.center, organization=self.org).update(role=restricted, scope_unit=None)
        self.assertFalse(check_exam_score_sheet_access(type(self.center).objects.get(pk=self.center.pk), path))

    def test_cross_field_identifier_collision_is_rejected_in_both_roster_orders(self):
        from types import SimpleNamespace

        first = SimpleNamespace(
            username="shared",
            profile=SimpleNamespace(fin="", institutional_identifier=""),
            get_full_name=lambda: "First student",
        )
        second = SimpleNamespace(
            username="other",
            profile=SimpleNamespace(fin="", institutional_identifier="shared"),
            get_full_name=lambda: "Second student",
        )
        rows = [{"student": first}, {"student": second}]
        for ordered in (rows, rows[::-1]):
            index = importer._roster_index({"rows": ordered})
            resolved, error, _ = importer._resolve_student({"key": "shared"}, index)
            self.assertIsNone(resolved)
            self.assertTrue(error)

    def test_formula_safe_template_identifier_round_trips(self):
        roster = service.roster_for_offering(offering=self.offering)
        roster["rows"][0]["student"].username = "+CODEX_TEST_student"
        for fmt in ("csv", "xlsx"):
            payload, _, _ = importer.build_template(roster=roster, fmt=fmt)
            rows = importer.read_rows(SimpleUploadedFile(f"CODEX_TEST_roundtrip.{fmt}", payload))
            self.assertEqual(rows[0]["key"], "'+CODEX_TEST_student")
            rows[0]["score"] = "32"
            plan = importer.build_plan(roster=roster, rows=rows)
            self.assertEqual(plan[0]["status"], importer.STATUS_NEW, plan)
