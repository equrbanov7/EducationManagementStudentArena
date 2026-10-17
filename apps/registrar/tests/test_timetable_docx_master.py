"""Magistratura dərs cədvəli formatı (sahib 2026-10-17, «Magistratura şöbəsi» I/II kurs).

«213/6BM ing (1)» başlıqları (qəbul kodu, tələbə sayı, «(ing)», «intensiv»), «QRUPLAR», axşam saatları,
«Fənn (45 saat) [dərəcə] Ad Soyad» hüceyrələri (soyad əvvəl də ola bilər), bir neçə qrupun «aud» xanasını
örtən BİRGƏ dərs və tam adla müəllim tutuşdurması (hər iki sıra, diakritikasız, çox mənalı hal)."""

from __future__ import annotations

import datetime
import pathlib
import tempfile
from types import SimpleNamespace

from django.test import SimpleTestCase

import docx

from apps.registrar.schedule_docx_match import Resolver, default_kind, is_full_name, label_key, teacher_matches
from apps.registrar.schedule_docx_parse import parse_group_label, parse_subject_text, parse_tables


class MasterGroupLabelTests(SimpleTestCase):
    def test_headers(self):
        expected = {
            "213/6BM ing (1)": ("213/6", "BM", "ing", ""),
            "206/6OT (1)": ("206/6", "OT", "az", ""),
            "409/6MBA (11)": ("409/6", "MBA", "az", ""),
            "213/5 BM (1) intensiv": ("213/5", "BM", "az", ""),
            "209/5 SP (10)": ("209/5", "SP", "az", ""),
            "709/5 Sbio": ("709/5", "Sbio", "az", ""),
            "209/5 KP (ing) (3)": ("209/5", "KP", "ing", ""),
            "803/5 T ing (1) intensiv": ("803/5", "T", "ing", ""),
            "631/5 K ing (7)": ("631/5", "K", "ing", ""),
            "403/6MN (1) ing": ("403/6", "MN", "ing", ""),
        }
        for raw, values in expected.items():
            label = parse_group_label(raw)
            self.assertIsNotNone(label, raw)
            self.assertEqual((label.digits, label.code, label.sector, label.sub), values, raw)

    def test_prod_group_names_match_headers_regardless_of_spaces_and_slash_spacing(self):
        for header, name in (
            ("631/5 K (8)", "631/5K"),
            ("631/5 K ing (7)", "631/5K ing"),
            ("505/6MB (10)", "505/6 MB"),
            ("709/5 Sbio", "709/5 SBİO"),
            ("632/6İ (10)", "632 / 6 İ"),
        ):
            self.assertEqual(label_key(parse_group_label(header)), label_key(parse_group_label(name)), header)
        self.assertNotEqual(label_key(parse_group_label("631/5 K")), label_key(parse_group_label("631/6 K")))
        self.assertNotEqual(label_key(parse_group_label("631/5 K")), label_key(parse_group_label("631/5 K ing")))


class MasterCellTests(SimpleTestCase):
    def test_cells(self):
        cases = {
            "Tədqiqat metodları (45 saat) t.ü.f.d. Günel Rəhimli": ("Tədqiqat metodları", 45, "Günel Rəhimli"),
            "AMTSF: Ailə psixoterapiyası (45 saat) p.ü.f.d.,dos. Afərin Abbasova": (
                "Ailə psixoterapiyası",
                45,
                "Afərin Abbasova",
            ),
            "Investisiya portfelinin idarə edilməsi (45 saat) Qasımzadə Nail": (
                "Investisiya portfelinin idarə edilməsi",
                45,
                "Qasımzadə Nail",
            ),
            "Gen və zülal mühəndisliyi (30 saat) / Firəngiz Azadəliyeva": (
                "Gen və zülal mühəndisliyi",
                30,
                "Firəngiz Azadəliyeva",
            ),
            "Beynəlxalq münasibətlərdə siyasi mədəniyyət (45 saat intensiv) Murad Əsgərov": (
                "Beynəlxalq münasibətlərdə siyasi mədəniyyət",
                45,
                "Murad Əsgərov",
            ),
            "Maliyyə nəzarəti (15 saat) intensiv Aydan Hüseynli": ("Maliyyə nəzarəti", 15, "Aydan Hüseynli"),
            "Beynəlxalq təşkilatlar (15saat ) Səbinə Manafova": ("Beynəlxalq təşkilatlar", 15, "Səbinə Manafova"),
            "Beynəlxalq turizm resursları (45 saat Hüseyn İdriszadə": (
                "Beynəlxalq turizm resursları",
                45,
                "Hüseyn İdriszadə",
            ),
            "Genetik məlumat (60 saəat) Firəngiz Azadəliyeva": ("Genetik məlumat", 60, "Firəngiz Azadəliyeva"),
            "Bioekologiyanın müasir problemləri(45 saat) k.e.n Cavadova Leyla": (
                "Bioekologiyanın müasir problemləri",
                45,
                "Cavadova Leyla",
            ),
            "Ağqoyunlu dövləti ( XV əsr) (45 saat) t.ü.f.d. Günay Həsənova": (
                "Ağqoyunlu dövləti ( XV əsr)",
                45,
                "Günay Həsənova",
            ),
            "Maliyyə auditi i.ü.f.d Əşrəf Həsənov": ("Maliyyə auditi", 0, "Əşrəf Həsənov"),
            "Sitogenetika (45 saat)": ("Sitogenetika", 45, ""),
        }
        for raw, (subject, hours, teacher) in cases.items():
            with self.subTest(raw=raw):
                parsed = parse_subject_text(raw)
                self.assertEqual((parsed.subject, parsed.hours, parsed.teacher), (subject, hours, teacher))
                self.assertEqual(parsed.kind, "")  # «mühəndisliyi» növ işarəsi deyil
                self.assertEqual(parsed.teacher_full, bool(teacher))
        self.assertTrue(parse_subject_text("Maliyyə nəzarəti (15 saat) intensiv Aydan Hüseynli").intensive)
        # Bakalavr hüceyrəsi dəyişmir.
        old = parse_subject_text("Kompüter şəbəkələri müh. Nobari S.")
        self.assertEqual(
            (old.subject, old.kind, old.teacher, old.hours, old.teacher_full),
            ("Kompüter şəbəkələri", "lecture", "Nobari S.", 0, False),
        )


def _document(path):
    document = docx.Document()
    header = ["QRUPLAR", "QRUPLAR", "", "632/6İ (10)", "", "", "631/6K (9)", ""]
    hours = ["Saatlar", "Saatlar", "aud", "Fənlər", "aud", "aud", "Fənlər", "Fənlər"]
    table = document.add_table(rows=4, cols=8)
    for col, text in enumerate(header):
        table.cell(0, col).text = text
    for col, text in enumerate(hours):
        table.cell(1, col).text = text
    for row, (day, time) in ((2, ("Bazar ertəsi", "18:40-20:00")), (3, ("Bazar ertəsi", "20:10-21:30"))):
        table.cell(row, 0).text, table.cell(row, 1).text = day, time
    table.cell(0, 0).merge(table.cell(0, 1)).text = "QRUPLAR"
    table.cell(0, 4).merge(table.cell(0, 5)).text = ""
    table.cell(0, 6).merge(table.cell(0, 7)).text = "631/6K (9)"
    # 2 — iki qrupun BİRGƏ dərsi: İ fənni + K-nın «aud» xanaları + K fənni birləşib; otaq soldakı «38».
    table.cell(2, 2).text = "38"
    table.cell(2, 3).merge(table.cell(2, 7)).text = "Tədqiqat metodları (45 saat) Qardyni Şokrollah"
    # 3 — İ-nin dərsi K-nın 1-ci «aud» xanasını örtür; K-nın öz otağı «206» (fənn sütununa yayılıb).
    table.cell(3, 2).text = "38"
    table.cell(3, 3).merge(table.cell(3, 4)).text = "İnformasiya texnologiyalarının tarixi (30 saat) Qardyni Şokrollah"
    table.cell(3, 5).merge(table.cell(3, 6)).text = "206"
    table.cell(3, 7).text = "Kompüter mühəndisliyinin tarixi (30 saat) t.ü.f.d.,dos. Qoşqar Əliyev"
    document.save(str(path))


class MasterTableTests(SimpleTestCase):
    def test_shared_lesson_and_rooms(self):
        with tempfile.TemporaryDirectory() as folder:
            path = pathlib.Path(folder) / "mag.docx"
            _document(path)
            result = parse_tables(docx.Document(str(path)).tables)
        self.assertEqual(result.skipped, [])
        self.assertEqual(result.groups, ["632/6 İ az", "631/6 K az"])
        rows = {(e.row, e.group.display, e.start, e.room, e.text.subject, e.text.teacher) for e in result.entries}
        evening = datetime.time(18, 40)
        self.assertIn((2, "632/6 İ az", evening, "38", "Tədqiqat metodları", "Qardyni Şokrollah"), rows)
        self.assertIn((2, "631/6 K az", evening, "38", "Tədqiqat metodları", "Qardyni Şokrollah"), rows)
        late = datetime.time(20, 10)
        self.assertIn((3, "632/6 İ az", late, "38", "İnformasiya texnologiyalarının tarixi", "Qardyni Şokrollah"), rows)
        self.assertIn((3, "631/6 K az", late, "206", "Kompüter mühəndisliyinin tarixi", "Qoşqar Əliyev"), rows)
        self.assertEqual(len(result.entries), 4)


def _user(pk, first, last):
    return SimpleNamespace(pk=pk, first_name=first, last_name=last)


class FullNameTeacherTests(SimpleTestCase):
    def test_full_name_either_order_diacritic_tolerant(self):
        users = [_user(1, "Günel", "Rəhimli"), _user(2, "Nail", "Qasımzadə"), _user(3, "Günay", "Həsənova")]
        self.assertTrue(is_full_name("Günel Rəhimli"))
        self.assertFalse(is_full_name("Rəhimli G."))
        self.assertEqual([u.pk for u in teacher_matches(users, "Günel Rəhimli", fuzzy=False)], [1])
        self.assertEqual([u.pk for u in teacher_matches(users, "Qasımzadə Nail", fuzzy=False)], [2])
        self.assertEqual([u.pk for u in teacher_matches(users, "Gunel Rehimli", fuzzy=False)], [1])
        self.assertEqual(teacher_matches(users, "Günel Rəhimova", fuzzy=False), [])
        # Kiçik yazı xətası («Qardyni» ↔ «Qardyani») yalnız ``fuzzy`` ilə.
        pool = [_user(4, "Şokrollah", "Qardyani")]
        self.assertEqual(teacher_matches(pool, "Qardyni Şokrollah", fuzzy=False), [])
        self.assertEqual([u.pk for u in teacher_matches(pool, "Qardyni Şokrollah", fuzzy=True)], [4])
        # Köhnə «Soyad İ.» forması dəyişmir.
        self.assertEqual([u.pk for u in teacher_matches(users, "Həsənova G.", fuzzy=False)], [3])

    def test_ambiguous_full_name_returns_every_account(self):
        users = [_user(1, "Rail", "Səfiyev"), _user(2, "Rail", "Səfiyev")]
        self.assertEqual(len(teacher_matches(users, "Səfiyev Rail", fuzzy=False)), 2)


class DefaultKindTests(SimpleTestCase):
    """2026-10-17 rəyi: magistratura hüceyrəsinin defolt mühazirəsi yalnız planda mühazirə VARSA (və ya plan
    məlum deyilsə) — workload tapşırıq sətri ilə eyni qayda (``default_kind``)."""

    def _resolver(self, hours):
        resolver = Resolver.__new__(Resolver)
        resolver._hours = {1: hours}
        return resolver

    def test_default_lecture_requires_lecture_in_planned_kinds(self):
        self.assertEqual(default_kind([], "lecture"), ("lecture", "kind_default_lecture"))
        self.assertEqual(default_kind(["lecture", "seminar"], "lecture"), ("lecture", "kind_default_lecture"))
        self.assertEqual(default_kind(["seminar", "lab"], "lecture"), ("", "kind_unknown"))
        self.assertEqual(default_kind(["lecture"], ""), ("", "kind_unknown"))
        offering = SimpleNamespace(pk=1)
        infer = Resolver.infer_kind
        self.assertEqual(
            infer(self._resolver({"seminar": 30, "lab": 15}), offering, None, default="lecture"), ("", "kind_unknown")
        )
        self.assertEqual(
            infer(self._resolver({"lab": 15}), offering, None, default="lecture"), ("lab", "kind_from_hours")
        )
        self.assertEqual(
            infer(self._resolver({"lecture": 30, "seminar": 15}), offering, None, default="lecture"),
            ("lecture", "kind_default_lecture"),
        )
        self.assertEqual(
            infer(self._resolver({}), offering, None, default="lecture"), ("lecture", "kind_default_lecture")
        )
