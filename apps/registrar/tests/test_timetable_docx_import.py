"""Word dərs cədvəli idxalı (``import_timetable_docx``) — sahib 2026-10-12.

Sintetik .docx (python-docx ilə test içində qurulur) real sənədin çətinliklərini daşıyır: «Qruplar» /
«Saatlar» başlıqları, başlığı boş «aud» sütunları, alt qrup sütunları («902 KE-1», «902 KE-2»), üfüqi
və şaquli birləşmələr, eyni saatın iki sətri (üst/alt həftə), «müh./sem./lab/sem» növləri, «alt»
işarəsi, TƏCRÜBƏ/LEVEL, tanınmayan fənn, növsüz hüceyrə və müəllim toqquşması.
"""

from __future__ import annotations

import datetime
import io
import pathlib
import tempfile

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase
from django.utils import timezone

import docx

from apps.audit.models import AuditLog
from apps.organizations.models import AcademicPeriod, Membership, Organization, OrgUnit
from apps.registrar import schedule, services
from apps.registrar.models import (
    AcademicStatus,
    Curriculum,
    Enrollment,
    Program,
    ScheduleSlot,
    SlotKind,
    StudentAcademicRecord,
    Subject,
    WeekType,
)
from apps.registrar.schedule_docx_import import build_plan, run_import
from apps.registrar.schedule_docx_parse import parse_docx, parse_group_label, parse_subject_text, parse_tables
from core.constants import AcademicPeriodType, OrganizationType, OrgUnitType
from core.rls import bypass_rls

User = get_user_model()

# Sütunlar: 0 gün · 1 saat · 2 aud(K) · 3 K · 4 aud(İ ing) · 5 İ ing · 6 aud(KE) · 7 KE-1 · 8 KE-2
HEADER = ["Qruplar", "Qruplar", "", "III KURS  901 K 20", "", "III KURS 901 İ ing 10", "", "II KURS 902 KE 30", ""]
HOURS = ["Saatlar", "Saatlar", "aud", "Fənlər", "aud", "Fənlər", "", "902 KE-1  15", "902 KE-2 15"]
ROWS = [
    # 2 — tək sətir: hamısı hər həftə; KE alt qruplarında EYNİ dərs (üfüqi birləşmiş) → qrupun özü
    [
        "Bazar ertəsi",
        "8:30-\n10:00",
        "11",
        "Verilənlər bazası müh.\nƏliyeva A.",
        "12",
        "Şəbəkələr sem. Həsənov B.",
        "13",
        "Fizika lab/sem\nKərimov C.",
        None,
    ],
    # 3/4 — eyni saat iki sətir (üst/alt): K fərqli dərslər; İ ing şaquli birləşmiş (hər həftə); KE alt qruplar
    [
        "Bazar ertəsi",
        "10:10-11:40",
        "21",
        "Verilənlər bazası sem.\nƏliyeva A.",
        "22",
        "Şəbəkələr müh. Həsənov B.",
        "",
        "Proqramlaşdırma sem.\nKərimov C.",
        "",
    ],
    [
        "Bazar ertəsi",
        "10:10-11:40",
        "21",
        "Şəbəkələr müh.\nHəsənov B.",
        "22",
        None,
        "",
        "",
        "Proqramlaşdırma sem.\nKərimov C.",
    ],
    # 5 — tanınmayan fənn, təcrübə, LEVEL
    ["Çərşənbə axşamı", "11:50-13:20", "31", "Astronomiya müh.\nƏliyeva A.", "", "TƏCRÜBƏ", "", "LEVEL", "LEVEL"],
    # 6 — «alt» işarəsi (tək sətir → alt həftə) və növsüz hüceyrə
    ["Cümə", "8:30-10:00", "41", "Verilənlər bazası sem alt\nƏliyeva A.", "42", "Şəbəkələr\nHəsənov B.", "", "", ""],
]


def build_docx(path):
    document = docx.Document()
    document.add_paragraph("Əsas dərs cədvəli")
    table = document.add_table(rows=2 + len(ROWS), cols=9)
    for col, text in enumerate(HEADER):
        table.cell(0, col).text = text
    for col, text in enumerate(HOURS):
        table.cell(1, col).text = text
    for offset, row in enumerate(ROWS, start=2):
        for col, text in enumerate(row):
            if text is not None:
                table.cell(offset, col).text = text
    # üfüqi birləşmələr: «Qruplar»/«Saatlar» (0–1), KE başlığı (7–8), KE-nin 2-ci sətri (7–8)
    for row, left, right, text in ((0, 0, 1, "Qruplar"), (1, 0, 1, "Saatlar"), (0, 7, 8, "II KURS 902 KE 30")):
        table.cell(row, left).merge(table.cell(row, right)).text = text
    table.cell(2, 7).merge(table.cell(2, 8)).text = "Fizika lab/sem\nKərimov C."
    # şaquli birləşmələr: gün (3–4), saat (3–4), K-nın otağı (3–4), İ ing dərsi (3–4)
    for col, text in ((0, "Bazar ertəsi"), (1, "10:10-11:40"), (2, "21"), (5, "Şəbəkələr müh. Həsənov B.")):
        table.cell(3, col).merge(table.cell(4, col)).text = text
    document.save(str(path))


class ParsingTests(SimpleTestCase):
    def test_subject_cells(self):
        cases = {
            "Kompüter şəbəkələri müh. Nobari S.": ("Kompüter şəbəkələri", "lecture", "Nobari S."),
            "İnformasiya təhlükəsizliyi (müh.) \nXancanova R.": (
                "İnformasiya təhlükəsizliyi",
                "lecture",
                "Xancanova R.",
            ),
            "Riyazi Məntiq məsələləri muh\nƏzizli N\n": ("Riyazi Məntiq məsələləri", "lecture", "Əzizli N."),
            "Zərərverici proqram vasitələrinin təhlili sem. .Məmmədova L": (
                "Zərərverici proqram vasitələrinin təhlili",
                "seminar",
                "Məmmədova L.",
            ),
            "Kompüter şəbəkələri müh.\nƏzizov  a S.": ("Kompüter şəbəkələri", "lecture", "Əzizova S."),
            "Kompüter mühəndisliyinin əsasları müh\nƏzizova S.": (
                "Kompüter mühəndisliyinin əsasları",
                "lecture",
                "Əzizova S.",
            ),
            "Fizika l+sem.\nNəbiyeva S.": ("Fizika", "lab", "Nəbiyeva S."),
            "Mülki müdafiə müh    .\nƏmrahov E.": ("Mülki müdafiə", "lecture", "Əmrahov E."),
            "Etika Və estetika\nHəsənli R.": ("Etika Və estetika", "", "Həsənli R."),
            "Sabah Hub sem.": ("Sabah Hub", "seminar", ""),
        }
        for raw, (subject, kind, teacher) in cases.items():
            with self.subTest(raw=raw):
                parsed = parse_subject_text(raw)
                self.assertEqual((parsed.subject, parsed.kind, parsed.teacher), (subject, kind, teacher))
        self.assertEqual(parse_subject_text("Psixologiya müh/sem.\nBayramova S.").kind_note, "combined")
        self.assertEqual(parse_subject_text("Kompüter şəbəkələri sem alt \nƏzizova S.").week_marker, "even")
        self.assertEqual(parse_subject_text("TƏCRÜBƏ").special, "practice")
        self.assertEqual(parse_subject_text("Təcrubə").special, "practice")
        self.assertEqual(parse_subject_text("LEVEL").special, "level")

    def test_group_labels(self):
        expected = {
            "IV KURS  233 İTex  ing 12": ("233", "İTex", "ing", ""),
            "IV KURS  233 KE - 13": ("233", "KE", "az", ""),
            "234 K-1    20": ("234", "K", "az", "1"),
            "I KURS  236 KEaz-40": ("236", "KE", "az", ""),
            "II KURS  235 KE az-20+19": ("235", "KE", "az", ""),
            "235 K İNG": ("235", "K", "ing", ""),
            "235 İT-1 az": ("235", "İT", "az", "1"),
        }
        for raw, values in expected.items():
            label = parse_group_label(raw)
            self.assertEqual((label.digits, label.code, label.sector, label.sub), values, raw)

    def test_table_structure(self):
        with tempfile.TemporaryDirectory() as folder:
            path = pathlib.Path(folder) / "t.docx"
            build_docx(path)
            result = parse_tables(docx.Document(str(path)).tables)
        rows = {(e.row, e.group.display, e.week_type, e.text.subject, e.text.kind, e.room) for e in result.entries}
        self.assertIn((2, "902 KE az", "all", "Fizika", "lab", "13"), rows)  # birləşmiş alt qrup → qrupun özü
        self.assertIn((3, "901 K az", "odd", "Verilənlər bazası", "seminar", "21"), rows)
        self.assertIn((4, "901 K az", "even", "Şəbəkələr", "lecture", "21"), rows)
        self.assertIn((3, "901 İ ing", "all", "Şəbəkələr", "lecture", "22"), rows)  # şaquli birləşmə → bir dəfə
        self.assertIn((3, "902 KE-1 az", "odd", "Proqramlaşdırma", "seminar", ""), rows)
        self.assertIn((4, "902 KE-2 az", "even", "Proqramlaşdırma", "seminar", ""), rows)
        self.assertIn((6, "901 K az", "even", "Verilənlər bazası", "seminar", "41"), rows)  # «alt» işarəsi
        ing_rows = [e for e in result.entries if e.group.display == "901 İ ing" and e.weekday == 1 and e.row in (3, 4)]
        self.assertEqual(len(ing_rows), 1)
        self.assertEqual(len(result.entries), 13)


class TimetableImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        today = timezone.localdate()
        cls.owner = User.objects.create_user("ttd_owner", "ttd_owner@qku.edu.az", "pw")
        with bypass_rls():
            cls.org = Organization.objects.create(
                name="TTD Univ",
                slug="ttd-univ",
                org_type=OrganizationType.UNIVERSITY,
                owner=cls.owner,
                status="active",
                is_active=True,
            )
            spec = OrgUnit.objects.create(
                organization=cls.org, name="İxtisas", slug="ttd-spec", unit_type=OrgUnitType.SPECIALTY
            )
            cls.groups = {
                name: OrgUnit.objects.create(
                    organization=cls.org,
                    parent=spec,
                    name=name,
                    slug=f"ttd-{index}",
                    unit_type=OrgUnitType.GROUP,
                )
                for index, name in enumerate(("901 K az", "901 İ ing", "902 KE", "902 KE-1", "902 KE-2"))
            }
            cls.period = AcademicPeriod.objects.create(
                organization=cls.org,
                name="Payız",
                period_type=AcademicPeriodType.SEMESTER,
                academic_year="2026/2027",
                start_date=today - datetime.timedelta(days=20),
                end_date=today + datetime.timedelta(days=100),
                is_current=True,
            )
            subjects = {
                code: Subject.objects.create(organization=cls.org, code=code, name=name)
                for code, name in (
                    ("VB1", "Verilənlərin bazası"),
                    ("SB1", "Şəbəkələr"),
                    ("FZ1", "Fizika"),
                    ("PR1", "Proqramlaşdırma"),
                )
            }
            cls.aliyeva = cls._teacher("a", "Aytən", "Əliyeva")
            cls.hasanov = cls._teacher("b", "Bəhruz", "Həsənov")
            cls.karimov = cls._teacher("c", "Cavid", "Kərimov")
            g = cls.groups
            cls.k_vb = cls._offering(subjects["VB1"], g["901 K az"], cls.aliyeva)
            cls.k_sb = cls._offering(subjects["SB1"], g["901 K az"], None)
            cls.i_sb = cls._offering(subjects["SB1"], g["901 İ ing"], cls.hasanov)
            cls.ke_fz = cls._offering(subjects["FZ1"], g["902 KE"], cls.karimov)
            cls.ke1_pr = cls._offering(subjects["PR1"], g["902 KE-1"], cls.karimov)
            cls.ke2_pr = cls._offering(subjects["PR1"], g["902 KE-2"], None)
            program = Program.objects.create(organization=cls.org, code="TTD-CS", name="Kompüter elmləri")
            curriculum = Curriculum.objects.create(organization=cls.org, program=program, admission_year=2025)
            cls.student = User.objects.create_user("ttd_student", "ttd_s@qku.edu.az", "pw")
            Membership.objects.create(
                user=cls.student,
                organization=cls.org,
                role=cls.org.roles.get(name="student"),
                scope_unit=g["902 KE-1"],
                is_primary=True,
                is_active=True,
            )
            StudentAcademicRecord.objects.create(
                organization=cls.org,
                student=cls.student,
                program=program,
                curriculum=curriculum,
                group=g["902 KE-1"],
                admission_year=2025,
                status=AcademicStatus.ENROLLED,
            )
            # «alt qrupdan əlavə» — tələbə 902 KE-1-dədir, Fizika isə birləşik 902 KE-nin açılışındadır
            Enrollment.objects.create(
                organization=cls.org, student=cls.student, offering=cls.ke_fz, source_group=g["902 KE-1"]
            )

    @classmethod
    def _teacher(cls, suffix, first, last):
        user = User.objects.create_user(
            f"ttd_t{suffix}", f"ttd_t{suffix}@qku.edu.az", "pw", first_name=first, last_name=last
        )
        Membership.objects.create(
            user=user, organization=cls.org, role=cls.org.roles.get(name="teacher"), is_primary=True, is_active=True
        )
        return user

    @classmethod
    def _offering(cls, subject, group, instructor):
        offering = services.get_or_create_offering(
            organization=cls.org, subject=subject, period=cls.period, group=group
        )
        offering.instructor = instructor
        offering.save(update_fields=["instructor"])
        return offering

    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.path = pathlib.Path(self.folder.name) / "cedvel.docx"
        build_docx(self.path)
        self.parsed = parse_docx(self.path)

    def tearDown(self):
        self.folder.cleanup()

    def _run(self, **kwargs):
        with bypass_rls():
            return run_import(organization=self.org, period=self.period, parsed=self.parsed, actor=self.owner, **kwargs)

    def _live(self):
        with bypass_rls():
            return list(ScheduleSlot.objects.filter(organization=self.org).select_related("offering"))

    def test_dry_run_writes_nothing_and_classifies_every_cell(self):
        plan = self._run()
        counts = plan.counts
        self.assertEqual(self._live(), [])
        self.assertEqual(counts["create"], 8)
        self.assertEqual(counts["clash"], 1)
        self.assertEqual(counts["unmatched"], 1)
        self.assertEqual(counts["ambiguous"], 1)
        self.assertEqual(counts["skipped"], 2)
        reasons = {reason for row in plan.rows for reason in row.reasons}
        self.assertTrue({"clash_teacher", "subject_not_in_group", "kind_unknown", "practice", "level"} <= reasons)

    def test_apply_creates_slots_with_kind_week_room_and_instructor(self):
        plan = self._run(apply=True)
        self.assertEqual(len(plan.created_ids), 8)
        slots = {(s.offering_id, s.weekday, s.start_time.strftime("%H:%M"), s.week_type): s for s in self._live()}
        fz = slots[(self.ke_fz.pk, 1, "08:30", WeekType.ALL)]
        self.assertEqual((fz.kind, fz.room, fz.instructor_id), (SlotKind.LAB, "13", None))  # jurnal sahibi → NULL
        self.assertEqual(slots[(self.k_vb.pk, 1, "10:10", WeekType.ODD)].kind, SlotKind.SEMINAR)
        even = slots[(self.k_sb.pk, 1, "10:10", WeekType.EVEN)]
        self.assertEqual((even.kind, even.instructor_id), (SlotKind.LECTURE, self.hasanov.pk))  # slot müəllimi
        self.assertIn((self.ke1_pr.pk, 1, "10:10", WeekType.ODD), slots)
        self.assertEqual(slots[(self.ke2_pr.pk, 1, "10:10", WeekType.EVEN)].instructor_id, self.karimov.pk)
        self.assertIn((self.k_vb.pk, 5, "08:30", WeekType.EVEN), slots)  # «alt» işarəsi
        self.assertNotIn((self.i_sb.pk, 1, "10:10", WeekType.ALL), slots)  # müəllim toqquşması yazılmadı
        with bypass_rls():
            self.assertTrue(AuditLog.objects.filter(resource_type="registrar.ScheduleSlot").exists())

    def test_second_run_is_idempotent(self):
        self._run(apply=True)
        plan = self._run(apply=True)
        self.assertEqual(plan.counts["create"], 0)
        self.assertEqual(plan.counts["exists"], 8)
        self.assertEqual(len(self._live()), 8)

    def test_slots_reach_teacher_and_student_schedules(self):
        self._run(apply=True)
        with bypass_rls():
            teacher_slots = schedule.get_teacher_schedule(
                organization=self.org, teacher=self.hasanov, period=self.period
            )
            student_slots = schedule.get_student_schedule(self.org, self.student, self.groups["902 KE-1"], self.period)
        self.assertIn(self.k_sb.pk, {slot.offering_id for slot in teacher_slots})
        self.assertIn(self.i_sb.pk, {slot.offering_id for slot in teacher_slots})
        # öz alt qrupunun dərsi + birləşik qrupun «alt qrupdan əlavə» Fizikası
        self.assertEqual({slot.offering_id for slot in student_slots}, {self.ke1_pr.pk, self.ke_fz.pk})

    def test_existing_conflict_is_reported_not_overwritten(self):
        with bypass_rls():
            ScheduleSlot.objects.create(
                organization=self.org,
                offering=self.k_sb,
                weekday=1,
                start_time=datetime.time(8, 30),
                end_time=datetime.time(10, 0),
                room="99",
            )
        plan = self._run(apply=True)
        clashes = [row for row in plan.rows if row.action == "clash" and "clash_group" in row.reasons]
        self.assertEqual(len(clashes), 1)
        self.assertEqual(clashes[0].offering, self.k_vb)

    def test_replace_group_slots_soft_deletes_other_slots_of_file_groups(self):
        with bypass_rls():
            old = ScheduleSlot.objects.create(
                organization=self.org,
                offering=self.k_vb,
                weekday=3,
                start_time=datetime.time(15, 15),
                end_time=datetime.time(16, 45),
            )
        plan = self._run(apply=True, replace_group_slots=True)
        self.assertEqual(plan.replaced_ids, [old.pk])
        with bypass_rls():
            self.assertTrue(ScheduleSlot.all_objects.get(pk=old.pk).is_deleted)
        again = self._run(apply=True, replace_group_slots=True)
        self.assertEqual((again.counts["create"], again.replaced_ids), (0, []))

    def test_command_dry_run_writes_report_and_prints_counts_only(self):
        out = io.StringIO()
        with tempfile.TemporaryDirectory() as report:
            call_command(
                "import_timetable_docx",
                "--file",
                str(self.path),
                "--org",
                self.org.slug,
                "--period",
                "Payız 2026/2027",
                "--report",
                report,
                stdout=out,
            )
            csv_text = (pathlib.Path(report) / "timetable_import.csv").read_text(encoding="utf-8")
        text = out.getvalue()
        self.assertIn("DRY-RUN", text)
        self.assertIn("create     8", text)
        self.assertNotIn("Həsənov", text)
        self.assertNotIn(self.hasanov.username, text)
        self.assertIn("ttd_tb", csv_text)  # tam hesabat yalnız CSV-də
        self.assertEqual(self._live(), [])

    def test_build_plan_is_read_only(self):
        with bypass_rls():
            build_plan(organization=self.org, period=self.period, parsed=self.parsed, replace_group_slots=True)
        self.assertEqual(self._live(), [])
