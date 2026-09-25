"""«Sorğu nəticələri» ixracı — yalnız aqreqatlar, səbətli saylar, audit jurnalı, icazə."""

from __future__ import annotations

import csv
import io

from django.test import TestCase
from django.urls import reverse

from apps.audit.models import AuditLog
from core.constants import AuditAction

from .factories import client_for
from .results_world import SUGGESTIONS, build_results_world


def _csv(response):
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


class ResultsExportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.w = build_results_world("svex")

    def _get(self, user, fmt, query=""):
        return client_for(self.w["org"], user).get(reverse("surveys:results_export", args=[fmt]) + query)

    def test_xlsx_contains_every_aggregate_sheet_and_is_audited(self):
        from openpyxl import load_workbook

        before = AuditLog.objects.filter(action=AuditAction.EXPORT, resource_type="surveys.results").count()
        response = self._get(self.w["rector"], "xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(book.sheetnames), 7)
        cells = {str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row}
        for text in SUGGESTIONS[:5]:
            self.assertNotIn(text, cells)  # şərh/təklif mətni ixrac olunmur
        self.assertNotIn("Mövzuları aydın izah edir 0", cells)
        log = AuditLog.objects.filter(action=AuditAction.EXPORT, resource_type="surveys.results").latest("created_at")
        self.assertEqual(
            AuditLog.objects.filter(action=AuditAction.EXPORT, resource_type="surveys.results").count(), before + 1
        )
        self.assertEqual(log.new_values["format"], "xlsx")
        self.assertEqual(log.user, self.w["rector"])

    def test_csv_teachers_are_bucketed_and_hidden_rows_have_no_numbers(self):
        rows = _csv(self._get(self.w["rector"], "csv", "?dataset=teachers"))
        header, body = rows[0], rows[1:]
        n_col, status_col, overall_col = (
            header.index("Cavab sayı"),
            header.index("Vəziyyət"),
            header.index("Orta ümumi bal (1–10)"),
        )
        by_name = {row[1]: row for row in body}
        self.assertEqual(by_name[self.w["teacher_a"].username][n_col], "5+")  # 6 cavab → «5+»
        hidden = by_name[self.w["teacher_e"].username]
        self.assertEqual((hidden[n_col], hidden[overall_col]), ("", ""))
        self.assertTrue(hidden[status_col].startswith("gizli"))
        self.assertEqual({row[n_col] for row in body} - {"", "<5", "5+"}, set())

    def test_questions_export_has_shares_not_raw_counts(self):
        rows = _csv(self._get(self.w["rector"], "csv", "?dataset=questions"))
        self.assertIn("1 (%)", rows[0])
        self.assertNotIn("Cavab sayı", rows[0])

    def test_filters_are_recorded_in_the_audit_entry(self):
        response = self._get(self.w["rector"], "csv", f"?dataset=departments&er_faculty={self.w['faculty'].pk}")
        self.assertEqual(response.status_code, 200)
        log = AuditLog.objects.filter(resource_type="surveys.results").latest("created_at")
        self.assertEqual(log.new_values["filters"].get("er_faculty"), str(self.w["faculty"].pk))
        self.assertEqual(log.new_values["dataset"], "departments")

    def test_permission_and_unknown_dataset(self):
        self.assertEqual(self._get(self.w["dean"], "csv", "?dataset=teachers").status_code, 403)
        self.assertEqual(self._get(self.w["rector"], "csv", "?dataset=raw").status_code, 404)
        self.assertEqual(self._get(self.w["rector"], "pdf").status_code, 404)
