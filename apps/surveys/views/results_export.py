"""«Sorğu nəticələri» ixracı — CSV (bir cədvəl) və XLSX (bütün cədvəllər).

``GET /sorgu/neticeler/ixrac/<csv|xlsx>/?er_*&dataset=…`` — cari filtrlərin
AQREQATLARI (bax ``results_datasets``); xam cavab/şərh YOXDUR. Hər ixrac audit
jurnalına yazılır (``AuditAction.EXPORT``, ``resource_type="surveys.results"``,
filtrlər və sətir sayı ilə). Mətn xanaları formula-neytrallaşdırılır
(``core.export_safety``). İcazə FAIL-CLOSED — əhatəsiz istifadəçi 403 alır; davam edən
kampaniya 409 (canlı nəticə yoxdur, M-1). Saylar ekranda olduğu kimi səbətlədir.
"""

from __future__ import annotations

import io

from django.contrib.auth.decorators import login_required
from django.http import Http404, HttpResponse, JsonResponse
from django.utils import timezone
from django.utils.translation import pgettext
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET

from core.export_safety import safe_csv_writer, sheet_append

from .results_datasets import CSV_DATASETS, DATASETS, build_tables
from .results_filters import effective_params, resolve
from .results_panel import scope_label

CTX = "surveys.results"

AUDIT_RESOURCE = "surveys.results"
_HEADER_FILL = "2563EB"  # --ems-primary-600
_HEADER_FONT = "FFFFFF"


def _filename(fmt, dataset="") -> str:
    stamp = timezone.localdate().isoformat()
    suffix = f"-{dataset}" if dataset else ""
    return f"sorgu-neticeleri{suffix}-{stamp}.{fmt}"


def _csv_bytes(table) -> bytes:
    buffer = io.StringIO()
    buffer.write("\ufeff")  # BOM — Excel UTF-8 Azərbaycan hərflərini düzgün oxusun
    writer = safe_csv_writer(buffer)
    writer.writerow(table["header"])
    for row in table["rows"]:
        writer.writerow(["" if value is None else value for value in row])
    return buffer.getvalue().encode("utf-8")


def _xlsx_bytes(tables) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    fill = PatternFill("solid", fgColor=_HEADER_FILL)
    font = Font(name="Calibri", color=_HEADER_FONT, bold=True)
    for index, table in enumerate(tables):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = str(table["title"])[:31]
        sheet_append(sheet, table["header"])
        for cell in sheet[1]:
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in table["rows"]:
            sheet_append(sheet, ["" if value is None else value for value in row])
        sheet.freeze_panes = "A2"
        for column, header in enumerate(table["header"], start=1):
            values = [header, *(row[column - 1] for row in table["rows"] if column - 1 < len(row))]
            width = min(max(len(str(value or "")) for value in values) + 2, 60)
            sheet.column_dimensions[get_column_letter(column)].width = max(width, 6)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _audit(request, resolved, fmt, dataset, rows):
    from core.audit import log_action
    from core.constants import AuditAction

    log_action(
        AuditAction.EXPORT,
        user=request.user,
        organization=resolved.organization,
        request=request,
        resource_type=AUDIT_RESOURCE,
        resource_repr=f"{fmt.upper()} · {dataset or 'all'}",
        reason=pgettext(CTX, "Sorğu nəticələrinin ixracı (aqreqatlar)"),
        new_values={
            "format": fmt,
            "dataset": dataset or "all",
            "period": resolved.query.period.label,
            "filters": effective_params(resolved.query, resolved.filters),
            "rows": rows,
        },
    )


@never_cache
@login_required
@require_GET
def export(request, fmt):
    if fmt not in ("csv", "xlsx"):
        raise Http404
    resolved = resolve(request, with_choices=True)
    if resolved is None:
        return JsonResponse({"ok": False, "error": "forbidden"}, status=403)
    if not resolved.campaign_ids:
        raise Http404
    if resolved.live:
        # M-1: davam edən kampaniyanın nəticəsi (o cümlədən ixracı) yoxdur.
        return JsonResponse({"ok": False, "error": "campaign_open"}, status=409)
    label = scope_label(resolved.scope)
    if fmt == "csv":
        dataset = request.GET.get("dataset") or "teachers"
        if dataset not in CSV_DATASETS:
            raise Http404
        table = build_tables(resolved, resolved.filters, [dataset], scope_label=label)[0]
        payload, content_type = _csv_bytes(table), "text/csv; charset=utf-8"
        rows = len(table["rows"])
    else:
        dataset = ""
        tables = build_tables(resolved, resolved.filters, DATASETS, scope_label=label)
        payload = _xlsx_bytes(tables)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        rows = sum(len(table["rows"]) for table in tables)
    _audit(request, resolved, fmt, dataset, rows)
    response = HttpResponse(payload, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{_filename(fmt, dataset)}"'
    response["X-Content-Type-Options"] = "nosniff"
    return response
